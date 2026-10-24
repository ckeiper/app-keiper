import unittest
from decimal import Decimal
from ledger_arithmetic import empty,accumulate,render,render_groups

class LedgerGroupsTests(unittest.TestCase):
 def test_known_types_are_exact_normalized_and_unknown_stays_other(self):
  r=empty(True)
  for t in ['Venda','  ANÚNCIOS MClicks ','Armazenamento Full','Coleta Full','Envio','Cobrança de DIFAL','Pagamento de faturas','E-Shop','Ajuste desconhecido',None,'=x']:
   accumulate(r,100,10,90,t)
  self.assertEqual({k:v['valid_rows'] for k,v in r['groups'].items()},dict(vendas=1,ads=1,full=2,envios=1,tributos=1,pagamentos=1,outros=4))
 def test_all_group_counts_and_amounts_reconcile_to_same_total(self):
  r=empty(True)
  for t,row in [('Venda',(100,10,90)),('Venda',(-100,-10,-90)),('Anúncios MClicks',(0,10,-10)),('Envio',(100,10,'89.99')),('E-Shop',(None,1,1))]:accumulate(r,*row,t)
  for key in ['valid_rows','invalid_rows','inconsistent_rows']:
   self.assertEqual(sum(g[key] for g in r['groups'].values()),r[key])
  for key in ['gross_total','fees_total','net_total','difference_total']:
   self.assertEqual(sum(Decimal(g[key]) for g in r['groups'].values()),Decimal(r[key]))
 def test_no_valid_rows_remain_unavailable_in_table(self):
  r=empty(True);accumulate(r,None,0,0,'Venda')
  html=render(r);self.assertIn('Movimentos por grupo',html)
  self.assertEqual(html.count('<td>Indisponível</td>'),3)
 def test_unknown_input_cannot_inject_or_grow_categories(self):
  r=empty(True)
  for i in range(100):accumulate(r,1,0,1,'<script>'+str(i))
  self.assertEqual(set(r['groups']),{'outros'})
  self.assertNotIn('<script>',render_groups(r))
 def test_legacy_has_no_fabricated_groups(self):
  self.assertEqual(render_groups(dict(valid_rows=10)),'')
 def test_parser_passes_type_and_preserves_invalid_group(self):
  import test_freight_reconciliation as fixtures
  from freight_reconciliation import FIELDS,inspect_sales_reconciliation
  from openpyxl import load_workbook
  from io import BytesIO
  payload=fixtures.SalesReconciliationTests().workbook([dict(operation_type='Venda',total_fees=10,net_after_fees=90),dict(operation_type='Anúncios MClicks',total_fees=None,net_after_fees=90)])
  w=load_workbook(BytesIO(payload));w.active.cell(3,len(FIELDS)+1,'Valor bruto (=)')
  w.active.cell(4,len(FIELDS)+1,100);w.active.cell(5,len(FIELDS)+1,100)
  out=BytesIO();w.save(out)
  r=inspect_sales_reconciliation('test.xlsx',out.getvalue())['financial']
  self.assertEqual(r['groups']['vendas']['valid_rows'],1)
  self.assertEqual(r['groups']['ads']['invalid_rows'],1)

 def test_bounded_review_uses_excel_line_and_never_copies_raw_values(self):
  r=empty(True)
  for i in range(30):accumulate(r,'=PRIVATE',10,90,'Venda',i+4)
  self.assertEqual(len(r['review_rows']),20)
  self.assertEqual(r['review_rows'][0]['line'],4)
  self.assertNotIn('PRIVATE',str(r))
  self.assertIn('Bruto',r['review_rows'][0]['reason'])
 def test_parser_line_numbers_include_empty_rows(self):
  import test_freight_reconciliation as fixtures
  from freight_reconciliation import FIELDS,inspect_sales_reconciliation
  from openpyxl import load_workbook
  from io import BytesIO
  payload=fixtures.SalesReconciliationTests().workbook([dict(operation_type='Venda',total_fees=10,net_after_fees=90)])
  w=load_workbook(BytesIO(payload));w.active.cell(3,len(FIELDS)+1,'Valor bruto (=)')
  w.active.insert_rows(4);w.active.cell(5,len(FIELDS)+1,100);w.active.cell(5,list(FIELDS).index('net_after_fees')+1,89)
  out=BytesIO();w.save(out)
  r=inspect_sales_reconciliation('test.xlsx',out.getvalue())['financial']
  self.assertEqual(r['review_rows'][0]['line'],5)
