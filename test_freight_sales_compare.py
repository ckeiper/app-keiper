import unittest
from freight_reconciliation import FIELDS, schema_columns
from freight_sales_compare import compare_sales_rows, audit_against_sales
from unittest.mock import Mock, patch
import contextlib


class SalesFreightComparisonTests(unittest.TestCase):
    def compare(self, movements, refs=None, payments=None, shipments=None):
        rows=[]
        for values in movements:
            row=dict(operation_type='Venda',operation_id='payment',shipment_id='shipment',
                     gross_shipping=20,shipping_discount=5,seller_shipping=15)
            row.update(values)
            rows.append([row.get(key) for key in FIELDS])
        return compare_sales_rows(rows,schema_columns(list(FIELDS.values())),
            refs if refs is not None else {'order':dict(frete=12,frete_fonte='estimado',frete_verificavel=True)},
            payments if payments is not None else [dict(operation_id='payment',pedido_id='order')],
            shipments or [])

    def test_exact_payment_link_and_estimated_source(self):
        result=self.compare([{}])
        self.assertEqual(result['above'],1)
        self.assertEqual(result['estimated_orders'],1)
        self.assertEqual(result['above_total'],'3.00')
        self.assertNotIn('payment',str(result))
        self.assertNotIn('shipment',str(result))

    def test_shipment_link_and_actual_zero_are_valid(self):
        result=self.compare([dict(gross_shipping=0,shipping_discount=0,seller_shipping=0)],
            refs={'order':dict(frete=0,frete_fonte='real',frete_verificavel=True)},payments=[],
            shipments=[dict(envio_id='shipment',pedido_id='order')])
        self.assertEqual(result['within'],1)
        self.assertEqual(result['real_orders'],1)

    def test_exact_order_id_and_conflicting_payment_are_checked_together(self):
        result=self.compare([dict(operation_id='order')],payments=[])
        self.assertEqual(result['compared_orders'],1)
        result=self.compare([dict(operation_id='order')],
                            payments=[dict(operation_id='order',pedido_id='other')])
        self.assertEqual(result['compared_orders'],0)
        self.assertEqual(result['ambiguous_rows'],1)

    def test_conflicting_links_and_multiple_movements_are_never_summed(self):
        result=self.compare([{},{}])
        self.assertEqual(result['ambiguous_orders'],1)
        self.assertEqual(result['compared_orders'],0)
        result=self.compare([{}],shipments=[dict(envio_id='shipment',pedido_id='other')])
        self.assertEqual(result['ambiguous_rows'],1)
        self.assertEqual(result['compared_orders'],0)

    def test_missing_sources_invalid_charges_and_non_sales_are_separate(self):
        result=self.compare([{}],refs={'order':dict(frete=0,frete_fonte='estimado',frete_verificavel=False)})
        self.assertEqual(result['unavailable_orders'],1)
        self.assertEqual(self.compare([{}],payments=[])['unmatched_rows'],1)
        for values in [dict(seller_shipping=None),dict(seller_shipping='=15'),
                       dict(seller_shipping=16),dict(gross_shipping=-2)]:
            self.assertEqual(self.compare([values])['invalid_orders'],1)
        self.assertEqual(self.compare([dict(operation_type='Anúncios MClicks')])['candidate_rows'],0)

    def test_wrapper_uses_store_scope_and_sales_formula_beyond_latest_100(self):
        from test_freight_reconciliation import SalesReconciliationTests
        payload=SalesReconciliationTests().workbook([dict(operation_date='01/10/2026',
                                                        operation_type='Venda')])
        calculate=Mock(return_value=(None,[dict(pedido_id='synthetic-order',frete=15,
                                                frete_fonte='real',frete_verificavel=True)]))
        query=Mock(side_effect=[[dict(pedido_id='synthetic-order',operation_id='synthetic-operation')],[]])
        result=audit_against_sales(5,'synthetic.xlsx',payload,{'period_start':'2026-10-01'},
                                   lambda:contextlib.nullcontext(None),query,calculate)
        calculate.assert_called_once_with(5,inicio='2026-10-01')
        self.assertTrue(all(call.args[2][0]==5 for call in query.call_args_list))
        self.assertEqual(result['within'],1)

    def test_sales_marks_missing_weights_unavailable_but_real_zero_verified(self):
        from test_listing_details import main_function
        for shipping, verified in [([],False),([dict(pedido_id='order',custo_frete=0)],True)]:
            query=Mock(side_effect=[
                [dict(aliquota_imposto=None,comissao_pct=None)],
                [dict(pedido_id='order',total=100,fechado_em='2026-10-01',status='paid')],
                [dict(pedido_id='order',quantidade=1,preco_unitario=100,custo=None,peso=None)],
                shipping])
            env=dict(banco=lambda:contextlib.nullcontext(None),consultar=query,
                     frete_esperado=lambda price,weight:None)
            calculate=main_function('calcular_vendas',env)
            _, rows=calculate(5,inicio='2026-10-01')
            self.assertEqual(rows[0]['frete_verificavel'],verified)

    def test_shared_bundle_does_not_reopen_xlsx_for_sales_comparison(self):
        from openpyxl import load_workbook
        from freight_reconciliation import inspect_sales_reconciliation
        from test_freight_reconciliation import SalesReconciliationTests
        payload=SalesReconciliationTests().workbook([dict(operation_date='01/10/2026',
                                                        operation_type='Venda')])
        calculate=Mock(return_value=(None,[dict(pedido_id='synthetic-operation',frete=15,
                                                frete_fonte='real',frete_verificavel=True)]))
        with patch('openpyxl.load_workbook',wraps=load_workbook) as reader:
            summary,preview,rows,columns=inspect_sales_reconciliation('synthetic.xlsx',payload,bundle=True)
            result=audit_against_sales(5,'synthetic.xlsx',payload,summary,
                lambda:contextlib.nullcontext(None),Mock(return_value=[]),calculate,
                parsed_rows=rows,parsed_columns=columns)
        self.assertEqual(reader.call_count,1)
        self.assertEqual(result['within'],1)
        self.assertEqual(len(preview),1)
        self.assertNotIn('synthetic-operation',str(summary))
