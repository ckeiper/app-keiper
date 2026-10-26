import os, time, html, json, datetime, csv, io, threading, smtplib, secrets, hashlib
from email.mime.text import MIMEText
from functools import wraps
from urllib.parse import urlencode
from flask import Flask, request, redirect, jsonify, Response, session, send_from_directory, g
import requests
import psycopg2
from psycopg2.extras import RealDictCursor
from apscheduler.schedulers.background import BackgroundScheduler
from werkzeug.security import generate_password_hash, check_password_hash
from account_mfa import init_mfa_schema, register_mfa, issue_challenge
from ml_token_crypto import seal_token, reveal_token, migrate_tokens, validate_token_config, enabled as token_encryption_enabled
from session_security import init_session_schema, install_session_guard
from web_security import session_key, install_security, init_security_schema, auth_attempt_allowed
from listing_details import (registered_weights, legacy_weight, render_weights,
                             render_promotions, active_promotions, read_json, get_json,
                             render_weight_column, effective_promotion)
from listing_sync import sync_promotions
from full_monitor import init_full_schema, full_metadata, sync_full
from full_routes import register_full
from competitor_monitor import init_competitor_schema, collect_due
from competitor_routes import register_competitors, tabs as listing_tabs
from listing_diagnostics import inspect_listing
from listing_health import diagnose_listing
from listing_sales import listing_sales_summary
from listing_trends import comparison
from listing_ranking import ranking_metrics
from listing_kpis import listing_kpis
from listing_opportunities import listing_opportunities, opportunity_actions
from freight_import import (preview_reconciliation_upload, render_reconciliation_preview,
                            PreviewError, MAX_UPLOAD_BYTES, render_reconciliation_column_hints)
from freight_comparison import render_freight_comparison_demo
from freight_reconciliation import inspect_sales_reconciliation, render_sales_reconciliation
from freight_store import (init_freight_schema, save_report, report_history, render_history,
                           consolidate_report, consolidated_history, render_consolidated)
from freight_sales_compare import audit_against_sales, render_sales_comparison
from visual_theme import VISUAL_CSS
from crm import init_crm_schema, register_crm, HOME_WIDGET
from page_performance import ranking_query_bounds
from sales_sync import init_sales_schema, sync_orders
from overview_store import init_overview_schema
from sync_runner import run_stages
from sync_exclusion import store_sync_lock
from listing_pagination import listing_page
from sales_view import render_sales
from overview import load_overview, render_overview, overview_layout
from management_routes import register_management
try:
    import openpyxl
except Exception:
    openpyxl = None

CLIENT_ID = os.getenv("ML_CLIENT_ID")
CLIENT_SECRET = os.getenv("ML_CLIENT_SECRET")
REDIRECT_URI = os.getenv("ML_REDIRECT_URI")
DATABASE_URL = os.getenv("DATABASE_URL")
SECRET_KEY = session_key()
validate_token_config()

API = "https://api.mercadolibre.com"
AUTH_URL = "https://auth.mercadolivre.com.br/authorization"
TOKEN_URL = "https://api.mercadolibre.com/oauth/token"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
LOGO_PATH = os.path.join(BASE_DIR, "logo_Keiper_Consultoria.png")

app = Flask(__name__)
app.secret_key = SECRET_KEY
install_session_guard(app, lambda uid: carregar_usuario_sessao(uid))
install_security(app)
app.logger.warning("ml_token_encryption_configuration enabled=%s key_present=%s",
                   token_encryption_enabled(), bool(os.getenv("ML_TOKEN_ENCRYPTION_KEY")))
app.logger.warning("mfa_configuration key_present=%s", bool(os.getenv("MFA_ENCRYPTION_KEY")))
DUMMY_PASSWORD_HASH = generate_password_hash(secrets.token_urlsafe(32))

@app.before_request
def registrar_inicio_requisicao():
    g.inicio_pagina = time.perf_counter()


@app.after_request
def registrar_pagina_lenta(response):
    """Only log slow endpoint names; no URLs, tokens or personal data."""
    inicio = getattr(g, "inicio_pagina", None)
    if inicio is not None:
        duracao_ms = (time.perf_counter() - inicio) * 1000
        if duracao_ms >= 750:
            app.logger.warning(
                "pagina_lenta endpoint=%s duracao_ms=%.0f status=%s",
                request.endpoint or "desconhecido", duracao_ms,
                response.status_code)
    return response
esc = html.escape

ATUALIZANDO = set()
LOCK_ATUALIZACAO = threading.Lock()
GATILHO_AUTO_SEG = 1800

FRETE_PRECO_BANDAS = [
    (0, 18.99), (19, 48.99), (49, 78.99), (79, 99.99),
    (100, 119.99), (120, 149.99), (150, 199.99), (200, 10**9)
]
FRETE_PESO_BANDAS = [
    (0, 0.3), (0.3, 0.5), (0.5, 1), (1, 1.5), (1.5, 2), (2, 3), (3, 4), (4, 5),
    (5, 6), (6, 7), (7, 8), (8, 9), (9, 10), (10, 11), (11, 13), (13, 15),
    (15, 17), (17, 20), (20, 25), (25, 30), (30, 40), (40, 50), (50, 60),
    (60, 70), (70, 80), (80, 90), (90, 100), (100, 125), (125, 150), (150, 10**9)
]
FRETE_TABELA = [
    [5.65, 6.85, 8.15, 12.95, 14.95, 16.95, 19.05, 21.65],
    [5.95, 6.95, 8.25, 13.85, 16.15, 18.15, 20.45, 23.25],
    [6.05, 7.15, 8.45, 14.45, 16.85, 19.05, 21.35, 24.45],
    [6.15, 7.35, 8.65, 14.75, 17.15, 19.45, 21.75, 25.45],
    [6.25, 7.45, 8.75, 15.05, 17.65, 19.85, 22.25, 25.55],
    [6.35, 8.65, 9.15, 16.45, 19.15, 21.65, 24.35, 27.05],
    [6.45, 8.75, 9.75, 17.85, 20.75, 23.35, 26.35, 29.25],
    [6.55, 8.85, 10.25, 19.75, 22.85, 26.05, 29.25, 32.45],
    [6.65, 8.95, 10.35, 25.95, 29.15, 33.35, 36.45, 40.85],
    [6.75, 9.05, 10.45, 27.55, 31.65, 36.75, 40.85, 45.25],
    [6.85, 9.25, 10.55, 29.45, 34.35, 39.25, 44.15, 49.35],
    [6.95, 9.35, 10.65, 30.25, 35.25, 40.35, 45.35, 50.75],
    [7.05, 9.45, 10.85, 38.25, 45.05, 51.95, 58.75, 65.85],
    [7.05, 9.65, 11.05, 41.65, 48.55, 55.45, 62.35, 69.35],
    [7.15, 10.05, 11.45, 42.55, 49.75, 56.85, 63.85, 70.95],
    [7.25, 10.25, 11.65, 45.55, 52.95, 60.55, 68.15, 75.65],
    [7.35, 10.45, 11.85, 48.95, 56.55, 64.05, 71.35, 79.35],
    [7.45, 10.65, 12.05, 55.15, 64.35, 73.55, 82.75, 91.95],
    [7.65, 11.05, 12.25, 64.55, 75.75, 85.45, 96.25, 106.85],
    [7.75, 11.25, 12.45, 66.45, 76.05, 86.25, 97.15, 107.85],
    [7.85, 11.45, 12.65, 68.35, 79.65, 89.75, 100.05, 107.95],
    [7.95, 11.65, 12.85, 70.95, 81.85, 92.85, 103.45, 111.65],
    [8.05, 11.85, 13.05, 75.55, 87.25, 99.05, 110.25, 119.05],
    [8.15, 12.05, 13.25, 80.95, 93.75, 105.95, 118.05, 127.45],
    [8.25, 12.25, 13.45, 84.65, 97.95, 110.75, 123.35, 133.15],
    [8.35, 12.45, 13.65, 94.05, 108.35, 122.95, 136.95, 147.85],
    [8.45, 12.65, 13.85, 107.45, 124.85, 140.45, 156.45, 168.85],
    [8.55, 12.85, 14.05, 120.15, 138.95, 156.95, 174.85, 188.85],
    [8.65, 12.85, 14.25, 127.45, 147.05, 166.55, 185.55, 200.35],
    [8.75, 12.85, 14.45, 167.05, 193.35, 218.45, 243.45, 262.85],
]

def frete_esperado(preco, peso_kg):
    if preco is None or peso_kg is None or peso_kg <= 0:
        return None
    if preco < 19:
        return round(min(preco / 2.0, 5.65), 2)
    col = None
    for i, (lo, hi) in enumerate(FRETE_PRECO_BANDAS):
        if lo <= preco < hi:
            col = i
            break
    if col is None:
        col = len(FRETE_PRECO_BANDAS) - 1
    linha = None
    for i, (lo, hi) in enumerate(FRETE_PESO_BANDAS):
        if lo <= peso_kg < hi:
            linha = i
            break
    if linha is None:
        linha = len(FRETE_PESO_BANDAS) - 1
    return FRETE_TABELA[linha][col]


TAGS_BOAS = {"good_quality_picture", "real_size", "immediate_delivery", "catalog_product",
             "good_quality_thumbnail", "fast_shipping", "free_shipping", "brand_quality"}
TAGS_RUINS = {"dragged_bids", "poor_quality_picture", "out_of_stock", "under_review",
              "suspended", "no_picture", "real_size_mismatch", "loses_rank", "low_quality"}


def montar_demandas_anuncio(a, em_promocao, conversao):
    demandas = []
    qtd = a["quantidade"] or 0
    fotos = a["fotos"] or 0
    if qtd <= 0:
        demandas.append(("Repor estoque (está zerado)", "urgente"))
    elif qtd < 5:
        demandas.append(("Repor estoque (baixo: " + str(qtd) + ")", "atencao"))
    if fotos <= 0:
        demandas.append(("Adicionar fotos (está sem nenhuma)", "urgente"))
    elif fotos < 10:
        demandas.append(("Adicionar " + str(10 - fotos) + " fotos (ideal 10)", "atencao"))
    clips = a["clips"]
    if clips is None:
        demandas.append(("Verificar se tem vídeo (clips não confirmado)", "atencao"))
    elif clips < 2:
        demandas.append(("Adicionar vídeo/clip (ideal 2)", "atencao"))
    tags = []
    try:
        tags = json.loads(a["tags"]) if a["tags"] else []
    except Exception:
        tags = []
    for t in tags:
        if t in TAGS_RUINS:
            demandas.append(("Resolver tag ruim: " + t, "urgente"))
    st = str(a["status"])
    if st not in ("active", "ACTIVE"):
        demandas.append(("Reativar anúncio (status: " + st + ")", "urgente"))
    if conversao is not None and conversao < 5 and (a.get("visitas") or 0) > 0:
        demandas.append(("Melhorar conversão (%.2f%% com %d visitas) — revise preço, título e fotos" % (conversao, a.get("visitas") or 0), "atencao"))
    return demandas


def img_logo(altura=34):
    if os.path.exists(LOGO_PATH):
        return ("<img src='/logo.png' alt='Keiper Consultoria' style='height:"
                + str(altura) + "px;vertical-align:middle;"
                "margin-right:10px;border-radius:8px'>")
    return ""


CSS = """
<style>
* { margin:0; padding:0; box-sizing:border-box; font-family:'Segoe UI',Arial,sans-serif; }
body { background:#f4f6f9; color:#1a1a2e; }
.top { background:#3483FA; color:#fff; padding:18px 32px; display:flex; align-items:center; justify-content:space-between; flex-wrap:wrap; gap:10px; }
.top h1 { font-size:22px; font-weight:700; display:flex; align-items:center; }
.top .brand { font-size:14px; opacity:.9; }
.top a.btn { background:#FFE600; color:#1a1a2e; text-decoration:none; padding:9px 18px; border-radius:8px; font-weight:600; font-size:14px; }
.top a.btn.sair { background:#d64545; color:#fff; }
.wrap { max-width:1300px; margin:28px auto; padding:0 20px; }
.layout { display:flex; gap:20px; align-items:flex-start; }
.menu-lateral { width:210px; flex-shrink:0; background:#fff; border-radius:14px; box-shadow:0 2px 10px rgba(0,0,0,.06); padding:14px; position:sticky; top:16px; }
.menu-lateral .voltar { display:block; font-size:13px; color:#6b7280; text-decoration:none; margin-bottom:10px; font-weight:600; }
.menu-lateral a.item { display:block; padding:9px 12px; border-radius:8px; color:#1a1a2e; text-decoration:none; font-size:14px; font-weight:600; margin-bottom:2px; }
.menu-lateral a.item:hover { background:#f0f4ff; }
.menu-lateral a.item.atual { background:#3483FA; color:#fff; }
.menu-lateral a.item.perigo { color:#d64545; }
.conteudo { flex:1; min-width:0; }
@media (max-width:820px) { .layout { flex-direction:column; } .menu-lateral { width:100%; position:static; } }
.card { background:#fff; border-radius:14px; box-shadow:0 2px 10px rgba(0,0,0,.06); padding:22px; margin-bottom:22px; }
.grid { display:grid; grid-template-columns:repeat(auto-fit,minmax(200px,1fr)); gap:16px; }
.stat { background:#fff; border-radius:14px; box-shadow:0 2px 10px rgba(0,0,0,.06); padding:18px; border-left:4px solid #3483FA; }
.stat .num { font-size:24px; font-weight:700; color:#3483FA; }
.stat .lab { font-size:13px; color:#6b7280; margin-top:4px; }
.stat .var { font-size:12px; font-weight:600; margin-top:6px; }
.var.subiu { color:#1a9d5c; }
.var.caiu { color:#d64545; }
.var.estavel { color:#6b7280; }
table { width:100%; border-collapse:collapse; margin-top:12px; font-size:13.5px; }
th { background:#f0f4ff; color:#1a1a2e; text-align:left; padding:10px 12px; font-weight:600; white-space:nowrap; }
td { padding:9px 12px; border-bottom:1px solid #eef1f5; vertical-align:top; }
tr:hover td { background:#fafbff; }
.tag { display:inline-block; padding:3px 10px; border-radius:20px; font-size:12px; font-weight:600; }
.tag.ativa { background:#e6f7ee; color:#1a9d5c; }
.tag.pausado { background:#fdeaea; color:#d64545; }
.tag.expirada { background:#fdeaea; color:#d64545; }
.tag.normal { background:#eef1f5; color:#6b7280; }
.tag.respondida { background:#e6f7ee; color:#1a9d5c; }
.tag.pendente { background:#fff7e0; color:#8a6d00; }
.tag.boa { background:#e6f7ee; color:#1a9d5c; }
.tag.ruim { background:#fdeaea; color:#d64545; }
.tag.promocao { background:#FFE600; color:#1a1a2e; }
.tag.sem { background:#eef1f5; color:#9aa1ad; }
.alerta { padding:2px 8px; border-radius:6px; font-size:12px; font-weight:600; }
.alerta.urgente { background:#fdeaea; color:#d64545; }
.alerta.atencao { background:#fff7e0; color:#8a6d00; }
.alerta.ok { background:#e6f7ee; color:#1a9d5c; }
.alerta.subiu { background:#e6f7ee; color:#1a9d5c; }
.alerta.caiu { background:#fdeaea; color:#d64545; }
.alerta.estavel { background:#eef1f5; color:#6b7280; }
h2 { font-size:18px; color:#1a1a2e; margin-bottom:4px; }
.sub { color:#6b7280; font-size:13px; margin-bottom:14px; }
a.link { color:#3483FA; text-decoration:none; font-weight:600; }
a.link:hover { text-decoration:underline; }
.icon-link { text-decoration:none; font-size:16px; margin-left:6px; }
.input-inline { width:220px; max-width:60%; padding:6px 10px; border:1px solid #d5dbe5; border-radius:6px; font-size:15px; }
.btn-pequeno { background:#3483FA; color:#fff; border:0; border-radius:6px; padding:7px 14px; font-size:13px; font-weight:600; cursor:pointer; margin-left:6px; }
.badge { background:#FFE600; color:#1a1a2e; border-radius:20px; padding:2px 10px; font-size:12px; font-weight:700; }
.aviso { background:#fff7e0; border:1px solid #ffe28a; color:#8a6d00; border-radius:10px; padding:12px 16px; font-size:14px; margin-bottom:14px; }
.sucesso { background:#e6f7ee; border:1px solid #b7e4c7; color:#1a9d5c; border-radius:10px; padding:12px 16px; font-size:14px; margin-bottom:14px; }
.muted { color:#6b7280; font-size:12px; margin-top:14px; }
.btn-danger { display:inline-block; background:#d64545; color:#fff; text-decoration:none; padding:6px 14px; border-radius:8px; font-size:13px; font-weight:600; }
.btn-danger:hover { background:#b93a3a; }
.btn-salvar { background:#3483FA; color:#fff; border:0; border-radius:8px; padding:12px 20px; font-size:15px; font-weight:600; cursor:pointer; }
.btn-acoes { display:inline-block; background:#3483FA; color:#fff; text-decoration:none; padding:9px 16px; border-radius:8px; font-size:13px; font-weight:600; margin-right:8px; }
.btn-acoes:hover { background:#2a6fd6; }
.btn-acoes.verde { background:#1a9d5c; }
.btn-acoes.verde:hover { background:#15804a; }
.btn-acoes.laranja { background:#f59e0b; }
.btn-acoes.laranja:hover { background:#d97706; }
.btn-acoes.cinza { background:#6b7280; }
.btn-acoes.cinza:hover { background:#4b5563; }
.btn-acoes.rosa { background:#ec4899; }
.btn-acoes.rosa:hover { background:#db2777; }
.btn-acoes.roxo { background:#8b5cf6; }
.btn-acoes.roxo:hover { background:#7c3aed; }
.input { width:100%; padding:12px; border:1px solid #d5dbe5; border-radius:8px; font-size:15px; margin-bottom:12px; }
.filtro { display:inline-block; padding:6px 12px; border:1px solid #d5dbe5; border-radius:6px; font-size:13px; margin-right:8px; margin-bottom:8px; }
.filtro-btn { background:#3483FA; color:#fff; border:0; border-radius:6px; padding:7px 16px; font-size:13px; font-weight:600; cursor:pointer; }
.grafico-box { position:relative; height:300px; width:100%; }
.demanda-item { border:1px solid #eef1f5; border-radius:10px; padding:14px 16px; margin-bottom:12px; }
.demanda-item.urgente { border-left:4px solid #d64545; }
.demanda-item.atencao { border-left:4px solid #f59e0b; }
.demanda-item .titulo { font-size:15px; font-weight:700; margin-bottom:6px; }
.demanda-item .acao { padding:4px 0; font-size:13.5px; }
.demanda-item .acao.urgente { color:#d64545; }
.demanda-item .acao.atencao { color:#b45309; }
.promo-item { border:1px solid #eef1f5; border-radius:10px; padding:14px 16px; margin-bottom:12px; }
.promo-item.ativa { border-left:4px solid #1a9d5c; }
.promo-item.candidata { border-left:4px solid #f59e0b; }
.promo-item.finalizada { border-left:4px solid #9aa1ad; }
.promo-item .titulo { font-size:15px; font-weight:700; margin-bottom:6px; }
.promo-item .meta { font-size:12.5px; color:#6b7280; margin-bottom:8px; }
.promo-item table { margin-top:6px; }
.campanha-item { border:1px solid #eef1f5; border-radius:10px; padding:14px 16px; margin-bottom:12px; }
.campanha-item .titulo { font-size:15px; font-weight:700; margin-bottom:6px; }
.campanha-item .meta { font-size:12.5px; color:#6b7280; margin-bottom:8px; }
.campanha-item table { margin-top:6px; }
.barra-topo { display:flex; justify-content:space-between; align-items:center; flex-wrap:wrap; gap:10px; margin-bottom:14px; }
.barra-topo h2 { margin:0; }
/* A navegação das lojas deve permanecer em uma coluna, mesmo com o tema global. */
.layout>.menu-lateral{
  display:flex;flex-direction:column;align-items:stretch;flex-wrap:nowrap;
  gap:3px;min-width:0;
}
.layout>.menu-lateral a.item,.layout>.menu-lateral a.voltar{
  display:block;flex:0 0 auto;width:100%;max-width:100%;
  box-sizing:border-box;min-width:0;white-space:normal;
  overflow-wrap:anywhere;height:auto;min-height:0;
}
/* Navegação lateral persistente no desktop e gaveta lateral no celular. */
.menu-lateral { max-height:calc(100vh - 32px); overflow-y:auto; }
.sidebar-toggle, .sidebar-overlay { display:none; }
@media (max-width:820px) {
  .layout { display:flex; flex-direction:row; }
  .layout .conteudo { width:100%; min-width:0; }
  .layout .menu-lateral {
    position:fixed; top:0; left:0; bottom:0; z-index:1102;
    width:min(280px, 82vw) !important; max-height:100dvh;
    border-radius:0; overflow-y:auto; box-sizing:border-box;
    transform:translateX(-105%); transition:transform .2s ease;
  }
  body.sidebar-open .layout .menu-lateral { transform:translateX(0); }
  .sidebar-toggle {
    display:inline-flex; align-items:center; gap:7px; cursor:pointer;
    border:0; border-radius:8px; padding:9px 12px;
    background:#3483FA; color:#fff; font-weight:700; font-size:14px;
  }
  .sidebar-overlay {
    position:fixed; inset:0; z-index:1101; background:rgba(0,0,0,.5);
  }
  body.sidebar-open .sidebar-overlay { display:block; }
  body.sidebar-open { overflow:hidden; }
  .top { flex-wrap:wrap; }
}
@media print { .top, .menu-lateral, .btn-acoes { display:none !important; } .layout { display:block; } body { background:#fff; } }
</style>
"""

SCRIPT = """
<script>
function editarNome(id){
  var span = document.getElementById('nome-' + id);
  if (!span || span.querySelector('input')) { return; }
  var atual = span.textContent.trim();
  var caixa = span.parentNode;
  var input = document.createElement('input');
  input.value = atual;
  input.className = 'input-inline';
  input.maxLength = 120;
  var btn = document.createElement('button');
  btn.textContent = 'Salvar';
  btn.className = 'btn-pequeno';
  span.style.display = 'none';
  caixa.appendChild(input);
  caixa.appendChild(btn);
  input.focus();
  function salvar(){
    var novo = input.value.trim();
    if (!novo) { alert('O nome não pode ficar vazio.'); return; }
    fetch('/renomear/' + id, {
      method: 'POST',
      headers: {'Content-Type': 'application/x-www-form-urlencoded'},
      body: 'cliente=' + encodeURIComponent(novo)
    }).then(function(r){ return r.json(); }).then(function(d){
      if (d.ok) {
        span.textContent = novo;
        span.style.display = '';
        input.remove();
        btn.remove();
      } else { alert('Não foi possível salvar. Tente de novo.'); }
    }).catch(function(){ alert('Erro ao salvar. Tente de novo.'); });
  }
  btn.onclick = salvar;
  input.addEventListener('keydown', function(e){
    if (e.key === 'Enter') { salvar(); }
    if (e.key === 'Escape') { span.style.display = ''; input.remove(); btn.remove(); }
  });
}
function editarCusto(cid, itemId){
  var span = document.getElementById('custo-' + cid + '-' + itemId);
  if (!span || span.querySelector('input')) { return; }
  var atual = span.textContent.trim().replace('R$','').trim();
  var caixa = span.parentNode;
  var input = document.createElement('input');
  input.value = atual;
  input.className = 'input-inline';
  input.type = 'number';
  input.step = '0.01';
  input.style.width = '90px';
  var btn = document.createElement('button');
  btn.textContent = 'OK';
  btn.className = 'btn-pequeno';
  span.style.display = 'none';
  caixa.appendChild(input);
  caixa.appendChild(btn);
  input.focus();
  function salvar(){
    var novo = parseFloat(input.value.replace(',','.'));
    if (isNaN(novo) || novo < 0) { alert('Digite um custo válido.'); return; }
    fetch('/salvar_custo/' + cid + '/' + itemId, {
      method: 'POST',
      headers: {'Content-Type': 'application/x-www-form-urlencoded'},
      body: 'custo=' + novo
    }).then(function(r){ return r.json(); }).then(function(d){
      if (d.ok) {
        span.textContent = 'R$ ' + novo.toFixed(2);
        span.style.display = '';
        input.remove();
        btn.remove();
      } else { alert('Não foi possível salvar.'); }
    }).catch(function(){ alert('Erro ao salvar.'); });
  }
  btn.onclick = salvar;
  input.addEventListener('keydown', function(e){
    if (e.key === 'Enter') { salvar(); }
    if (e.key === 'Escape') { span.style.display = ''; input.remove(); btn.remove(); }
  });
}
function editarRoas(cid, campanhaId){
  var span = document.getElementById('roas-' + cid + '-' + campanhaId);
  if (!span || span.querySelector('input')) { return; }
  var atual = span.textContent.trim();
  var caixa = span.parentNode;
  var input = document.createElement('input');
  input.value = atual;
  input.className = 'input-inline';
  input.type = 'number';
  input.step = '0.1';
  input.style.width = '80px';
  var btn = document.createElement('button');
  btn.textContent = 'OK';
  btn.className = 'btn-pequeno';
  span.style.display = 'none';
  caixa.appendChild(input);
  caixa.appendChild(btn);
  input.focus();
  function salvar(){
    var novo = parseFloat(input.value.replace(',','.'));
    if (isNaN(novo) || novo < 0) { alert('Digite um ROAS válido.'); return; }
    fetch('/salvar_roas/' + cid + '/' + campanhaId, {
      method: 'POST',
      headers: {'Content-Type': 'application/x-www-form-urlencoded'},
      body: 'roas=' + novo
    }).then(function(r){ return r.json(); }).then(function(d){
      if (d.ok) {
        span.textContent = novo.toFixed(1);
        span.style.display = '';
        input.remove();
        btn.remove();
      } else { alert('Não foi possível salvar.'); }
    }).catch(function(){ alert('Erro ao salvar.'); });
  }
  btn.onclick = salvar;
  input.addEventListener('keydown', function(e){
    if (e.key === 'Enter') { salvar(); }
    if (e.key === 'Escape') { span.style.display = ''; input.remove(); btn.remove(); }
  });
}
function gerarSenha(){
  var campo = document.getElementById('campo-senha');
  if (!campo) { return; }
  var chars = 'abcdefghjkmnpqrstuvwxyzABCDEFGHJKMNPQRSTUVWXYZ23456789!@#$%';
  var s = '';
  var bytes = new Uint32Array(16);
  crypto.getRandomValues(bytes);
  for (var i = 0; i < bytes.length; i++) { s += chars.charAt(bytes[i] % chars.length); }
  campo.value = s;
}
</script>
</script>
<script>
document.addEventListener("DOMContentLoaded", function () {
  var cel = document.querySelectorAll("td, .num");
  for (var i = 0; i < cel.length; i++) {
    var t = cel[i].textContent.trim();
    var ehDinheiro = t.indexOf("R$") === 0;
    var s = t.replace("R$", "").trim();
    var n = null;
    if (/^\d+\.\d{2}$/.test(s)) { n = parseFloat(s); }
    else if (/^\d+$/.test(s) && !ehDinheiro) { n = parseInt(s, 10); }
    if (n === null) { continue; }
    var f = n.toLocaleString("pt-BR", {
      minimumFractionDigits: ehDinheiro ? 2 : 0,
      maximumFractionDigits: ehDinheiro ? 2 : 0
    });
    cel[i].textContent = (ehDinheiro ? "R$ " : "") + f;
  }
});
</script>
"""

def pagina(titulo, corpo):
    is_home = "<main class='kh-home'>" in corpo
    is_progress = "<main class='ap-progress'>" in corpo
    use_sidebar = is_home or is_progress
    topo = "<div class='top'><div><h1>" + img_logo(34) + "Keiper Consultoria</h1>"
    marca = "<div class='brand'>Painel de lojas Mercado Livre</div></div>"
    links = ""
    if session.get("tipo") == "admin" and not use_sidebar:
        links += "<a class='btn' href='/seguranca'>Segurança</a>"
    links += "<a class='btn' href='/conectar'>+ Conectar loja</a></div>"
    cab = "<!doctype html><html><head><meta charset='utf-8'>"
    cab += "<meta name='viewport' content='width=device-width,initial-scale=1'>"
    cab += "<title>" + esc(titulo) + "</title>" + CSS + VISUAL_CSS + "</head><body>"
    fim = "<div class='wrap'>" + corpo + "</div>" + SCRIPT + "</body></html>"
    sidebar = ""
    if use_sidebar:
        from home_sidebar import render_home_sidebar
        sidebar = render_home_sidebar(admin=session.get("tipo") == "admin", active="progress" if is_progress else "home")
    return cab + sidebar + topo + marca + links + fim


def pagina_loja(titulo, cid, cliente, secao, corpo):
    from section_refresh import SCOPES, render as render_refresh
    if secao in SCOPES and "class='k-refresh'" not in corpo:
        corpo = render_refresh(cid, secao) + corpo
    from store_sidebar import render_store_sidebar
    sidebar = render_store_sidebar(cid, cliente, secao, admin=session.get("tipo") == "admin")
    cab = "<!doctype html><html lang='pt-BR'><head><meta charset='utf-8'>"
    cab += "<meta name='viewport' content='width=device-width,initial-scale=1'>"
    cab += "<title>" + esc(titulo) + "</title>" + CSS + VISUAL_CSS + "</head><body>"
    cab += sidebar + "<header class='top'><div><h1>" + img_logo(34) + "Keiper Consultoria</h1></div>"
    cab += "<div class='top-store' title='" + esc(str(cliente), quote=True) + "'>" + esc(str(cliente)) + "</div>"
    cab += "<a class='btn sair' href='/logout'>Sair</a></header>"
    cab += "<div class='wrap'><div class='layout'><main class='conteudo'>" + corpo
    from chart_personalization import CONTROLS as chart_controls
    return cab + "</main></div></div>" + chart_controls + SCRIPT + "</body></html>"


FUSO_BR = datetime.timezone(datetime.timedelta(hours=-3))


def hora_br(ts):
    if not ts:
        return "nunca"
    return datetime.datetime.fromtimestamp(ts, FUSO_BR).strftime(
        "%d/%m/%Y %H:%M")


def barra_atualizar(cid, ultima, tipo="resumo"):
    from section_refresh import SCOPES, render as render_refresh
    if tipo in SCOPES:
        return render_refresh(cid, tipo, hora_br(ultima))
    ult_txt = hora_br(ultima)
    return ("<div style='display:flex;gap:12px;align-items:center;"
            "flex-wrap:wrap;margin-bottom:14px'>"
            "<a class='btn-acoes' href='/atualizar/" + tipo + "/" + str(cid) + "'>"
            "🔄 Atualizar agora</a>"
            "<span style='color:#6b7280;font-size:13px'>"
            "Última atualização: " + ult_txt + "</span>"
            + ("<span class='muted'>Atualização em andamento; recarregue em instantes.</span>"
               if cid in ATUALIZANDO else "") + "</div>")


def ultima_atualizacao(cid):
    with banco() as conn:
        ult = consultar(conn, "SELECT MAX(atualizado_em) AS ult FROM anuncios WHERE conexao_id=%s", (cid,))
    return ult[0]["ult"] if ult and ult[0]["ult"] else None


def banco():
    return psycopg2.connect(DATABASE_URL)


def carregar_usuario_sessao(uid):
    conn = banco()
    try:
        rows = consultar(conn, "SELECT id, tipo, conexao_id, auth_version FROM usuarios WHERE id=%s", (uid,))
        return rows[0] if rows else None
    finally:
        conn.close()


def consultar(conn, sql, params=None):
    cur = conn.cursor(cursor_factory=RealDictCursor)
    cur.execute(sql, params or ())
    return cur.fetchall()


def executar(conn, sql, params=None):
    cur = conn.cursor()
    cur.execute(sql, params or ())
    return cur


def iniciar_banco():
    from startup_schema import schema_executor
    execute = schema_executor(executar)
    with banco() as conn:
        init_security_schema(conn, execute)
        execute(conn, """CREATE TABLE IF NOT EXISTS conexoes (
            id SERIAL PRIMARY KEY,
            cliente TEXT NOT NULL,
            user_id BIGINT,
            access_token TEXT,
            refresh_token TEXT,
            expires_at BIGINT,
            status TEXT DEFAULT 'ativa'
        )""")
        execute(conn, "ALTER TABLE conexoes ADD COLUMN IF NOT EXISTS aliquota_imposto REAL DEFAULT 0")
        execute(conn, "ALTER TABLE conexoes ADD COLUMN IF NOT EXISTS comissao_pct REAL DEFAULT 12")
        execute(conn, "ALTER TABLE conexoes ADD COLUMN IF NOT EXISTS mc_meta_pct REAL")
        execute(conn, "ALTER TABLE conexoes ADD COLUMN IF NOT EXISTS mc_meta_atualizada_em BIGINT")
        execute(conn, """CREATE TABLE IF NOT EXISTS usuarios (
            id SERIAL PRIMARY KEY,
            username TEXT UNIQUE NOT NULL,
            senha_hash TEXT NOT NULL,
            nome TEXT,
            tipo TEXT DEFAULT 'cliente',
            conexao_id INTEGER,
            criado_em BIGINT
        )""")
        execute(conn, "ALTER TABLE usuarios ADD COLUMN IF NOT EXISTS email TEXT")
        init_session_schema(conn, execute)
        init_mfa_schema(conn, execute)
        admin = consultar(conn, "SELECT id FROM usuarios WHERE username=%s", ("admin",))
        if not admin:
            # Existing accounts are untouched. New installations require a secret.
            bootstrap_password = os.getenv("KEIPER_ADMIN_INITIAL_PASSWORD", "")
            if len(bootstrap_password) >= 16:
                execute(conn, "INSERT INTO usuarios (username, senha_hash, nome, tipo, criado_em) VALUES (%s,%s,%s,%s,%s)",
                         ("admin", generate_password_hash(bootstrap_password), "Administrador", "admin", int(time.time())))
            else:
                app.logger.error("Admin bootstrap skipped: configure KEIPER_ADMIN_INITIAL_PASSWORD (16+ chars).")
        execute(conn, """CREATE TABLE IF NOT EXISTS reset_tokens (
            token_hash TEXT PRIMARY KEY,
            usuario_id INTEGER,
            expira_em BIGINT
        )""")
        execute(conn, """CREATE TABLE IF NOT EXISTS dados_conta (
            conexao_id INTEGER PRIMARY KEY,
            user_id BIGINT, nickname TEXT, nome TEXT, sobrenome TEXT,
            reputacao TEXT, pontos INTEGER, atualizado_em BIGINT
        )""")
        execute(conn, """CREATE TABLE IF NOT EXISTS anuncios (
            id SERIAL PRIMARY KEY,
            conexao_id INTEGER, item_id TEXT, titulo TEXT, preco REAL,
            quantidade INTEGER, status TEXT, vendidos INTEGER, atualizado_em BIGINT,
            fotos INTEGER DEFAULT 0, clips INTEGER DEFAULT 0, tags TEXT, peso REAL,
            custo REAL, sku TEXT,
            UNIQUE(conexao_id, item_id)
        )""")
        execute(conn, "ALTER TABLE anuncios ADD COLUMN IF NOT EXISTS fotos INTEGER DEFAULT 0")
        execute(conn, "ALTER TABLE anuncios ADD COLUMN IF NOT EXISTS clips INTEGER DEFAULT 0")
        execute(conn, "ALTER TABLE anuncios ADD COLUMN IF NOT EXISTS tags TEXT")
        execute(conn, "ALTER TABLE anuncios ADD COLUMN IF NOT EXISTS peso REAL")
        execute(conn, "ALTER TABLE anuncios ADD COLUMN IF NOT EXISTS custo REAL")
        execute(conn, "ALTER TABLE anuncios ADD COLUMN IF NOT EXISTS sku TEXT")
        execute(conn, "ALTER TABLE anuncios ADD COLUMN IF NOT EXISTS original_price REAL")
        execute(conn, "ALTER TABLE anuncios ADD COLUMN IF NOT EXISTS thumbnail TEXT")
        execute(conn, "ALTER TABLE anuncios ADD COLUMN IF NOT EXISTS pesos_json TEXT")
        execute(conn, "ALTER TABLE anuncios ADD COLUMN IF NOT EXISTS promocoes_json TEXT")
        execute(conn, """CREATE TABLE IF NOT EXISTS pedidos (
            id SERIAL PRIMARY KEY,
            conexao_id INTEGER, pedido_id TEXT, status TEXT,
            total REAL, moeda TEXT, fechado_em TEXT, atualizado_em BIGINT,
            UNIQUE(conexao_id, pedido_id)
        )""")
        execute(conn, """CREATE TABLE IF NOT EXISTS pedidos_itens (
            conexao_id INTEGER, pedido_id TEXT, item_id TEXT,
            quantidade INTEGER, preco_unitario REAL, atualizado_em BIGINT,
            PRIMARY KEY (conexao_id, pedido_id, item_id)
        )""")
        init_sales_schema(conn, execute)
        from profitability_store import init_schema as init_profitability
        init_profitability(conn, execute)
        init_full_schema(conn, execute)
        init_competitor_schema(conn, execute)
        from offer_history import init_schema as init_offer_history
        init_offer_history(conn, execute)
        from listing_interventions import init_schema as init_interventions
        init_interventions(conn, execute)
        from listing_tasks import init_schema as init_listing_tasks
        init_listing_tasks(conn, execute)
        from listing_content import init_schema as init_listing_content
        init_listing_content(conn, execute)
        from automatic_history import init_schema as init_automatic_history
        init_automatic_history(conn, execute)
        from automatic_history import seed_existing
        # Backfilling legacy observations is optional. It must not prevent
        # Gunicorn from booting when a live collector holds conflicting locks.
        # Savepoints keep the rest of the schema transaction healthy after a
        # PostgreSQL deadlock or lock timeout.
        execute(conn, "SAVEPOINT automatic_history_backfill")
        try:
            execute(conn, "SET LOCAL lock_timeout = '3s'")
            lock = consultar(conn, "SELECT pg_try_advisory_xact_lock(783420, 1) AS acquired")
            if lock and lock[0]["acquired"]:
                seed_existing(conn, execute)
            else:
                app.logger.warning("automatic_history_seed_skipped concurrent_migration")
            execute(conn, "SET LOCAL lock_timeout = DEFAULT")
        except (psycopg2.errors.DeadlockDetected,
                psycopg2.errors.LockNotAvailable,
                psycopg2.errors.QueryCanceled) as exc:
            execute(conn, "ROLLBACK TO SAVEPOINT automatic_history_backfill")
            app.logger.warning("automatic_history_seed_deferred error=%s", type(exc).__name__)
        finally:
            execute(conn, "RELEASE SAVEPOINT automatic_history_backfill")
        init_freight_schema(conn, execute)
        execute(conn, """CREATE TABLE IF NOT EXISTS metricas (
            conexao_id INTEGER, item_id TEXT, vendidos INTEGER,
            visitas INTEGER, conversao REAL, atualizado_em BIGINT,
            PRIMARY KEY (conexao_id, item_id)
        )""")
        execute(conn, """CREATE TABLE IF NOT EXISTS desempenho_historico (
            conexao_id INTEGER, item_id TEXT, data TEXT,
            visitas INTEGER DEFAULT 0, vendidos INTEGER DEFAULT 0, atualizado_em BIGINT,
            PRIMARY KEY (conexao_id, item_id, data)
        )""")
        execute(conn, """CREATE TABLE IF NOT EXISTS perguntas (
            conexao_id INTEGER, pergunta_id TEXT, item_id TEXT, texto TEXT,
            status TEXT, resposta TEXT, data TEXT, atualizado_em BIGINT,
            PRIMARY KEY (conexao_id, pergunta_id)
        )""")
        execute(conn, "CREATE INDEX IF NOT EXISTS perguntas_listing_coleta ON perguntas(conexao_id,item_id,atualizado_em DESC NULLS LAST,pergunta_id DESC)")
        execute(conn, """CREATE TABLE IF NOT EXISTS envios (
            conexao_id INTEGER, envio_id TEXT, status TEXT, tracking TEXT,
            pedido_id TEXT, data TEXT, custo_frete REAL, atualizado_em BIGINT,
            PRIMARY KEY (conexao_id, envio_id)
        )""")
        execute(conn, "ALTER TABLE envios ADD COLUMN IF NOT EXISTS custo_frete REAL")
        execute(conn, """CREATE TABLE IF NOT EXISTS promocoes (
            conexao_id INTEGER, promocao_id TEXT, tipo TEXT, nome TEXT,
            status TEXT, inicio TEXT, fim TEXT, qtd_itens INTEGER, atualizado_em BIGINT,
            PRIMARY KEY (conexao_id, promocao_id)
        )""")
        execute(conn, """CREATE TABLE IF NOT EXISTS promocoes_itens (
            conexao_id INTEGER, item_id TEXT, promocao_id TEXT,
            preco_promocional REAL, atualizado_em BIGINT,
            PRIMARY KEY (conexao_id, item_id, promocao_id)
        )""")
        execute(conn, """CREATE TABLE IF NOT EXISTS promocoes_adesoes (
            conexao_id INTEGER NOT NULL,
            promocao_id TEXT NOT NULL,
            item_id TEXT NOT NULL,
            usuario_id INTEGER NOT NULL,
            preco DECIMAL(14,2) NOT NULL,
            margem_pct DECIMAL(7,2) NOT NULL,
            status TEXT NOT NULL,
            criado_em BIGINT NOT NULL,
            atualizado_em BIGINT NOT NULL,
            resposta TEXT,
            PRIMARY KEY(conexao_id,promocao_id,item_id)
        )""")
        execute(conn, "ALTER TABLE promocoes_itens ADD COLUMN IF NOT EXISTS preco_promocional REAL")
        execute(conn, """CREATE TABLE IF NOT EXISTS ads_campanhas (
            conexao_id INTEGER, campanha_id TEXT, nome TEXT, status TEXT,
            tipo TEXT, data_inicio TEXT, data_fim TEXT, orcamento REAL, atualizado_em BIGINT,
            PRIMARY KEY (conexao_id, campanha_id)
        )""")
        execute(conn, "ALTER TABLE ads_campanhas ADD COLUMN IF NOT EXISTS roas_objetivo REAL")
        from ads_change_history import CREATE_SQL, INDEX_SQL
        execute(conn, CREATE_SQL)
        execute(conn, INDEX_SQL)
        execute(conn, """CREATE TABLE IF NOT EXISTS ads_metricas (
            conexao_id INTEGER, ad_id TEXT, campanha_id TEXT, item_id TEXT,
            impressoes INTEGER, cliques INTEGER, ctr REAL, gasto REAL, atualizado_em BIGINT,
            PRIMARY KEY (conexao_id, ad_id)
        )""")
        execute(conn, """CREATE TABLE IF NOT EXISTS ads_dia (
            conexao_id INTEGER, data TEXT,
            gasto REAL, faturamento REAL, impressoes INTEGER, cliques INTEGER, atualizado_em BIGINT,
            PRIMARY KEY (conexao_id, data)
        )""")
        execute(conn, """CREATE TABLE IF NOT EXISTS ads_orcamento_esgotado (conexao_id INTEGER NOT NULL, campanha_id TEXT NOT NULL, data TEXT NOT NULL, detectado_em BIGINT NOT NULL, PRIMARY KEY(conexao_id,campanha_id,data))""")
        execute(conn, """CREATE TABLE IF NOT EXISTS ads_campanha_dia (
            conexao_id INTEGER, campanha_id TEXT, data TEXT,
            gasto REAL, faturamento REAL, atualizado_em BIGINT,
            PRIMARY KEY (conexao_id, campanha_id, data)
        )""")
        execute(conn, """CREATE TABLE IF NOT EXISTS erros (
            conexao_id INTEGER, categoria TEXT, mensagem TEXT, quando BIGINT,
            PRIMARY KEY (conexao_id, categoria)
        )""")
        execute(conn, "ALTER TABLE erros ADD COLUMN IF NOT EXISTS resolvido_em BIGINT")
        from ads_official_metrics import init_schema as init_ads_sales_schema
        init_ads_sales_schema(conn, execute)
        from section_refresh import init_schema as init_section_refresh
        init_section_refresh(conn, execute)
        init_overview_schema(conn, execute)
        from keiper_ai_store import init_schema as init_keiper_ai
        init_keiper_ai(conn, execute)
        from kyra_reviews import init_schema as init_kyra_reviews
        init_kyra_reviews(conn, execute)
        init_crm_schema(conn, execute)
        execute(conn, """CREATE TABLE IF NOT EXISTS cliente_onboarding_financeiro (
            usuario_id INTEGER PRIMARY KEY,
            conexao_id INTEGER NOT NULL,
            custos_importados_em BIGINT,
            imposto_confirmado_em BIGINT,
            concluido_em BIGINT
        )""")
        execute(conn, """CREATE TABLE IF NOT EXISTS cliente_prioridades_acessos (
            conexao_id INTEGER PRIMARY KEY,
            habilitado BOOLEAN NOT NULL DEFAULT FALSE,
            atualizado_em BIGINT NOT NULL
        )""")
        # Existing clients are enrolled when the portal is explicitly enabled;
        # this avoids abruptly locking long-standing accounts before rollout.
        execute(conn, """INSERT INTO cliente_onboarding_financeiro(usuario_id,conexao_id)
            SELECT u.id,u.conexao_id FROM usuarios u
            JOIN cliente_prioridades_acessos a ON a.conexao_id=u.conexao_id
            WHERE u.tipo='cliente' AND a.habilitado=TRUE
            ON CONFLICT(usuario_id) DO NOTHING""")

        # ---------- AUTENTICAÇÃO ----------

def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if "usuario_id" not in session:
            return redirect("/login")
        return f(*args, **kwargs)
    return wrapper


def admin_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if "usuario_id" not in session:
            return redirect("/login")
        if session.get("tipo") != "admin":
            return pagina("Acesso negado", "<div class='card'><h2>Acesso restrito</h2></div>"), 403
        return f(*args, **kwargs)
    return wrapper


def pode_ver_loja(cid):
    if session.get("tipo") == "admin":
        return True
    return session.get("conexao_id") == cid


def smtp_configurado():
    return bool(os.getenv("SMTP_HOST") and os.getenv("SMTP_USER") and os.getenv("SMTP_PASS"))


def enviar_email(destino, assunto, corpo_texto):
    host = os.getenv("SMTP_HOST")
    user = os.getenv("SMTP_USER")
    senha = os.getenv("SMTP_PASS")
    if not host or not user or not senha or not destino:
        return False
    porta = int(os.getenv("SMTP_PORT", "587"))
    de = os.getenv("SMTP_FROM", user)
    msg = MIMEText(corpo_texto, "plain", "utf-8")
    msg["Subject"] = assunto
    msg["From"] = de
    msg["To"] = destino
    try:
        with smtplib.SMTP(host, porta) as s:
            s.starttls()
            s.login(user, senha)
            s.sendmail(de, [destino], msg.as_string())
        return True
    except Exception:
        return False


def criar_token_reset(usuario_id):
    token = secrets.token_urlsafe(32)
    th = hashlib.sha256(token.encode()).hexdigest()
    with banco() as conn:
        executar(conn, "INSERT INTO reset_tokens (token_hash, usuario_id, expira_em) VALUES (%s,%s,%s)",
                 (th, usuario_id, int(time.time()) + 1800))
    return token


@app.route("/logo.png")
def logo_arquivo():
    return send_from_directory(BASE_DIR, "logo_Keiper_Consultoria.png")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = (request.form.get("username") or "").strip()
        senha = request.form.get("senha") or ""
        if len(username) > 200 or len(senha) > 256:
            return "Credenciais inválidas.", 400
        if not auth_attempt_allowed(banco, executar, "login", username):
            return "Muitas tentativas. Aguarde 15 minutos e tente novamente.", 429, {"Retry-After": "900"}
        with banco() as conn:
            u = consultar(conn, "SELECT * FROM usuarios WHERE username=%s", (username,))
        u = u[0] if u else None
        valid_password = check_password_hash(u["senha_hash"] if u else DUMMY_PASSWORD_HASH, senha)
        if u and valid_password:
            if u.get("mfa_enabled"):
                if not issue_challenge(u, banco, consultar, executar):
                    return redirect("/login")
                return redirect("/login/verificar")
            session.clear()
            session.permanent = True
            session["usuario_id"] = u["id"]
            session["username"] = u["username"]
            session["nome"] = u["nome"]
            session["tipo"] = u["tipo"]
            session["conexao_id"] = u["conexao_id"]
            session["auth_version"] = u["auth_version"]
            if u["tipo"] == "cliente" and u["conexao_id"]:
                with banco() as conn:
                    enrolled=consultar(conn,"""SELECT 1 FROM cliente_onboarding_financeiro
                        WHERE usuario_id=%s AND conexao_id=%s""",(u["id"],u["conexao_id"]))
                session['onboarding_requerido']=bool(enrolled)
            if u["tipo"] == "cliente" and u["conexao_id"]:
                return redirect("/painel/" + str(u["conexao_id"]))
            return redirect("/painel")
        corpo = ("<div class='card' style='max-width:420px;margin:60px auto;text-align:center'>"
                 + img_logo(64) + "<h2>Login</h2>"
                 + "<div class='aviso'>Usuário ou senha incorretos.</div>"
                 + form_login() + "</div>")
        return pagina("Login", corpo)
    corpo = ("<div class='card' style='max-width:420px;margin:60px auto;text-align:center'>"
             + img_logo(64) + "<h2>Login</h2>"
             + "<div class='sub'>Acesse o painel Keiper Consultoria</div>"
             + form_login() + "</div>")
    return pagina("Login", corpo)


def form_login():
    return ("<form method='post' action='/login' style='margin-top:12px'>"
            "<input name='username' required placeholder='Usuário' class='input'>"
            "<input type='password' name='senha' required placeholder='Senha' class='input'>"
            "<button class='btn-salvar' style='width:100%'>Entrar</button></form>"
            "<p style='margin-top:12px'><a class='link' href='/recuperar'>Esqueci minha senha</a></p>")


@app.route("/recuperar", methods=["GET", "POST"])
def recuperar():
    if request.method == "POST":
        ident = (request.form.get("ident") or "").strip()
        if not auth_attempt_allowed(banco, executar, "reset", ident):
            return "Muitas solicitações. Aguarde 15 minutos.", 429, {"Retry-After": "900"}
        if not smtp_configurado():
            corpo = ("<div class='card' style='max-width:460px;margin:60px auto;text-align:center'>"
                     "<h2>Recuperar senha</h2>"
                     "<div class='aviso'>A recuperação por e-mail <b>não está configurada</b> "
                     "no servidor. Peça ao administrador para redefinir sua senha "
                     "na tela de Usuários.</div>"
                     "<p><a class='link' href='/login'>← Voltar ao login</a></p></div>")
            return pagina("Recuperar senha", corpo)
        with banco() as conn:
            u = consultar(conn, "SELECT * FROM usuarios WHERE email=%s OR username=%s",
                          (ident, ident))
        u = u[0] if u else None
        if u and u["email"]:
            token = criar_token_reset(u["id"])
            link = "https://app.keiperconsultoria.com.br/redefinir/" + token
            texto = ("Olá " + (u["nome"] or u["username"]) + ",\n\n"
                     "Recebemos um pedido de redefinição de senha do painel "
                     "Keiper Consultoria.\n"
                     "Use o link abaixo para criar uma nova senha "
                     "(válido por 30 minutos):\n\n"
                     + link + "\n\nSe não foi você quem pediu, ignore este e-mail.")
            enviar_email(u["email"], "Redefinição de senha — Keiper Consultoria", texto)
        corpo = ("<div class='card' style='max-width:460px;margin:60px auto;text-align:center'>"
                 "<h2>Recuperar senha</h2>"
                 "<div class='sucesso'>Se este e-mail estiver cadastrado, você receberá "
                 "o link de redefinição em instantes. Verifique também a caixa de spam.</div>"
                 "<p><a class='link' href='/login'>← Voltar ao login</a></p></div>")
        return pagina("Recuperar senha", corpo)
    corpo = ("<div class='card' style='max-width:460px;margin:60px auto;text-align:center'>"
             "<h2>Recuperar senha</h2>"
             "<div class='sub'>Digite seu usuário ou e-mail cadastrado. "
             "Enviaremos um link para criar uma nova senha.</div>"
             "<form method='post' action='/recuperar'>"
             "<input name='ident' required placeholder='Usuário ou e-mail' class='input'>"
             "<button class='btn-salvar' style='width:100%'>Enviar link</button></form>"
             "<p style='margin-top:12px'><a class='link' href='/login'>← Voltar ao login</a></p></div>")
    return pagina("Recuperar senha", corpo)


@app.route("/redefinir/<token>", methods=["GET", "POST"])
def redefinir(token):
    th = hashlib.sha256(token.encode()).hexdigest()
    with banco() as conn:
        r = consultar(conn, "SELECT * FROM reset_tokens WHERE token_hash=%s", (th,))
    r = r[0] if r else None
    if not r or r["expira_em"] < int(time.time()):
        corpo = ("<div class='card' style='max-width:460px;margin:60px auto;text-align:center'>"
                 "<h2>Link inválido</h2>"
                 "<div class='aviso'>Este link de redefinição é inválido ou expirou. "
                 "Solicite um novo.</div>"
                 "<p><a class='link' href='/recuperar'>Solicitar novo link</a></p></div>")
        return pagina("Redefinir senha", corpo)
    erro = ""
    if request.method == "POST":
        senha = request.form.get("senha") or ""
        if not 12 <= len(senha) <= 256:
            erro = "A senha precisa ter de 12 a 256 caracteres."
        else:
            with banco() as conn:
                # Lock the account before consuming the token: simultaneous resets
                # for this user cannot both pass, including two different tokens.
                consultar(conn, "SELECT id FROM usuarios WHERE id=%s FOR UPDATE", (r["usuario_id"],))
                consumed = consultar(conn, "DELETE FROM reset_tokens WHERE token_hash=%s AND expira_em >= %s RETURNING usuario_id",
                                     (th, int(time.time())))
                if not consumed:
                    return "Link inválido ou já utilizado.", 400
                executar(conn, "UPDATE usuarios SET senha_hash=%s WHERE id=%s",
                         (generate_password_hash(senha), r["usuario_id"]))
                executar(conn, "DELETE FROM reset_tokens WHERE usuario_id=%s", (r["usuario_id"],))
            corpo = ("<div class='card' style='max-width:460px;margin:60px auto;text-align:center'>"
                     "<h2>Senha alterada com sucesso!</h2>"
                     "<div class='sucesso'>Agora você pode entrar com a nova senha.</div>"
                     "<p><a class='link' href='/login'>Ir para o login</a></p></div>")
            return pagina("Senha alterada", corpo)
    corpo = ("<div class='card' style='max-width:460px;margin:60px auto;text-align:center'>"
             "<h2>Criar nova senha</h2>"
             + ("<div class='aviso'>" + esc(erro) + "</div>" if erro else "")
             + "<form method='post' action='/redefinir/" + esc(token) + "'>"
             "<input type='password' name='senha' required "
             "placeholder='Nova senha (mín. 12 caracteres)' class='input'>"
             "<button class='btn-salvar' style='width:100%'>Salvar senha</button></form></div>")
    return pagina("Redefinir senha", corpo)


@app.route("/logout")
def logout():
    session.clear()
    return redirect("/login")


@app.route("/admin/diagnostico-visitas-copas")
@admin_required
def diagnostico_visitas_copas():
    """Manual, read-only API probe for Copas. Never expose access tokens."""
    from scripts.copas_visits_render_probe import probe
    try:
        resultado = probe()
    except Exception:
        app.logger.exception("copas_visits_probe_failed")
        return jsonify({"erro": "Falha no diagnóstico de visitas da Copas"}), 503
    return jsonify({"loja": "Copas Ind", "item_id": "MLB5325240286", **resultado})


@app.route("/admin/progresso-visitas", methods=["GET"])
@admin_required
def progresso_visitas_todas():
    """Admin-only read-only progress by account; incomplete visits are not store totals."""
    import html
    from copas_visits_store import progress
    try:
        with banco() as conn:
            cur = executar(conn, """
                SELECT c.id, c.cliente FROM conexoes c
                WHERE c.status='ativa' AND EXISTS (
                    SELECT 1 FROM anuncios a WHERE a.conexao_id=c.id AND a.status='active'
                ) ORDER BY c.id
            """)
            lojas = cur.fetchall()
        linhas = []
        for conexao_id, nome in lojas:
            dados = progress(connection_id=conexao_id)
            ativos = dados["anuncios_ativos"]
            consultados = dados["anuncios_consultados"]
            percentual = round(100 * consultados / ativos, 1) if ativos else 0
            data = dados["ultima_consulta"] or "Ainda não consultado"
            linhas.append(
                "<tr><td>" + html.escape(str(nome)) + "</td><td>" + str(consultados)
                + " / " + str(ativos) + "</td><td>" + str(percentual)
                + "%</td><td>" + str(dados["visitas_registradas"])
                + "</td><td>" + html.escape(str(data)) + "</td></tr>"
            )
        corpo = (
            "<div class='card'><h2>Coleta de visitas por loja</h2>"
            "<p>Os números de visitas são parciais até que todos os anúncios tenham sido consultados. "
            "A coleta automática avança em lotes de até 10 anúncios.</p>"
            "<div style='overflow-x:auto'><table><thead><tr><th>Loja</th><th>Consultados</th>"
            "<th>Cobertura</th><th>Visitas (7 dias, parcial)</th><th>Última consulta (UTC)</th>"
            "</tr></thead><tbody>" + "".join(linhas) + "</tbody></table></div></div>"
        )
        return pagina("Progresso de visitas", corpo)
    except Exception:
        app.logger.exception("visits_all_accounts_progress_failed")
        return jsonify({"erro": "Falha ao consultar progresso das visitas"}), 503


@app.route("/admin/progresso-visitas-copas", methods=["GET"])
@admin_required
def progresso_visitas_copas():
    """Read-only administrative coverage report; never starts a collection."""
    from copas_visits_store import progress
    try:
        return jsonify(progress())
    except Exception:
        app.logger.exception("copas_visits_progress_failed")
        return jsonify({"erro": "Falha ao consultar cobertura de visitas"}), 503


@app.route("/admin/coleta-visitas-copas", methods=["GET", "POST"])
@admin_required
def coleta_visitas_copas():
    """Admin-confirmed bounded pilot; never run collection on GET."""
    import hmac
    if request.method == "GET":
        nonce = secrets.token_urlsafe(24)
        session["copas_visits_nonce"] = nonce
        corpo = ("<div class='card'><h2>Coleta piloto de visitas — Copas Ind</h2>"
                 "<p>Consulta um lote de até 10 anúncios ativos, apenas os últimos sete dias completos. "
                 "Não altera vendas nem outras lojas.</p>"
                 "<form method='post' action='/admin/coleta-visitas-copas'>"
                 "<input type='hidden' name='nonce' value='" + esc(nonce) + "'>"
                 "<label for='offset'>Início do lote (0, 10, 20...)</label>"
                 "<input id='offset' name='offset' type='number' min='0' step='10' value='0' required>"
                 "<button class='btn-acoes' type='submit'>Executar coleta piloto</button>"
                 "</form></div>")
        return pagina("Coleta piloto de visitas", corpo)
    nonce = session.pop("copas_visits_nonce", "")
    submitted = request.form.get("nonce", "")
    if not nonce or not hmac.compare_digest(nonce, submitted):
        return jsonify({"erro": "Confirmação inválida ou expirada"}), 403
    from copas_visits_store import collect
    try:
        raw_offset = request.form.get("offset", "0")
        if not raw_offset.isascii() or not raw_offset.isdecimal():
            return jsonify({"erro": "Lote inválido"}), 400
        offset = int(raw_offset)
        if offset > 10000 or offset % 10:
            return jsonify({"erro": "Lote inválido"}), 400
        result = collect(limit=10, offset=offset)
    except Exception:
        app.logger.exception("copas_visits_collection_failed")
        return jsonify({"erro": "Falha na coleta piloto"}), 503
    return jsonify(result)


@app.route("/")
def home():
    if "usuario_id" not in session:
        return redirect("/login")
    return redirect("/painel")

@app.route("/conectar")
def conectar():
    if not CLIENT_ID or not REDIRECT_URI:
        return pagina("Erro",
                      "<div class='card'><h2>Configuração faltando</h2>"
                      "<div class='aviso'>As variáveis ML_CLIENT_ID e "
                      "ML_REDIRECT_URI precisam estar no Render "
                      "(Environment).</div></div>")
    state = secrets.token_urlsafe(32)
    session["ml_oauth"] = {"state": state, "created_at": int(time.time()),
                           "usuario_id": session.get("usuario_id")}
    url = AUTH_URL + "?" + urlencode({
        "response_type": "code", "client_id": CLIENT_ID,
        "redirect_uri": REDIRECT_URI, "state": state})
    return redirect(url)


@app.route("/callback")
@app.route("/oauth")
def callback_ml():
    pending = session.get("ml_oauth") or {}
    state = request.args.get("state") or ""
    age = int(time.time()) - pending.get("created_at", 0)
    if (not state or not secrets.compare_digest(state.encode(), pending.get("state", "").encode())
            or not 0 <= age <= 900
            or pending.get("usuario_id") != session.get("usuario_id")):
        return pagina("Autorização", "<div class='card'><h2>Autorização inválida ou expirada</h2>"
                      "<p>Inicie novamente a conexão neste navegador.</p>"
                      "<p><a class='link' href='/conectar'>Conectar loja</a></p></div>"), 400
    session.pop("ml_oauth", None)
    code = request.args.get("code")
    if not code:
        return pagina("Autorização",
                      "<div class='card'><h2>Autorização não concluída</h2>"
                      "<div class='sub'>O Mercado Livre não devolveu o código. "
                      "Tente conectar de novo.</div></div>")
    try:
        resp = requests.post(TOKEN_URL, data={
            "grant_type": "authorization_code",
            "client_id": CLIENT_ID,
            "client_secret": CLIENT_SECRET,
            "code": code,
            "redirect_uri": REDIRECT_URI,
        }, timeout=25)
        dados = resp.json()
        if resp.status_code != 200 or not all(dados.get(k) for k in
                ("access_token", "refresh_token", "expires_in")):
            raise ValueError("oauth_token_failed")
        r = requests.get(API + "/users/me",
                         headers={"Authorization": "Bearer "
                                  + dados["access_token"]}, timeout=25)
        eu = r.json() if r.status_code == 200 else {}
        if not eu.get("id") or str(eu["id"]) != str(dados.get("user_id")):
            raise ValueError("oauth_identity_failed")
    except (requests.RequestException, ValueError, TypeError):
        return pagina("Erro",
                      "<div class='card'><h2>Falha ao autorizar</h2>"
                      "<div class='aviso'>Não foi possível concluir a conexão com o Mercado Livre. "
                      "Tente novamente.</div><p><a class='link' href='/conectar'>"
                      "Conectar loja</a></p></div>"), 502
    with banco() as conn:
        existe = consultar(conn, "SELECT id FROM conexoes WHERE user_id=%s",
                           (eu.get("id"),))
        if existe:
            executar(conn, "UPDATE conexoes SET access_token=%s, "
                           "refresh_token=%s, expires_at=%s, status='ativa' "
                           "WHERE id=%s",
                     (seal_token(dados["access_token"]), seal_token(dados["refresh_token"]),
                      int(time.time()) + dados["expires_in"],
                      existe[0]["id"]))
            cid = existe[0]["id"]
        else:
            cur = executar(conn,
                           """INSERT INTO conexoes
                              (cliente, user_id, access_token, refresh_token,
                               expires_at, status)
                              VALUES (%s,%s,%s,%s,%s,'ativa') RETURNING id""",
                           (eu.get("nickname") or "Nova loja", eu.get("id"),
                            seal_token(dados["access_token"]), seal_token(dados["refresh_token"]),
                            int(time.time()) + dados["expires_in"]))
            cid = cur.fetchone()[0]
    if session.get("usuario_id") and pode_ver_loja(cid):
        return redirect("/painel/" + str(cid))
    return pagina("Loja conectada",
                  "<div class='card'><h2>Loja conectada com sucesso</h2>"
                  "<p>A autorização do Mercado Livre foi concluída. "
                  "A Keiper Consultoria poderá finalizar a configuração do seu acesso.</p>"
                  "<p>Se você já recebeu seu usuário e senha, entre no painel.</p>"
                  "<p><a class='link' href='/login'>Ir para o login</a></p></div>")
@app.route("/status")
@login_required
def status():
    with banco() as conn:
        if session.get("tipo") == "admin":
            linhas = consultar(conn, "SELECT cliente, user_id, status, expires_at FROM conexoes")
        else:
            linhas = consultar(conn, "SELECT cliente, user_id, status, expires_at "
                                     "FROM conexoes WHERE id=%s",
                               (session.get("conexao_id"),))
    linhas_html = ""
    for l in linhas:
        quando = "-"
        if l["expires_at"]:
            quando = hora_br(l["expires_at"])
        tag = "tag " + (l["status"] if l["status"] in ("ativa", "pausado", "expirada") else "normal")
        linhas_html += ("<tr><td>" + esc(str(l["cliente"])) + "</td><td>"
                        + str(l["user_id"]) + "</td><td><span class='" + tag + "'>"
                        + esc(l["status"]) + "</span></td><td>" + quando + "</td></tr>")
    corpo = ("<div class='card'><h2>Conexões</h2>"
             "<div class='sub'>Status de cada loja conectada</div>"
             "<table><thead><tr><th>Cliente</th><th>Loja</th><th>Status</th>"
             "<th>Expira em</th></tr></thead>"
             "<tbody>" + linhas_html + "</tbody></table></div>")
    return pagina("Status", corpo)


@app.route("/renomear/<int:cid>", methods=["POST"])
@login_required
def renomear(cid):
    if not pode_ver_loja(cid):
        return jsonify({"ok": False}), 403
    novo = (request.form.get("cliente") or "").strip()
    if not novo:
        return jsonify({"ok": False}), 400
    with banco() as conn:
        executar(conn, "UPDATE conexoes SET cliente=%s WHERE id=%s", (novo, cid))
    from overview_cache import invalidate_overview_cache
    invalidate_overview_cache()
    return jsonify({"ok": True})


@app.route("/salvar_custo/<int:cid>/<item_id>", methods=["POST"])
@login_required
def salvar_custo(cid, item_id):
    if not pode_ver_loja(cid):
        return jsonify({"ok": False}), 403
    try:
        custo = float((request.form.get("custo") or "0").replace(",", "."))
    except Exception:
        return jsonify({"ok": False}), 400
    if custo < 0:
        return jsonify({"ok": False}), 400
    with banco() as conn:
        executar(conn, "UPDATE anuncios SET custo=%s WHERE conexao_id=%s AND item_id=%s",
                 (custo, cid, item_id))
    return jsonify({"ok": True})


@app.route("/salvar_roas/<int:cid>/<campanha_id>", methods=["POST"])
@login_required
def salvar_roas(cid, campanha_id):
    if not pode_ver_loja(cid):
        return jsonify({"ok": False}), 403
    try:
        roas = float((request.form.get("roas") or "0").replace(",", "."))
    except Exception:
        return jsonify({"ok": False}), 400
    if roas < 0:
        return jsonify({"ok": False}), 400
    from ads_change_history import number, record_changes
    if number(roas) is None:
        return jsonify({"ok": False}), 400
    with banco() as conn:
        anterior = consultar(conn, """SELECT roas_objetivo FROM ads_campanhas
            WHERE conexao_id=%s AND campanha_id=%s FOR UPDATE""", (cid,campanha_id))
        if not anterior:
            return jsonify({"ok": False, "error":"Campanha não encontrada"}), 404
        record_changes(conn, executar, cid, campanha_id, anterior[0],
                       {"roas_objetivo":roas}, int(time.time()), "painel")
        executar(conn, "UPDATE ads_campanhas SET roas_objetivo=%s "
                       "WHERE conexao_id=%s AND campanha_id=%s",
                 (roas, cid, campanha_id))
    return jsonify({"ok": True})


def financeiro_status_usuario(uid,cid):
    """Always validate progress server-side, never rely on a client cookie."""
    from client_onboarding import status
    with banco() as conn:
        row=consultar(conn,"SELECT * FROM cliente_onboarding_financeiro WHERE usuario_id=%s AND conexao_id=%s",(uid,cid))
        if not row:
            return status(None,0,0)
        count=consultar(conn,"""SELECT COUNT(*) FILTER(WHERE LOWER(status)='active') AS total,
             COUNT(*) FILTER(WHERE LOWER(status)='active' AND custo IS NOT NULL AND custo>=0) AS com_custo
             FROM anuncios WHERE conexao_id=%s""",(cid,))
    metrics=count[0] if count else {}
    return status(row[0] if row else None,int(metrics.get('total') or 0),int(metrics.get('com_custo') or 0))


@app.before_request
def exigir_onboarding_financeiro_cliente():
    if session.get('tipo')!='cliente' or not session.get('usuario_id'):
        return None
    cid=session.get('conexao_id')
    if not isinstance(cid,int):
        return None
    if session.get('onboarding_financeiro_concluido')==cid:
        return None
    # The login records whether the client was enrolled. Legacy sessions
    # keep their previous access; entering the shared portal always rechecks.
    if not session.get('onboarding_requerido') and request.endpoint!='prioridades_loja':
        return None
    if request.endpoint in {'login','logout','logo_arquivo','static',
                            'primeiro_acesso','primeiro_acesso_imposto',
                            'modelo_custos','upload_custos'}:
        return None
    progress=financeiro_status_usuario(session['usuario_id'],cid)
    if not progress['required']:
        session['onboarding_requerido']=False
        return None
    session['onboarding_requerido']=True
    if progress['done']:
        session['onboarding_financeiro_concluido']=cid
        return None
    return redirect('/primeiro-acesso')


@app.route('/primeiro-acesso')
@login_required
def primeiro_acesso():
    if session.get('tipo')!='cliente' or not isinstance(session.get('conexao_id'),int):
        return redirect('/painel')
    cid=session['conexao_id']
    progress=financeiro_status_usuario(session['usuario_id'],cid)
    if not progress['required']:
        return redirect('/painel/'+str(cid))
    if progress['done']:
        session['onboarding_financeiro_concluido']=cid
        return redirect('/painel/'+str(cid))
    with banco() as conn:
        rows=consultar(conn,'SELECT cliente,aliquota_imposto FROM conexoes WHERE id=%s',(cid,))
    if not rows:
        return pagina('Configuração inicial','<div class="card">Loja não encontrada.</div>'),404
    from client_onboarding import onboarding_page
    message={'arquivo':'O arquivo precisa ser um Excel .xlsx válido.',
             'sem_itens':'A planilha não atualizou nenhum SKU desta loja. Confira o modelo.',
             'taxa':'Informe uma alíquota válida entre 0% e 100%.',
             'tamanho':'O arquivo é muito grande. Limite: 4 MB.',
             'custos':'A planilha contém custos inválidos ou negativos.'}.get(request.args.get('erro'),'')
    return pagina('Configuração inicial',onboarding_page(cid,rows[0]['cliente'],progress,
        rows[0]['aliquota_imposto'] if rows[0]['aliquota_imposto'] is not None else '',message))


@app.route('/primeiro-acesso/imposto',methods=['POST'])
@login_required
def primeiro_acesso_imposto():
    if session.get('tipo')!='cliente' or not isinstance(session.get('conexao_id'),int):
        return 'Acesso restrito',403
    from client_onboarding import validate_tax
    try:
        aliquota=validate_tax(request.form.get('aliquota'))
    except ValueError:
        return redirect('/primeiro-acesso?erro=taxa')
    cid=session['conexao_id']
    with banco() as conn:
        executar(conn,'UPDATE conexoes SET aliquota_imposto=%s WHERE id=%s',(aliquota,cid))
        executar(conn,"""INSERT INTO cliente_onboarding_financeiro
            (usuario_id,conexao_id,imposto_confirmado_em) VALUES (%s,%s,%s)
            ON CONFLICT (usuario_id) DO UPDATE SET
            conexao_id=EXCLUDED.conexao_id,
            imposto_confirmado_em=EXCLUDED.imposto_confirmado_em""",
            (session['usuario_id'],cid,int(time.time())))
    return redirect('/primeiro-acesso')


@app.route("/salvar_config/<int:cid>", methods=["POST"])
@login_required
def salvar_config(cid):
    if not pode_ver_loja(cid):
        return pagina("Acesso negado", "<div class='card'><h2>Acesso restrito</h2></div>")
    try:
        aliquota = float((request.form.get("aliquota") or "0").replace(",", "."))
    except Exception:
        aliquota = 0
    try:
        comissao = float((request.form.get("comissao") or "0").replace(",", "."))
    except Exception:
        comissao = 12
    with banco() as conn:
        executar(conn, "UPDATE conexoes SET aliquota_imposto=%s, comissao_pct=%s WHERE id=%s",
                 (aliquota, comissao, cid))
    return redirect("/custos/" + str(cid))


@app.route("/upload_custos/<int:cid>", methods=["POST"])
@login_required
def upload_custos(cid):
    if not pode_ver_loja(cid):
        return pagina("Acesso negado", "<div class='card'><h2>Acesso restrito</h2></div>"),403
    onboarding=(session.get('tipo')=='cliente' and session.get('onboarding_financeiro_concluido')!=cid
                and financeiro_status_usuario(session['usuario_id'],cid)['required'])
    base='/primeiro-acesso' if onboarding else '/custos/'+str(cid)
    if openpyxl is None:
        return redirect(base+'?erro=biblioteca')
    arquivo=request.files.get('arquivo')
    if not arquivo or not arquivo.filename:
        return redirect(base+'?erro=sem_arquivo')
    if not arquivo.filename.lower().endswith('.xlsx'):
        return redirect(base+'?erro=arquivo_invalido')
    from client_onboarding import parse_cost
    try:
        arquivo.stream.seek(0,2)
        size=arquivo.stream.tell()
        arquivo.stream.seek(0)
        if size>4*1024*1024:
            return redirect(base+'?erro=tamanho')
        wb=openpyxl.load_workbook(arquivo,data_only=True,read_only=True)
        ws=wb.active
    except Exception:
        return redirect(base+'?erro=arquivo_invalido')
    try:
        with banco() as conn:
            anuncios=consultar(conn,"SELECT item_id,sku FROM anuncios WHERE conexao_id=%s",(cid,))
        mapa={}
        for a in anuncios:
            for key in (a.get('sku'),a['item_id']):
                if key is not None:
                    mapa.setdefault(str(key).strip().upper(),set()).add(a['item_id'])
        atualizacoes={}
        nao_achados=set()
        for index,row in enumerate(ws.iter_rows(min_row=2,max_col=2,values_only=True)):
            if index>=20000:
                return redirect(base+'?erro=tamanho')
            sku=row[0] if len(row)>0 else None
            custo=row[1] if len(row)>1 else None
            if sku is None and custo is None:
                continue
            if sku is None or custo is None:
                return redirect(base+'?erro=custos')
            value=parse_cost(custo)
            if value is None:
                return redirect(base+'?erro=custos')
            ids=mapa.get(str(sku).strip().upper())
            if ids:
                for item_id in ids:
                    atualizacoes[item_id]=value
            else:
                nao_achados.add(str(sku))
    except Exception:
        app.logger.exception('cost_import_invalid cid=%s',cid)
        return redirect(base+'?erro=arquivo_invalido')
    finally:
        try:
            wb.close()
        except Exception:
            pass
    if not atualizacoes:
        return redirect(base+'?erro=sem_itens')
    from psycopg2.extras import execute_batch
    with banco() as conn:
        with conn.cursor() as cursor:
            execute_batch(cursor,"UPDATE anuncios SET custo=%s WHERE conexao_id=%s AND item_id=%s",
                          [(value,cid,item_id) for item_id,value in atualizacoes.items()],page_size=200)
        if onboarding:
            executar(conn,"""INSERT INTO cliente_onboarding_financeiro
                (usuario_id,conexao_id,custos_importados_em) VALUES (%s,%s,%s)
                ON CONFLICT (usuario_id) DO UPDATE SET
                conexao_id=EXCLUDED.conexao_id,
                custos_importados_em=EXCLUDED.custos_importados_em""",
                (session['usuario_id'],cid,int(time.time())))
    if onboarding:
        return redirect('/primeiro-acesso')
    return redirect('/custos/'+str(cid)+'?ok='+str(len(atualizacoes))+'&faltam='+str(len(nao_achados)))


@app.route("/modelo_custos/<int:cid>")
@login_required
def modelo_custos(cid):
    if not pode_ver_loja(cid):
        return pagina("Acesso negado", "<div class='card'><h2>Acesso restrito</h2></div>")
    if openpyxl is None:
        return pagina("Erro", "<div class='card'><h2>Biblioteca de Excel não instalada</h2>"
                      "<div class='sub'>Adicione 'openpyxl' no requirements.txt "
                      "e faça o deploy de novo.</div></div>")
    with banco() as conn:
        c = consultar(conn, "SELECT * FROM conexoes WHERE id=%s", (cid,))
        c = c[0] if c else None
        if not c:
            return "Loja não encontrada", 404
        anuncios = consultar(conn, "SELECT item_id, sku, titulo FROM anuncios "
                                   "WHERE conexao_id=%s ORDER BY titulo", (cid,))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Custos"
    ws.append(["SKU", "Custo"])
    for a in anuncios:
        ws.append([a["sku"] or a["item_id"], None])
    buf = io.BytesIO()
    wb.save(buf)
    nome = "modelo_custos_" + str(c["cliente"]).replace(" ", "_") + ".xlsx"
    return Response(buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=" + nome})


@app.route("/custos/<int:cid>")
@login_required
def custos(cid):
    if not pode_ver_loja(cid):
        return pagina("Acesso negado", "<div class='card'><h2>Acesso restrito</h2></div>")
    with banco() as conn:
        c = consultar(conn, "SELECT * FROM conexoes WHERE id=%s", (cid,))
        c = c[0] if c else None
        if c:
            anuncios = consultar(conn, "SELECT * FROM anuncios WHERE conexao_id=%s "
                                       "ORDER BY titulo", (cid,))
    if not c:
        return pagina("Não encontrada", "<div class='card'><h2>Loja não encontrada</h2></div>")
    ok = request.args.get("ok")
    faltam = request.args.get("faltam")
    erro = request.args.get("erro")
    aliquota = c["aliquota_imposto"] if c["aliquota_imposto"] is not None else 0
    comissao = c["comissao_pct"] if c["comissao_pct"] is not None else 12
    ultima = ultima_atualizacao(cid)
    corpo = "<h2 style='margin-bottom:10px'>Custos</h2>" + barra_atualizar(cid, ultima, "custos")
    if erro == "sem_arquivo":
        corpo += "<div class='aviso'>Nenhum arquivo selecionado. " \
                 "Escolha a planilha e clique em Importar.</div>"
    elif erro == "arquivo_invalido":
        corpo += "<div class='aviso'>Não consegui ler o arquivo. " \
                 "Use um Excel no formato .xlsx.</div>"
    elif erro == "biblioteca":
        corpo += "<div class='aviso'>A biblioteca de Excel não está instalada. " \
                 "Adicione 'openpyxl' no requirements.txt e faça o deploy de novo.</div>"
    if ok is not None:
        msg = "<b>Importação concluída:</b> " + ok + " custo(s) atualizado(s)."
        if faltam and faltam != "0":
            msg += (" " + faltam + " SKU(s) não encontrados nos anúncios desta loja "
                    "(confira a grafia ou use o código MLB).")
        corpo += "<div class='sucesso'>" + msg + "</div>"
    com_custo = sum(1 for a in anuncios if a["custo"] is not None)
    corpo += ("<div class='grid' style='margin-bottom:18px'>"
              "<div class='stat'><div class='num'>" + str(len(anuncios))
              + "</div><div class='lab'>Anúncios na loja</div></div>"
              "<div class='stat'><div class='num'>" + str(com_custo)
              + "</div><div class='lab'>Com custo cadastrado</div></div>"
              "<div class='stat'><div class='num'>" + str(len(anuncios) - com_custo)
              + "</div><div class='lab'>Falta cadastrar custo</div></div>"
              "</div>")
    corpo += ("<div class='card'><h2>Configuração da loja</h2>"
              "<div class='sub'>Usadas no cálculo da Margem de Contribuição "
              "na tela de Vendas</div>"
              "<form method='post' action='/salvar_config/" + str(cid)
              + "' style='display:flex;flex-wrap:wrap;gap:10px;align-items:center'>"
              "<label style='font-size:13px'>Alíquota de imposto (%):</label>"
              "<input type='number' step='0.01' min='0' name='aliquota' value='"
              + str(aliquota) + "' class='filtro' style='width:100px'>"
              "<label style='font-size:13px'>Comissão ML (%):</label>"
              "<input type='number' step='0.01' min='0' name='comissao' value='"
              + str(comissao) + "' class='filtro' style='width:100px'>"
              "<button class='filtro-btn'>Salvar</button>"
              "</form></div>")
    corpo += ("<div class='card'><h2>Importar custos em massa (Excel)</h2>"
              "<div class='sub'>Planilha com a <b>1ª coluna = SKU</b> e a "
              "<b>2ª coluna = custo (R$)</b>. Pode usar o SKU cadastrado no anúncio "
              "ou o código MLB do item. A primeira linha é o cabeçalho e é ignorada.</div>"
              "<form method='post' action='/upload_custos/" + str(cid)
              + "' enctype='multipart/form-data' "
              "style='display:flex;flex-wrap:wrap;gap:10px;align-items:center'>"
              "<input type='file' name='arquivo' accept='.xlsx' required "
              "class='filtro' style='padding:6px'>"
              "<button class='filtro-btn'>Importar</button>"
              "</form>"
              "<div style='margin-top:10px'><a class='link' href='/modelo_custos/"
              + str(cid) + "'>⬇ Baixar modelo (Excel) já com os SKUs da loja "
              "— preencha a coluna Custo e suba de volta</a></div>"
              "</div>")
    if not anuncios:
        corpo += ("<div class='card' style='text-align:center;padding:40px'>"
                  "<h2>Nenhum anúncio carregado ainda</h2>"
                  "<div class='sub'>Use 'Atualizar dados' no menu lateral "
                  "para buscar os anúncios da loja</div></div>")
        return pagina_loja("Custos", cid, c["cliente"], "custos", corpo)
    linhas_html = ""
    for a in anuncios:
        if a["sku"]:
            sku_txt = esc(str(a["sku"]))
        else:
            sku_txt = "<span class='tag sem'>—</span>"
        if a["custo"] is not None:
            custo_html = ("<span id='custo-" + str(cid) + "-" + str(a["item_id"])
                          + "'>R$ %.2f</span>" % a["custo"])
        else:
            custo_html = ("<span id='custo-" + str(cid) + "-" + str(a["item_id"])
                          + "'><a class='link' href='#' onclick='editarCusto("
                          + str(cid) + ", \"" + str(a["item_id"])
                          + "\"); return false;'>+ definir</a></span>")
        linhas_html += ("<tr><td>" + sku_txt + "</td>"
                        "<td><b>" + esc(str(a["titulo"]))
                        + "</b><br><small class='muted'>" + str(a["item_id"]) + "</small></td>"
                        "<td>R$ %.2f" % (a["preco"] or 0) + "</td>"
                        "<td>" + custo_html + "</td></tr>")
    corpo += ("<div class='card'><h2>Custos por anúncio <span class='badge'>"
              + str(len(anuncios)) + "</span></h2>"
              "<div class='sub'>Clique em '+ definir' ou no valor para editar "
              "o custo de cada anúncio</div>"
              "<div style='overflow-x:auto'><table>"
              "<thead><tr><th>SKU</th><th>Anúncio</th><th>Preço</th>"
              "<th>Custo</th></tr></thead>"
              "<tbody>" + linhas_html + "</tbody></table></div>"
              "<div class='muted'>O custo é usado no cálculo da Margem de Contribuição "
              "(tela de Vendas) e na análise de promoções.</div></div>")
    return pagina_loja("Custos", cid, c["cliente"], "custos", corpo)


def calcular_vendas(cid, inicio=None):
    """Detail: latest 100. Store summary: ALL orders in its bounded date range."""
    with banco() as conn:
        c = consultar(conn, "SELECT * FROM conexoes WHERE id=%s", (cid,))
        c = c[0] if c else None
        if not c:
            return None, []
        if inicio is None:
            pedidos = consultar(conn, "SELECT * FROM pedidos WHERE conexao_id=%s "
                                "AND fechado_em IS NOT NULL "
                                "ORDER BY fechado_em DESC LIMIT 100", (cid,))
        else:
            pedidos = consultar(conn, """
                SELECT * FROM pedidos
                WHERE conexao_id=%s AND fechado_em IS NOT NULL
                  AND LEFT(fechado_em,10) >= %s
                ORDER BY fechado_em DESC
            """, (cid, inicio))
        # Only load details for the selected orders, not the store's entire history.
        ids = [p["pedido_id"] for p in pedidos]
        if ids:
            itens = consultar(conn, """SELECT pi.pedido_id, pi.quantidade, pi.preco_unitario,
                                              a.custo, a.peso
                                       FROM pedidos_itens pi
                                       LEFT JOIN anuncios a
                                         ON a.conexao_id=pi.conexao_id
                                        AND a.item_id=pi.item_id
                                       WHERE pi.conexao_id=%s
                                         AND pi.pedido_id = ANY(%s)""", (cid, ids))
            envios = consultar(conn, """SELECT pedido_id, custo_frete FROM envios
                                       WHERE conexao_id=%s AND custo_frete IS NOT NULL
                                         AND pedido_id = ANY(%s)""", (cid, ids))
        else:
            itens, envios = [], []
    aliquota_pct = c["aliquota_imposto"] if c["aliquota_imposto"] is not None else 0
    comissao_pct = c["comissao_pct"] if c["comissao_pct"] is not None else 12
    aliquota = aliquota_pct / 100.0
    comissao = comissao_pct / 100.0
    itens_por_pedido = {}
    for it in itens:
        itens_por_pedido.setdefault(it["pedido_id"], []).append(it)
    frete_por_pedido = {}
    for e in envios:
        pid = e["pedido_id"]
        if pid and pid not in frete_por_pedido:
            frete_por_pedido[pid] = e["custo_frete"] or 0
    vendas = []
    for p in pedidos:
        receita = p["total"] or 0
        itens_p = itens_por_pedido.get(p["pedido_id"], [])
        custo_prod = 0.0
        frete = 0.0
        sem_custo = False
        sem_itens = not itens_p
        frete_verificavel = bool(itens_p)
        for it in itens_p:
            qtd = it["quantidade"] or 1
            pu = it["preco_unitario"] or 0
            if it["custo"] is None:
                sem_custo = True
            else:
                custo_prod += (it["custo"] or 0) * qtd
            f = frete_esperado(pu, it["peso"])
            if f is not None:
                frete += f * qtd
            else:
                frete_verificavel = False
        frete_fonte = "estimado"
        frete_real = frete_por_pedido.get(p["pedido_id"])
        if frete_real is not None:
            frete = frete_real
            frete_fonte = "real"
            frete_verificavel = True
        imposto = receita * aliquota
        comissao_v = receita * comissao
        mc = receita - custo_prod - imposto - comissao_v - frete
        mc_pct = (mc * 100.0 / receita) if receita else 0
        vendas.append({
            "pedido_id": p["pedido_id"],
            "data": str(p["fechado_em"])[:16].replace("T", " "),
            "status": str(p["status"]),
            "receita": receita,
            "custo_prod": custo_prod,
            "imposto": imposto,
            "comissao": comissao_v,
            "frete": frete,
            "frete_fonte": frete_fonte,
            "frete_verificavel": frete_verificavel,
            "mc": mc,
            "mc_pct": mc_pct,
            "sem_itens": sem_itens,
            "sem_custo": sem_custo,
        })
    return c, vendas


@app.route("/vendas/<int:cid>")
@login_required
def vendas(cid):
    if not pode_ver_loja(cid):
        return pagina("Acesso negado", "<div class='card'><h2>Acesso restrito</h2></div>")
    c, vendas = calcular_vendas(cid)
    if not c:
        return pagina("Não encontrada", "<div class='card'><h2>Loja não encontrada</h2></div>")
    aliquota_pct = c["aliquota_imposto"] if c["aliquota_imposto"] is not None else 0
    comissao_pct = c["comissao_pct"] if c["comissao_pct"] is not None else 12
    ultima = ultima_atualizacao(cid)
    corpo = ("<div class='barra-topo'><h2>Vendas e Margem de Contribuição</h2>"
             "<div style='display:flex;gap:10px;flex-wrap:wrap;align-items:center'>"
             "<a class='btn-acoes cinza' href='/exportar_vendas/"
             + str(cid) + "'>⬇ Exportar CSV</a>"
             "</div></div>"
             + barra_atualizar(cid, ultima, "vendas"))
    corpo += render_sales(cid, request.args, banco, consultar)
    if not vendas:
        corpo += ("<div class='card' style='text-align:center;padding:40px'>"
                  "<h2>Nenhuma venda carregada ainda</h2>"
                  "<div class='sub'>Use 'Atualizar agora' acima para buscar os pedidos "
                  "da loja. As vendas analisadas são os pedidos fechados.</div></div>")
        return pagina_loja("Vendas", cid, c["cliente"], "vendas", corpo)
    analisadas = sum(1 for v in vendas if not v["sem_itens"] and not v["sem_custo"])
    pendentes = len(vendas) - analisadas
    tot_receita = sum(v["receita"] for v in vendas
                      if not v["sem_itens"] and not v["sem_custo"])
    tot_mc = sum(v["mc"] for v in vendas
                 if not v["sem_itens"] and not v["sem_custo"])
    mc_medio = (tot_mc * 100.0 / tot_receita) if tot_receita else 0
    com_frete_real = sum(1 for v in vendas if v["frete_fonte"] == "real")
    corpo += ("<div class='grid' style='margin-bottom:18px'>"
              "<div class='stat'><div class='num'>R$ %.2f" % tot_receita
              + "</div><div class='lab'>Receita das vendas analisadas</div></div>"
              "<div class='stat'><div class='num'>R$ %.2f" % tot_mc
              + "</div><div class='lab'>Margem de Contribuição total</div></div>"
              "<div class='stat'><div class='num'>" + ("%.1f%%" % mc_medio)
              + "</div><div class='lab'>MG média (MC ÷ Receita)</div></div>"
              "<div class='stat'><div class='num'>" + str(pendentes)
              + "</div><div class='lab'>Vendas sem custo cadastrado</div></div>"
              "</div>")
    linhas_html = ""
    for v in vendas:
        if v["sem_itens"]:
            mc_html = "<span class='tag sem'>itens não registrados</span>"
        elif v["sem_custo"]:
            mc_html = "<a class='link' href='/custos/" + str(cid) + "'>cadastre o custo</a>"
        elif v["mc_pct"] <= 0:
            mc_html = "<span class='alerta urgente'>" + ("%.1f%% (R$ %.2f)"
                      % (v["mc_pct"], v["mc"])) + "</span>"
        elif v["mc_pct"] < 20:
            mc_html = "<span class='alerta atencao'>" + ("%.1f%% (R$ %.2f)"
                      % (v["mc_pct"], v["mc"])) + "</span>"
        else:
            mc_html = "<span class='alerta ok'>" + ("%.1f%% (R$ %.2f)"
                      % (v["mc_pct"], v["mc"])) + "</span>"
        fonte = "<small class='muted'>real</small>" if v["frete_fonte"] == "real" \
            else "<small class='muted'>est.</small>"
        if v["frete"]:
            frete_html = ("R$ %.2f " % v["frete"]) + fonte
        else:
            frete_html = "<span class='tag sem'>—</span>"
        if v["sem_itens"]:
            custo_html = "<span class='tag sem'>—</span>"
        else:
            custo_html = "R$ %.2f" % v["custo_prod"]
        linhas_html += ("<tr>"
                        "<td><b>" + str(v["pedido_id"]) + "</b><br>"
                        "<small class='muted'>" + esc(v["data"]) + "</small></td>"
                        "<td>" + esc(v["status"]) + "</td>"
                        "<td>R$ %.2f" % v["receita"] + "</td>"
                        "<td>" + custo_html + "</td>"
                        "<td>R$ %.2f" % v["imposto"] + "</td>"
                        "<td>R$ %.2f" % v["comissao"] + "</td>"
                        "<td>" + frete_html + "</td>"
                        "<td>" + mc_html + "</td>"
                        "</tr>")
    corpo += ("<div class='card'><h2>Vendas <span class='badge'>"
              + str(len(vendas)) + "</span></h2>"
              "<div class='sub'>MG = Receita − Custo dos produtos − Imposto ("
              + str(aliquota_pct) + "%) − Comissão ML (" + str(comissao_pct)
              + "%) − Frete. Entre parênteses, o lucro em R$ de cada venda.</div>"
              "<div style='overflow-x:auto'><table>"
              "<thead><tr><th>Pedido</th><th>Status</th><th>Receita</th>"
              "<th>Custo produtos</th><th>Imposto</th><th>Comissão</th>"
              "<th>Frete</th><th>MG (Lucro R$)</th></tr></thead>"
              "<tbody>" + linhas_html + "</tbody></table></div>"
              "<div class='muted'>Frete: usa o <b>valor real</b> do envio quando o "
              "Mercado Livre informa (" + str(com_frete_real)
              + " vendas com frete real); senão, estimativa pela tabela oficial do ML "
              "(marcado como 'est.'). Imposto e comissão são configurados na "
              "<a class='link' href='/custos/" + str(cid) + "'>tela de Custos</a>.</div></div>")
    return pagina_loja("Vendas", cid, c["cliente"], "vendas", corpo)


@app.route("/exportar_vendas/<int:cid>")
@login_required
def exportar_vendas(cid):
    if not pode_ver_loja(cid):
        return "Acesso negado", 403
    c, vendas = calcular_vendas(cid)
    if not c:
        return "Loja não encontrada", 404
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Pedido", "Data", "Status", "Receita", "Custo produtos",
                     "Imposto", "Comissao", "Frete", "Fonte frete", "MC %", "Lucro R$"])
    for v in vendas:
        writer.writerow([v["pedido_id"], v["data"], v["status"],
                         ("%.2f" % v["receita"]).replace(".", ","),
                         ("%.2f" % v["custo_prod"]).replace(".", ","),
                         ("%.2f" % v["imposto"]).replace(".", ","),
                         ("%.2f" % v["comissao"]).replace(".", ","),
                         ("%.2f" % v["frete"]).replace(".", ","),
                         v["frete_fonte"],
                         ("%.1f" % v["mc_pct"]).replace(".", ","),
                         ("%.2f" % v["mc"]).replace(".", ",")])
    nome_arq = "vendas_" + str(c["cliente"]).replace(" ", "_") + ".csv"
    return Response(buf.getvalue(), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": "attachment; filename=" + nome_arq})
def executar_atualizacao_bg(cid):
    try:
        executar_atualizacao(cid)
    except Exception as e:
        try:
            registrar_erro(cid, "atualizacao", "Falha na atualização: " + type(e).__name__)
        except Exception:
            pass
    finally:
        with LOCK_ATUALIZACAO:
            ATUALIZANDO.discard(cid)    
@app.route("/diagnostico")
@login_required
def diagnostico():
    with banco() as conn:
        if session.get("tipo") == "admin":
            lojas = consultar(conn, "SELECT id, cliente, status "
                                    "FROM conexoes ORDER BY cliente")
            erros = consultar(conn, "SELECT * FROM erros "
                                    "ORDER BY quando DESC")
        else:
            lojas = consultar(conn, "SELECT id, cliente, status "
                                    "FROM conexoes WHERE id=%s",
                              (session.get("conexao_id"),))
            erros = consultar(conn, "SELECT * FROM erros "
                                    "WHERE conexao_id=%s ORDER BY quando DESC",
                              (session.get("conexao_id"),))
    with banco() as conn:
        if session.get("tipo") == "admin":
            execucoes = consultar(conn, """SELECT s.inicio_em,s.fim_em,s.status,s.detalhe,c.cliente FROM sincronizacao_execucoes s JOIN conexoes c ON c.id=s.conexao_id ORDER BY s.id DESC LIMIT 30""")
        else:
            execucoes = consultar(conn, """SELECT s.inicio_em,s.fim_em,s.status,s.detalhe,c.cliente FROM sincronizacao_execucoes s JOIN conexoes c ON c.id=s.conexao_id WHERE s.conexao_id=%s ORDER BY s.id DESC LIMIT 30""", (session.get("conexao_id"),))
    corpo = ("<p style='margin-bottom:10px'><a class='link' href='/painel'>"
             "← Voltar ao painel</a></p>"
             "<h2 style='margin-bottom:14px'>Diagnóstico</h2>")
    if session.get('tipo')=='admin':
        from collection_health import render as render_health,CSS as health_css
        corpo += health_css+render_health(snapshot_saude_coletas())
    if FILA_ATUALIZACAO:
        fila_txt = str(len(FILA_ATUALIZACAO)) + " na fila"
    else:
        fila_txt = "vazia"
    corpo += ("<div class='grid' style='margin-bottom:18px'>"
              "<div class='stat'><div class='num'>" + fila_txt
              + "</div><div class='lab'>Fila de atualização</div></div>"
              "</div>")
    linhas = ""
    for l in lojas:
        linhas += ("<tr><td><b>" + esc(str(l["cliente"])) + "</b></td>"
                   "<td>" + esc(str(l["status"])) + "</td><td>"
                   + hora_br(ultima_atualizacao(l["id"])) + "</td></tr>")
    corpo += ("<div class='card'><h2>Última atualização por conta</h2>"
              "<table><thead><tr><th>Conta</th><th>Status</th>"
              "<th>Última atualização</th></tr></thead><tbody>"
              + linhas + "</tbody></table></div>")
    linhas_sync = ""
    for e in execucoes:
        duracao = "—" if e["fim_em"] is None else str(max(0,(e["fim_em"]-e["inicio_em"])//60))+" min"
        linhas_sync += ("<tr><td>"+esc(str(e["cliente"]))+"</td><td>"+hora_br(e["inicio_em"])+"</td><td>"+duracao+"</td><td>"+esc(str(e["status"]))+"</td><td>"+esc(str(e["detalhe"] or "—"))+"</td></tr>")
    if not linhas_sync:
        linhas_sync = "<tr><td colspan='5'>Nenhuma execução registrada</td></tr>"
    corpo += ("<div class='card'><h2>Histórico de sincronizações</h2><div style='overflow-x:auto'><table><thead><tr><th>Loja</th><th>Início</th><th>Duração</th><th>Resultado</th><th>Detalhe</th></tr></thead><tbody>"+linhas_sync+"</tbody></table></div></div>")
    linhas_e = ""
    nomes = {loja["id"]: loja["cliente"] for loja in lojas}
    for e in erros:
        estado = "Resolvido em " + hora_br(e["resolvido_em"]) if e.get("resolvido_em") else "Pendente"
        linhas_e += ("<tr><td>" + esc(str(nomes.get(e["conexao_id"], "Loja indisponível"))) + "</td><td>"
                     + esc(estado) + "</td><td>" + esc(str(e["categoria"])) + "</td>"
                     "<td>" + esc(str(e["mensagem"])) + "</td><td>"
                     + hora_br(e["quando"]) + "</td></tr>")
    if not linhas_e:
        linhas_e = ("<tr><td colspan='5'><span class='alerta ok'>"
                    "Nenhum erro registrado</span></td></tr>")
    corpo += ("<div class='card'><h2>Erros de atualização</h2>"
              "<table><thead><tr><th>Loja</th><th>Situação</th><th>Categoria</th><th>Mensagem</th>"
              "<th>Quando</th></tr></thead><tbody>"
              + linhas_e + "</tbody></table></div>")
    return pagina("Diagnóstico", corpo)
FILA_ATUALIZACAO = []
LOCK_FILA = threading.Lock()
WORKER_VIVO = False


def iniciar_heartbeat_tarefa(tarefa_id, execucao_id=None, intervalo=300):
    """Refresh a running task lease until its worker signals completion."""
    parar = threading.Event()

    def renovar():
        while not parar.wait(intervalo):
            try:
                if not renovar_lease_tarefa_persistente(tarefa_id):
                    app.logger.error("Lease da tarefa %s nao pode ser renovado", tarefa_id)
                    break
                if execucao_id is not None and not renovar_heartbeat_execucao(execucao_id):
                    app.logger.error("Heartbeat da execucao %s nao pode ser renovado", execucao_id)
                    break
            except Exception:
                app.logger.exception("Erro ao renovar lease da tarefa %s", tarefa_id)

    thread = threading.Thread(target=renovar, daemon=True)
    thread.start()
    return parar


def processar_tarefa_persistente_uma():
    """Process one durable job; never called by the legacy scheduler by default."""
    tarefa = reservar_tarefa_persistente()
    if tarefa is None:
        return False
    cid = tarefa["conexao_id"]
    tarefa_id = tarefa["id"]
    execucao_id = None
    try:
        with store_sync_lock(banco, cid) as acquired:
            if not acquired:
                reagendar_tarefa_persistente(tarefa_id, "Loja ja em execucao")
                return False
            # Only the holder of the full-collection lock may retire abandoned
            # telemetry. No active collector can hold this same lock now.
            encerrar_execucoes_abandonadas(cid)
            execucao_id = registrar_execucao_inicio(cid)
            if execucao_id is None:
                # Another process is already synchronizing this store.
                reagendar_tarefa_persistente(tarefa_id, "Loja ja em execucao")
                return False
            parar_heartbeat = iniciar_heartbeat_tarefa(tarefa_id, execucao_id)
            try:
                sucesso = executar_atualizacao(cid)
            finally:
                parar_heartbeat.set()
            registrar_execucao_fim(execucao_id, "concluida" if sucesso else "falhou")
            if sucesso:
                finalizar_tarefa_persistente(tarefa_id, True)
            else:
                reagendar_tarefa_persistente(tarefa_id, "Atualizacao retornou False")
            return bool(sucesso)
    except Exception as exc:
        if execucao_id is not None:
            try:
                registrar_execucao_fim(execucao_id, "falhou", type(exc).__name__)
            except Exception:
                pass
        reagendar_tarefa_persistente(tarefa_id, type(exc).__name__)
        return False


def encerrar_execucoes_abandonadas(cid):
    """Caller must hold store_sync_lock for the complete claim/collect cycle."""
    with banco() as conn:
        executar(conn, """UPDATE sincronizacao_execucoes
            SET status='interrompida', fim_em=%s,
                detalhe='Retomada sob trava exclusiva; coletor anterior ausente'
            WHERE conexao_id=%s AND status='executando'""",
            (int(time.time()), cid))


def loja_habilitada_fila_persistente(cid):
    """Explicit per-store opt-in; empty env keeps every store on legacy queue."""
    raw = os.getenv("DURABLE_QUEUE_STORE_IDS", "").strip()
    if not raw:
        return False
    ids = {part.strip() for part in raw.split(",") if part.strip().isdigit()}
    return str(cid) in ids


def enfileirar_atualizacao(cid):
    global WORKER_VIVO
    if loja_habilitada_fila_persistente(cid):
        # Persist the request, but do not start a second worker in this call.
        # The durable scheduler will be enabled separately after validation.
        return criar_tarefa_persistente(cid) is not None
    with LOCK_FILA:
        if cid in FILA_ATUALIZACAO:
            return False
        with LOCK_ATUALIZACAO:
            if cid in ATUALIZANDO:
                return False
            ATUALIZANDO.add(cid)
        FILA_ATUALIZACAO.append(cid)
        if WORKER_VIVO:
            return True
        WORKER_VIVO = True
    threading.Thread(target=processar_fila, daemon=True).start()
    return True


def registrar_execucao_inicio(cid):
    with banco() as conn:
        # Claim one active execution per store across Gunicorn workers.
        # A two-hour lease allows recovery after an abrupt worker shutdown.
        agora = int(time.time())
        executar(conn, "SELECT pg_advisory_xact_lock(72591, %s)", (cid,))
        cur = executar(conn, """INSERT INTO sincronizacao_execucoes
                         (conexao_id, inicio_em, status)
                         SELECT %s, %s, 'executando'
                         WHERE NOT EXISTS (
                           SELECT 1 FROM sincronizacao_execucoes
                           WHERE conexao_id=%s AND status='executando'
                             AND COALESCE(heartbeat_em, inicio_em) >= %s
                         ) RETURNING id""",
                       (cid, agora, cid, agora - 7200))
        row = cur.fetchone()
        return row[0] if row else None


def renovar_heartbeat_execucao(execucao_id, agora=None):
    """Record liveness without modifying the original start timestamp."""
    agora = int(time.time()) if agora is None else int(agora)
    with banco() as conn:
        cur = executar(conn, """UPDATE sincronizacao_execucoes
            SET heartbeat_em=%s WHERE id=%s AND status='executando'""",
            (agora, execucao_id))
        return cur.rowcount == 1


def registrar_execucao_fim(execucao_id, status, detalhe=None):
    with banco() as conn:
        executar(conn, """UPDATE sincronizacao_execucoes
                         SET fim_em=%s, status=%s, detalhe=%s
                         WHERE id=%s""",
                 (int(time.time()), status, (detalhe or "")[:300], execucao_id))


def processar_fila():
    global WORKER_VIVO
    while True:
        with LOCK_FILA:
            if not FILA_ATUALIZACAO:
                WORKER_VIVO = False
                return
            cid = FILA_ATUALIZACAO[0]
        execucao_id = None
        try:
            with store_sync_lock(banco, cid) as acquired:
                if not acquired:
                    continue  # A live collector owns this store.
                # A free collection lock proves previous telemetry is abandoned.
                encerrar_execucoes_abandonadas(cid)
                execucao_id = registrar_execucao_inicio(cid)
                if execucao_id is None:
                    continue  # Shared execution claim is required even under the lock.
                atualizacao_ok = executar_atualizacao(cid)
                if execucao_id is not None:
                    try:
                        registrar_execucao_fim(execucao_id, "concluida" if atualizacao_ok else "falhou", None if atualizacao_ok else "Atualizacao retornou False")
                    except Exception:
                        pass
        except Exception as exc:
            if execucao_id is not None:
                try:
                    registrar_execucao_fim(execucao_id, "falhou", type(exc).__name__)
                except Exception:
                    pass
            try:
                registrar_erro(cid, "atualizacao", "Falha na atualização: " + type(exc).__name__)
            except Exception:
                pass  # Keep draining other shops even when persistence is unavailable.
        finally:
            with LOCK_ATUALIZACAO:
                ATUALIZANDO.discard(cid)
            with LOCK_FILA:
                try:
                    FILA_ATUALIZACAO.remove(cid)
                except ValueError:
                    pass    
@app.route("/auditoria-frete/<int:cid>", methods=["GET", "POST"])
@login_required
def auditoria_frete(cid):
    if not pode_ver_loja(cid):
        return redirect("/painel")
    with banco() as conn:
        lojas = consultar(conn, "SELECT cliente FROM conexoes WHERE id=%s", (cid,))
    if not lojas:
        return "Loja não encontrada", 404
    if "freight_preview_csrf" not in session:
        session["freight_preview_csrf"] = secrets.token_urlsafe(32)
    corpo = (
        "<section class='card'><h2>🚚 Auditoria de Frete</h2>"
        "<p>Envie a planilha de <b>Conciliação Bancária do Mercado Livre</b> "
        "para conferir o formato e os lançamentos de frete.</p>"
        "<div class='aviso'>A pré-visualização não grava dados. O administrador pode importar os indicadores. O relatório não informa "
        "o frete esperado para determinar cobranças indevidas.</div>"
        "<form id='frete-upload-form' method='post' enctype='multipart/form-data'>"
        "<input type='hidden' name='csrf_token' value='"
        + esc(session["freight_preview_csrf"]) + "'>"
        "<input type='file' name='arquivo' accept='.xlsx,.csv' required>"
        "<button class='btn' type='submit'>Pré-visualizar planilha</button>"
        "</form><div id='frete-progresso' class='frete-progresso' role='status' aria-live='polite' hidden>"
        "<span class='frete-spinner' aria-hidden='true'></span>"
        "<span><b>Enviando e processando o arquivo…</b><small>Isso pode levar alguns instantes. "
        "Mantenha esta página aberta até aparecer o resultado.</small></span></div>"
        "<p>Formatos XLSX e CSV · máximo 2 MB.</p></section>"
    )
    corpo += """
    <style>
    .frete-progresso{margin:18px 0 0;padding:16px;border-radius:12px;background:#eaf1ff;
        color:#173b66;display:flex;gap:14px;align-items:center}
    .frete-progresso[hidden]{display:none}
    .frete-progresso>span:last-child{min-width:0;overflow-wrap:anywhere}
    .frete-progresso small{display:block;margin-top:5px;color:#45617f}
    #frete-upload-form[aria-busy=true] button{opacity:.65;cursor:wait}
    .frete-spinner{width:24px;height:24px;border:3px solid #bcd0e6;
        border-top-color:#2563eb;border-radius:50%;animation:frete-girar .8s linear infinite;flex-shrink:0}
    @keyframes frete-girar{to{transform:rotate(360deg)}}
    .frete-metricas{display:grid;grid-template-columns:repeat(auto-fit,minmax(165px,1fr));gap:12px;margin:15px 0}
    .frete-metrica{border:1px solid #e3e9f0;border-radius:12px;padding:15px;background:#f8fafc;display:flex;flex-direction:column;gap:7px}
    .frete-metrica span{font-size:13px;color:#64748b}
    .frete-metrica strong{font-size:21px;color:#142338;overflow-wrap:anywhere}
    .frete-detalhes{margin-top:14px;border-top:1px solid #e2e8f0;padding-top:14px}
    .frete-detalhes summary{cursor:pointer;font-weight:600;color:#20528d}
    .frete-resultado{margin-top:16px}
    </style>
    <script>
    document.addEventListener('DOMContentLoaded',function(){
      var form=document.getElementById('frete-upload-form');
      var status=document.getElementById('frete-progresso');
      if(!form||!status)return;
      form.addEventListener('submit',function(event){
        if(form.getAttribute('aria-busy')==='true'){
          event.preventDefault();return;
        }
        if(!form.checkValidity())return;
        form.setAttribute('aria-busy','true');
        form.querySelectorAll('button[type=submit]').forEach(function(button){
          button.setAttribute('aria-disabled','true');
        });
        status.hidden=false;
        var file=form.querySelector('input[type=file]');
        var name=file&&file.files.length?file.files[0].name:'arquivo';
        var msg=status.querySelector('b');
        if(msg)msg.textContent='Enviando e analisando: '+name;
        status.scrollIntoView({block:'nearest',behavior:'smooth'});
      });
      window.addEventListener('pageshow',function(){
        form.removeAttribute('aria-busy');
        form.querySelectorAll('button[type=submit]').forEach(function(button){
          button.removeAttribute('aria-disabled');
        });
        status.hidden=true;
      });
    });
    </script>
    """
    if session.get('tipo') == 'admin':
        corpo = corpo.replace("</form><div id=\'frete-progresso\'", "<button class='btn' type='submit' name='acao' value='importar'>Importar indicadores desta loja</button>"
            "<fieldset><legend>Consolidar relatório completo por dia</legend>"
            "<label>Início do período <input type='date' name='periodo_inicio'></label>"
            "<label>Fim do período <input type='date' name='periodo_fim'></label>"
            "<label><input type='checkbox' name='completo' value='sim'>Confirmo que o arquivo contém todos os lançamentos do período informado, sem filtros adicionais.</label>"
            "<p>Até 31 dias por arquivo. Substitui os indicadores dos dias informados, inclusive dias sem movimentos. Não use relatórios parciais.</p>"
            "<button class='btn' type='submit' name='acao' value='consolidar'>Atualizar consolidação diária</button></fieldset></form><div id=\'frete-progresso\'")
        with banco() as conn:
            corpo += render_history(report_history(conn, consultar, cid))
            consolidado_html = render_consolidated(consolidated_history(conn, consultar, cid))
            corpo += consolidado_html
    if request.method == "POST":
        if request.form.get('acao') in ('importar', 'consolidar') and session.get('tipo') != 'admin':
            return 'Acesso restrito ao administrador', 403
        valid_csrf = secrets.compare_digest(
            session.get("freight_preview_csrf", ""),
            request.form.get("csrf_token", "")) and bool(session.get("freight_preview_csrf"))
        if not valid_csrf:
            return "Token inválido", 403
        if request.content_length is None or request.content_length > MAX_UPLOAD_BYTES + 65536:
            return "Arquivo muito grande", 413
        arquivo = request.files.get("arquivo")
        if arquivo is None:
            return "Selecione um arquivo", 400
        try:
            payload = arquivo.read(MAX_UPLOAD_BYTES + 1)
            bundle = None
            if (arquivo.filename or '').lower().endswith('.xlsx'):
                try:
                    bundle = inspect_sales_reconciliation(arquivo.filename, payload, bundle=True)
                except PreviewError:
                    pass
            abas = bundle[1] if bundle is not None else preview_reconciliation_upload(arquivo.filename, payload)
        except PreviewError as exc:
            return esc(str(exc)), 400
        detalhes_planilha = ("<details class='card frete-detalhes'><summary>Ver pré-visualização e campos identificados</summary>"
            + render_reconciliation_preview(abas)
            + render_reconciliation_column_hints(abas) + "</details>")
        try:
            if bundle is None:
                raise PreviewError('Formato não reconhecido para conferência e comparação por vendas.')
            resumo, _, linhas_comparacao, colunas_comparacao = bundle
            corpo += render_sales_reconciliation(resumo)
            resumo['sales_comparison'] = audit_against_sales(
                cid, arquivo.filename, payload, resumo, banco, consultar, calcular_vendas,
                parsed_rows=linhas_comparacao, parsed_columns=colunas_comparacao)
            corpo += render_sales_comparison(resumo['sales_comparison'])
            corpo += detalhes_planilha
            if request.form.get('acao') in ('importar', 'consolidar'):
                try:
                    with banco() as conn:
                        if request.form.get('acao') == 'consolidar':
                            novo = consolidate_report(conn, consultar, cid, payload, resumo,
                                request.form.get('periodo_inicio'), request.form.get('periodo_fim'),
                                request.form.get('completo') == 'sim')
                            corpo = corpo.replace(consolidado_html,
                                render_consolidated(consolidated_history(conn, consultar, cid)))
                        else:
                            novo = save_report(conn, consultar, cid, payload, resumo)
                    corpo = corpo.replace('Nenhum dado foi salvo.', 'Os indicadores agregados deste relatório estão registrados nesta loja.')
                    corpo += "<div class='aviso'>" + ('Indicadores importados.' if novo else 'Este arquivo já foi importado nesta loja.') + '</div>'
                except ValueError as exc:
                    return esc(str(exc)), 400
        except PreviewError as exc:
            corpo += "<div class='aviso'>" + esc(str(exc)) + "</div>"
            corpo += detalhes_planilha
    response = Response(pagina_loja("Auditoria de Frete", cid, lojas[0]["cliente"],
                                   "auditoria_frete", corpo), mimetype="text/html")
    response.headers["Cache-Control"] = "no-store"
    return response


@app.route("/painel")
@login_required
def painel():
    # The landing page is a client directory, not an expensive sales report.
    with banco() as conn:
        if session.get("tipo") == "admin":
            lojas = consultar(conn, "SELECT id, cliente, status FROM conexoes ORDER BY cliente, id")
        else:
            lojas = consultar(conn, "SELECT id, cliente, status FROM conexoes WHERE id=%s ORDER BY cliente, id",
                              (session.get("conexao_id"),))
    lojas = [loja for loja in lojas if pode_ver_loja(loja["id"])]
    from client_home import render_client_home
    corpo = render_client_home(lojas, admin=session.get("tipo") == "admin")
    if session.get("tipo") == "admin":
        # Async collection-health snapshot keeps the client directory fast.
        corpo += """<div id="kh-collection-health" aria-live="polite"></div>
        <script>(function(){
          var panel=document.getElementById('kh-collection-health');
          fetch('/painel/saude',{credentials:'same-origin'})
            .then(function(r){if(!r.ok)throw Error();return r.text();})
            .then(function(html){if(panel)panel.innerHTML=html;})
            .catch(function(){if(panel)panel.innerHTML='<p class="muted">Saúde das coletas indisponível. <a href="/diagnostico">Abrir diagnóstico</a></p>';});
        })();</script>"""
        # CRM stays beneath the client cards and health summary.
        corpo += HOME_WIDGET
    return pagina("Suas lojas", corpo)



@app.route('/painel/alertas/<int:cid>')
@login_required
def painel_alertas(cid):
    if not pode_ver_loja(cid):
        return jsonify(error='Acesso restrito'), 403
    from home_alerts import load, LABELS, label
    try:
        groups, known = load(cid, banco, consultar)
    except Exception:
        app.logger.exception('Falha ao consultar alertas da loja %s', cid)
        return jsonify(error='Alertas indisponíveis'), 503
    response = jsonify(known=known, url='/painel/alertas/'+str(cid)+'/detalhes',
        alerts=[dict(code=code, severity=LABELS[code][0], label=label(code,len(groups[code])))
                for code in LABELS if groups.get(code)])
    response.headers['Cache-Control'] = 'no-store'
    return response


@app.route('/painel/alertas/<int:cid>/detalhes')
@login_required
def painel_alertas_detalhes(cid):
    if not pode_ver_loja(cid):
        return pagina('Acesso negado', '<p>Acesso restrito</p>'), 403
    from home_alerts import load, render
    try:
        groups, known = load(cid, banco, consultar)
    except Exception:
        app.logger.exception('Falha ao consultar detalhes de alertas da loja %s', cid)
        return pagina('Alertas indisponíveis', '<p>Não foi possível consultar os alertas. Tente novamente.</p>'), 503
    response = Response(pagina('Alertas da loja', render(groups, cid, known)), mimetype='text/html')
    response.headers['Cache-Control'] = 'no-store'
    return response


def snapshot_saude_coletas():
    """A few bounded reads; no order load or marketplace API calls."""
    from collection_health import assess
    from sales_metrics import reporting_today
    hoje=reporting_today()
    semana=(hoje-datetime.timedelta(days=7)).isoformat()
    ontem=(hoje-datetime.timedelta(days=1)).isoformat()
    with banco() as conn:
        shops=consultar(conn,"SELECT id,cliente,status FROM conexoes ORDER BY cliente")
        catalog=consultar(conn,"SELECT conexao_id,MAX(atualizado_em) AS latest FROM anuncios GROUP BY conexao_id")
        ads=consultar(conn,"""SELECT conexao_id,MIN(data) AS first_day,MAX(data) AS last_day,
            COUNT(*) FILTER(WHERE data >= %s AND data <= %s) AS days_7
            FROM ads_dia GROUP BY conexao_id""",(semana,ontem))
        sales=consultar(conn,"SELECT conexao_id,MAX(dia) AS latest_day FROM vendas_cobertura WHERE dia <= %s GROUP BY conexao_id",(ontem,))
        runs=consultar(conn,"""SELECT DISTINCT ON(conexao_id) conexao_id,status,inicio_em,heartbeat_em
            FROM sincronizacao_execucoes ORDER BY conexao_id,id DESC""")
        errors=consultar(conn,"SELECT conexao_id,categoria FROM erros WHERE resolvido_em IS NULL")
    return assess(shops,catalog,ads,sales,runs,errors,today=hoje)


@app.route("/painel/saude")
@login_required
@admin_required
def painel_saude():
    from collection_health import render,CSS
    response=Response(CSS+render(snapshot_saude_coletas(),compact=True),mimetype="text/html")
    response.headers['Cache-Control']='private, max-age=180'
    return response



@app.route("/prioridades")
@login_required
def central_prioridades():
    if session.get('tipo') != 'admin':
        cid=session.get('conexao_id')
        if not isinstance(cid,int):
            return pagina('Acesso negado', '<div class="card">Conta sem loja vinculada.</div>'),403
        return redirect('/prioridades/'+str(cid))
    from client_priorities import admin_index
    with banco() as conn:
        shops=consultar(conn,"SELECT id,cliente,status FROM conexoes ORDER BY cliente")
        enabled=consultar(conn,"SELECT conexao_id FROM cliente_prioridades_acessos WHERE habilitado=TRUE")
    csrf=session.get('prioridades_csrf')
    if not csrf:
        csrf=secrets.token_urlsafe(32)
        session['prioridades_csrf']=csrf
    body=admin_index(shops,{int(r['conexao_id']) for r in enabled},snapshot_saude_coletas(),csrf)
    response=Response(pagina('Central de Prioridades',body),mimetype='text/html')
    response.headers['Cache-Control']='no-store'
    return response


@app.route("/prioridades/<int:cid>")
@login_required
def prioridades_loja(cid):
    if not pode_ver_loja(cid):
        return pagina('Acesso negado','<div class="card">Acesso restrito</div>'),403
    is_admin=session.get('tipo')=='admin'
    if not is_admin:
        with banco() as conn:
            enabled=consultar(conn,"SELECT habilitado FROM cliente_prioridades_acessos WHERE conexao_id=%s",(cid,))
        if not enabled or not enabled[0]['habilitado']:
            return pagina('Acesso não liberado','<div class="card">A Central de Prioridades ainda não foi liberada para esta loja.</div>'),403
    from client_priorities import store_portal
    from management_data import load_management
    from overview_read import TIMEOUT_MESSAGE
    try:
        report=load_management(cid,{'dias':'7'},banco,consultar)
        if report is None:
            return pagina('Loja não encontrada','<div class="card">Loja não encontrada</div>'),404
        body=store_portal(report,cid,client_view=not is_admin)
    except ValueError:
        return pagina('Prioridades indisponíveis','<div class="card">'+esc(TIMEOUT_MESSAGE)+'</div>'),503
    response=Response(pagina('Central de Prioridades',body),mimetype='text/html')
    response.headers['Cache-Control']='no-store'
    return response


@app.route("/prioridades/<int:cid>/acesso",methods=['POST'])
@admin_required
def acesso_prioridades_cliente(cid):
    expected=session.get('prioridades_csrf')
    token=request.form.get('csrf_token','')
    if not expected or not secrets.compare_digest(expected,token):
        return 'Solicitação inválida',403
    flag=request.form.get('enabled')
    if flag not in ('0','1'):
        return 'Parâmetro inválido',400
    with banco() as conn:
        exists=consultar(conn,'SELECT id FROM conexoes WHERE id=%s',(cid,))
        if not exists:
            return 'Loja não encontrada',404
        executar(conn,"""INSERT INTO cliente_prioridades_acessos (conexao_id,habilitado,atualizado_em)
            VALUES (%s,%s,%s) ON CONFLICT (conexao_id) DO UPDATE SET
            habilitado=EXCLUDED.habilitado,atualizado_em=EXCLUDED.atualizado_em""",
            (cid,flag=='1',int(time.time())))
        if flag=='1':
            executar(conn,"""INSERT INTO cliente_onboarding_financeiro(usuario_id,conexao_id)
                SELECT id,conexao_id FROM usuarios WHERE tipo='cliente' AND conexao_id=%s
                ON CONFLICT(usuario_id) DO NOTHING""",(cid,))
    return redirect('/prioridades')


@app.route("/painel/tendencias")
@login_required
def painel_tendencias():
    """Lightweight, separately loaded seven-day sales direction per authorized shop."""
    from zoneinfo import ZoneInfo
    hoje = datetime.datetime.now(ZoneInfo("America/Sao_Paulo")).date()
    fim = hoje.isoformat()
    inicio_atual = (hoje - datetime.timedelta(days=7)).isoformat()
    inicio_anterior = (hoje - datetime.timedelta(days=14)).isoformat()
    with banco() as conn:
        if session.get("tipo") == "admin":
            lojas = consultar(conn, "SELECT id FROM conexoes")
        else:
            lojas = consultar(conn, "SELECT id FROM conexoes WHERE id=%s",
                               (session.get("conexao_id"),))
        ids = [int(loja["id"]) for loja in lojas if pode_ver_loja(loja["id"])]
        if not ids:
            response = jsonify({})
            response.headers["Cache-Control"] = "private, max-age=60"
            return response
        linhas = consultar(conn, """
            SELECT conexao_id,
                   COALESCE(SUM(CASE WHEN LEFT(fechado_em,10) >= %s
                                      AND LEFT(fechado_em,10) < %s
                                THEN total ELSE 0 END),0) AS atual,
                   COALESCE(SUM(CASE WHEN LEFT(fechado_em,10) >= %s
                                      AND LEFT(fechado_em,10) < %s
                                THEN total ELSE 0 END),0) AS anterior
            FROM pedidos
            WHERE conexao_id = ANY(%s)
              AND fechado_em IS NOT NULL
              AND LEFT(fechado_em,10) >= %s
              AND LEFT(fechado_em,10) < %s
              AND (status IS NULL OR status NOT IN ('cancelled', 'cancelada'))
            GROUP BY conexao_id
        """, (inicio_atual, fim, inicio_anterior, inicio_atual,
              ids, inicio_anterior, fim))
    tendencias = {}
    for linha in linhas:
        atual = float(linha["atual"] or 0)
        anterior = float(linha["anterior"] or 0)
        tendencias[str(linha["conexao_id"])] = (
            "up" if atual > anterior else "down" if atual < anterior else "flat"
        )
    response = jsonify(tendencias)
    response.headers["Cache-Control"] = "private, max-age=60"
    return response


@app.route("/painel/analise")
@login_required
def painel_analise():
    with banco() as conn:
        if session.get("tipo") == "admin":
            lojas = consultar(conn, "SELECT id, cliente, status FROM conexoes ORDER BY cliente, id")
        else:
            lojas = consultar(conn, "SELECT id, cliente, status FROM conexoes WHERE id=%s ORDER BY cliente, id",
                              (session.get("conexao_id"),))
    # Keep the existing access policy and apply it before any sales/coverage read.
    lojas = [loja for loja in lojas if pode_ver_loja(loja["id"])]
    admin = session.get("tipo") == "admin"
    inicio_painel = time.perf_counter()
    try:
        from overview_cache import load_cached_overview
        dados = load_cached_overview(lojas, request.args, banco, consultar, load_overview)
        apos_vendas = time.perf_counter()
        from visits_overview import load_visits_card
        visitas_card = load_visits_card(banco, lojas)
        apos_visitas = time.perf_counter()
        corpo = render_overview(dados, request.args, admin=admin, visits_card=visitas_card)
        apos_render = time.perf_counter()
        if (apos_render - inicio_painel) > 0.5:
            app.logger.warning(
                "painel_etapas lojas=%d vendas_ms=%.0f visitas_ms=%.0f render_ms=%.0f",
                len(lojas),(apos_vendas-inicio_painel)*1000,
                (apos_visitas-apos_vendas)*1000,(apos_render-apos_visitas)*1000)
    except ValueError as erro:
        from sales_view import filters
        corpo = ("<div class='aviso'>" + esc(str(erro))
                 + "</div>" + filters(None, request.args, action="/painel"))
    if admin:
        corpo = HOME_WIDGET + corpo
    return pagina("Visão Geral 2.0", overview_layout(lojas, corpo, admin=admin))


@app.route("/painel/legado")
@login_required
def painel_legado():
    hoje_iso = datetime.date.today().isoformat()
    ontem_iso = (datetime.date.today()
                 - datetime.timedelta(days=1)).isoformat()
    ini7 = (datetime.date.today()
            - datetime.timedelta(days=6)).isoformat()
    with banco() as conn:
        if session.get("tipo") == "admin":
            lojas = consultar(conn, "SELECT id, cliente, status "
                                    "FROM conexoes ORDER BY cliente")
        else:
            lojas = consultar(conn, "SELECT id, cliente, status "
                                    "FROM conexoes WHERE id=%s "
                                    "ORDER BY cliente",
                              (session.get("conexao_id"),))
        disparou = False
    agora = int(time.time())
    for l in lojas:
        ult = ultima_atualizacao(l["id"])
        if ult is None or (agora - ult) > GATILHO_AUTO_SEG:
            if enfileirar_atualizacao(l["id"]):
                disparou = True
    if not lojas:
        corpo = ("<div class='card' style='text-align:center;padding:50px'>"
                 "<h2>Nenhuma loja conectada ainda</h2>"
                 "<div class='sub'>Comece conectando a primeira loja</div>"
                 "<p><a class='link' href='/conectar'>Conectar loja</a></p></div>")
        return pagina("Painel", corpo)
    ids = [l["id"] for l in lojas]
    fat_hoje = fat_ontem = 0.0
    vendas_hoje = vendas_ontem = 0
    visitas_total = None
    gasto_hoje = None
    gasto_total = 0.0
    fat_geral = 0.0
    g_fat = []
    if ids:
        ph = ",".join(["%s"] * len(ids))
        with banco() as conn:
            r = consultar(conn, "SELECT COALESCE(SUM(total),0) AS t, "
                                "COUNT(*) AS n FROM pedidos "
                                "WHERE conexao_id IN (" + ph + ") "
                                "AND fechado_em IS NOT NULL "
                                "AND SUBSTRING(fechado_em,1,10)=%s",
                          (*ids, hoje_iso))
            fat_hoje = r[0]["t"] or 0
            vendas_hoje = r[0]["n"] or 0
            r = consultar(conn, "SELECT COALESCE(SUM(total),0) AS t, "
                                "COUNT(*) AS n FROM pedidos "
                                "WHERE conexao_id IN (" + ph + ") "
                                "AND fechado_em IS NOT NULL "
                                "AND SUBSTRING(fechado_em,1,10)=%s",
                          (*ids, ontem_iso))
            fat_ontem = r[0]["t"] or 0
            vendas_ontem = r[0]["n"] or 0
            r = consultar(conn, "SELECT COALESCE(SUM(faturamento),0) AS g "
                                "FROM ads_dia "
                                "WHERE conexao_id IN (" + ph + ") "
                                "AND data=%s", (*ids, hoje_iso))
            gasto_hoje = r[0]["g"]
            r = consultar(conn, "SELECT COALESCE(SUM(gasto),0) AS g "
                                "FROM ads_dia "
                                "WHERE conexao_id IN (" + ph + ")", ids)
            gasto_total = r[0]["g"] or 0
            r = consultar(conn, "SELECT COALESCE(SUM(total),0) AS t "
                                "FROM pedidos "
                                "WHERE conexao_id IN (" + ph + ") "
                                "AND fechado_em IS NOT NULL", ids)
            fat_geral = r[0]["t"] or 0
            g_fat = consultar(conn, "SELECT SUBSTRING(fechado_em,1,10) AS dia, "
                                    "COALESCE(SUM(total),0) AS t "
                                    "FROM pedidos "
                                    "WHERE conexao_id IN (" + ph + ") "
                                    "AND fechado_em IS NOT NULL "
                                    "AND SUBSTRING(fechado_em,1,10) >= %s "
                                    "GROUP BY SUBSTRING(fechado_em,1,10)",
                              (*ids, ini7))
    fat_map = {r["dia"]: r["t"] for r in g_fat}
    rotulos, serie_fat, serie_vis = [], [], []
    d = datetime.date.today() - datetime.timedelta(days=6)
    while d <= datetime.date.today():
        iso = d.isoformat()
        rotulos.append(d.strftime("%d/%m"))
        serie_fat.append(round(fat_map.get(iso, 0), 2))
        serie_vis.append(None)
        d += datetime.timedelta(days=1)
    tacos = None
    if fat_geral and fat_geral > 0:
        tacos = round(gasto_total * 100.0 / fat_geral, 1)
    lista = "<div class='menu-lateral' style='width:250px'>"
    lista += "<div class='voltar'>Contas conectadas</div>"
    for l in lojas:
        if str(l["status"]) == "ativa":
            cor = "#1a9d5c"
        else:
            cor = "#d64545"
        lista += ("<div style='display:flex;align-items:center'>"
                  "<a class='item' style='flex:1' href='/painel/"
                  + str(l["id"]) + "'>"
                  "<span style='display:inline-block;width:8px;height:8px;"
                  "border-radius:50%;background:" + cor
                  + ";margin-right:8px'></span>"
                  + esc(str(l["cliente"])) + "</a>"
                  "<a class='icon-link' href='#' onclick='editarNome("
                  + str(l["id"]) + "); return false;'>✏️</a></div>")
    if session.get("tipo") == "admin":
        lista += ("<div style='border-top:1px solid #eef1f5;"
                  "margin-top:10px;padding-top:8px'></div>")
        lista += "<a class='item' href='/usuarios'>👥 Usuários</a>"
        lista += "<a class='item' href='/diagnostico'>🔍 Diagnóstico</a>"
    lista += "</div>"
    corpo = "<h2 style='margin-bottom:14px'>Visão geral</h2>"
    corpo += "<div class='grid' style='margin-bottom:18px'>"
    corpo += ("<div class='stat'><div class='num'>R$ " + br(fat_hoje)
              + "</div><div class='lab'>Faturamento hoje (todas as contas)</div>"
              + var_html(fat_hoje, fat_ontem) + "</div>")
    corpo += ("<div class='stat'><div class='num'>" + br(vendas_hoje, 0)
              + "</div><div class='lab'>Vendas hoje</div>"
              + "<div class='var estavel'>ontem: "
              + br(vendas_ontem, 0) + "</div></div>")
    corpo += ("<div class='stat'><div class='num'>Indisponível"
              "</div><div class='lab'>Visitas totais (API indisponível)"
              "</div></div>")
    if gasto_hoje is not None:
        gasto_txt = "R$ " + br(gasto_hoje)
    else:
        gasto_txt = "<span class='tag sem'>—</span>"
    corpo += ("<div class='stat'><div class='num'>" + gasto_txt
              + "</div><div class='lab'>Faturamento Ads hj</div></div>")
    corpo += ("<div class='stat'><div class='num'>R$ " + br(gasto_total)
              + "</div><div class='lab'>Investimento total em Ads</div></div>")
    if tacos is not None:
        tacos_txt = br(tacos, 1) + "%"
    else:
        tacos_txt = "<span class='tag sem'>—</span>"
    corpo += ("<div class='stat'><div class='num'>" + tacos_txt
              + "</div><div class='lab'>TACOS geral (Ads ÷ faturamento)"
              + "</div></div>")
    corpo += "</div>"
    corpo += ("<div class='card'><h2>Evolução dos últimos 7 dias</h2>"
              "<div class='sub'>Faturamento e visitas somando todas as contas"
              "</div><div class='grafico-box'>"
              "<canvas id='graficoGeral'></canvas></div></div>")
    corpo += ("<script src='https://cdn.jsdelivr.net/npm/"
              "chart.js@4.4.1/dist/chart.umd.min.js'></script>")
    corpo += ("<script>var rot = " + json.dumps(rotulos) + ";"
              "var sf = " + json.dumps(serie_fat) + ";"
              "new Chart(document.getElementById('graficoGeral'), {"
              "type:'line',"
              "data:{labels:rot,datasets:["
              "{label:'Faturamento registrado (R$)',data:sf,borderColor:'#3483FA',"
              "backgroundColor:'rgba(52,131,250,0.12)',tension:0.3,fill:true}]},"
              "options:{responsive:true,maintainAspectRatio:false,"
              "plugins:{legend:{position:'top'}},"
              "scales:{y:{beginAtZero:true,title:{display:true,text:'R$'}}}}});</script>")
    try:
        tenta_r = int(request.args.get("r", 0))
    except Exception:
        tenta_r = 0
    if disparou:
         if tenta_r < 10:
            corpo += ("<div class='aviso'>Atualizando contas em segundo plano "
                      "(restam " + str(len(FILA_ATUALIZACAO))
                      + " na fila) — a página recarrega sozinha.</div>")
            corpo += ("<script>setTimeout(function(){ location.href = "
                      "'/painel/legado?r=" + str(tenta_r + 1) + "'; }, 60000);</script>")
         else:
            corpo += ("<div class='aviso'>A atualização ainda está rodando em "
                      "segundo plano. Recarregue a página em instantes.</div>")
    corpo += ("<div class='muted'>TACOS = investimento em Ads dividido pelo "
              "faturamento total. Clique em uma conta na lista à esquerda "
              "para abrir o Resumo dela.</div>")
    corpo_final = ("<div class='layout'>" + lista
                   + "<div class='conteudo'>" + corpo + "</div></div>")
    return pagina("Painel", corpo_final)


@app.route("/usuarios")
@admin_required
def usuarios():
    with banco() as conn:
        us = consultar(conn, """SELECT u.id, u.username, u.nome, u.email,
                                u.tipo, u.conexao_id, c.cliente
                                FROM usuarios u
                                LEFT JOIN conexoes c ON c.id=u.conexao_id
                                ORDER BY u.id""")
        lojas = consultar(conn, "SELECT id, cliente FROM conexoes ORDER BY id")
    msg = request.args.get("msg", "")
    corpo = ("<p style='margin-bottom:10px'><a class='link' href='/painel'>"
             "← Voltar ao painel</a></p>"
             "<h2 style='margin-bottom:14px'>Gerenciar usuários</h2>")
    if msg:
        corpo += "<div class='sucesso'>" + esc(msg) + "</div>"
    corpo += ("<div class='card'><h2>Novo usuário</h2>"
              "<div class='sub'>Crie um acesso. Cliente vê apenas a loja "
              "vinculada; admin vê tudo. Se preencher o e-mail, o usuário "
              "poderá recuperar a senha sozinho.</div>"
              "<form method='post' action='/usuarios/criar' "
              "style='display:flex;flex-wrap:wrap;gap:10px;align-items:center'>"
              "<input name='username' required placeholder='Usuário' class='filtro'>"
              "<input type='password' name='senha' id='campo-senha' "
              "required placeholder='Senha' class='filtro'>"
              "<button type='button' class='filtro-btn' onclick='gerarSenha()' "
              "style='background:#6b7280'>🎲 Gerar senha</button>"
              "<input name='nome' placeholder='Nome (opcional)' class='filtro'>"
              "<input name='email' type='email' "
              "placeholder='E-mail (p/ recuperação)' class='filtro'>"
              "<select name='tipo' class='filtro'>"
              "<option value='cliente'>Cliente</option>"
              "<option value='admin'>Admin</option></select>"
              "<select name='conexao_id' class='filtro'>"
              "<option value=''>— sem loja —</option>"
              + "".join("<option value='" + str(l["id"]) + "'>"
                        + esc(str(l["cliente"])) + "</option>" for l in lojas)
              + "</select>"
              "<button class='filtro-btn'>Criar</button>"
              "</form></div>")
    linhas_html = ""
    for u in us:
        if u["tipo"] == "admin":
            tipo_txt = "Admin"
            tag = "tag ativa"
        else:
            tipo_txt = "Cliente"
            tag = "tag pendente"
        if u["cliente"]:
            loja_txt = esc(str(u["cliente"]))
        else:
            loja_txt = "<span class='tag sem'>—</span>"
        if u["email"]:
            email_txt = esc(str(u["email"]))
        else:
            email_txt = "<span class='tag sem'>—</span>"
        linhas_html += ("<tr><td><b>" + esc(str(u["username"])) + "</b></td>"
                        "<td>" + esc(str(u["nome"] or "-")) + "</td>"
                        "<td>" + email_txt + "</td>"
                        "<td><span class='" + tag + "'>" + tipo_txt + "</span></td>"
                        "<td>" + loja_txt + "</td>"
                        "<td><a class='btn-acoes cinza' href='/usuarios/editar/"
                        + str(u["id"]) + "'>Editar</a>"
                        " <a class='btn-danger' href='/usuarios/excluir/"
                        + str(u["id"]) + "'>Excluir</a></td></tr>")
    corpo += ("<div class='card'><h2>Usuários <span class='badge'>"
              + str(len(us)) + "</span></h2>"
              "<table><thead><tr><th>Usuário</th><th>Nome</th><th>E-mail</th>"
              "<th>Tipo</th><th>Loja</th><th>Ações</th></tr></thead>"
              "<tbody>" + linhas_html + "</tbody></table></div>")
    return pagina("Usuários", corpo)


@app.route("/usuarios/criar", methods=["POST"])
@admin_required
def usuarios_criar():
    username = (request.form.get("username") or "").strip()
    senha = request.form.get("senha") or ""
    nome = (request.form.get("nome") or "").strip()
    email = (request.form.get("email") or "").strip()
    tipo = request.form.get("tipo") or "cliente"
    conexao_id = request.form.get("conexao_id") or None
    if not username or not 12 <= len(senha) <= 256:
        return redirect("/usuarios?msg=Informe usuário e senha de 12 a 256 caracteres")
    if tipo not in ("admin", "cliente"):
        return "Tipo de usuário inválido", 400
    try:
        conexao_id = int(conexao_id) if conexao_id else None
    except Exception:
        conexao_id = None
    with banco() as conn:
        executar(conn, "INSERT INTO usuarios "
                       "(username, senha_hash, nome, email, tipo, conexao_id, criado_em) "
                       "VALUES (%s,%s,%s,%s,%s,%s,%s)",
                 (username, generate_password_hash(senha), nome, email,
                  tipo, conexao_id, int(time.time())))
        if tipo=='cliente' and conexao_id is not None:
            new_user=consultar(conn,"SELECT id FROM usuarios WHERE username=%s",(username,))
            executar(conn,"""INSERT INTO cliente_onboarding_financeiro(usuario_id,conexao_id)
                VALUES (%s,%s) ON CONFLICT(usuario_id) DO NOTHING""",
                (new_user[0]['id'],conexao_id))
    return redirect("/usuarios?msg=Usuário criado com sucesso")


@app.route("/usuarios/editar/<int:uid>", methods=["GET", "POST"])
@admin_required
def usuarios_editar(uid):
    with banco() as conn:
        u = consultar(conn, "SELECT * FROM usuarios WHERE id=%s", (uid,))
        u = u[0] if u else None
        lojas = consultar(conn, "SELECT id, cliente FROM conexoes ORDER BY id")
    if not u:
        return pagina("Não encontrado",
                      "<div class='card'><h2>Usuário não encontrado</h2></div>")
    if request.method == "POST":
        nome = (request.form.get("nome") or "").strip()
        email = (request.form.get("email") or "").strip()
        tipo = request.form.get("tipo") or "cliente"
        conexao_id = request.form.get("conexao_id") or None
        nova_senha = request.form.get("senha") or ""
        if (nova_senha and not 12 <= len(nova_senha) <= 256) or tipo not in ("admin", "cliente"):
            return "Informe uma senha de 12 a 256 caracteres e um tipo válido.", 400
        try:
            conexao_id = int(conexao_id) if conexao_id else None
        except Exception:
            conexao_id = None
        with banco() as conn:
            if nova_senha:
                executar(conn, "UPDATE usuarios SET nome=%s, email=%s, tipo=%s, "
                               "conexao_id=%s, senha_hash=%s WHERE id=%s",
                         (nome, email, tipo, conexao_id,
                          generate_password_hash(nova_senha), uid))
            else:
                executar(conn, "UPDATE usuarios SET nome=%s, email=%s, tipo=%s, "
                               "conexao_id=%s WHERE id=%s",
                         (nome, email, tipo, conexao_id, uid))
        return redirect("/usuarios?msg=Usuário atualizado")
    corpo = ("<h2 style='margin-bottom:14px'>Editar usuário: "
             + esc(str(u["username"])) + "</h2>"
             "<div class='card'><form method='post' action='/usuarios/editar/"
             + str(uid) + "' style='display:flex;flex-wrap:wrap;gap:10px;"
             "align-items:center'>"
             "<input name='nome' placeholder='Nome' value='"
             + esc(str(u["nome"] or "")) + "' class='filtro'>"
             "<input name='email' type='email' "
             "placeholder='E-mail (p/ recuperação)' value='"
             + esc(str(u["email"] or "")) + "' class='filtro'>"
             "<input type='password' name='senha' id='campo-senha' "
             "placeholder='Nova senha (vazio p/ manter)' class='filtro'>"
             "<button type='button' class='filtro-btn' onclick='gerarSenha()' "
             "style='background:#6b7280'>🎲 Gerar senha</button>"
             "<select name='tipo' class='filtro'>"
             "<option value='cliente'"
             + (" selected" if u["tipo"] != "admin" else "")
             + ">Cliente</option>"
             "<option value='admin'"
             + (" selected" if u["tipo"] == "admin" else "")
             + ">Admin</option></select>"
             "<select name='conexao_id' class='filtro'>"
             "<option value=''>— sem loja —</option>"
             + "".join("<option value='" + str(l["id"]) + "'"
                       + (" selected" if u["conexao_id"] == l["id"] else "")
                       + ">" + esc(str(l["cliente"])) + "</option>" for l in lojas)
             + "</select>"
             "<button class='filtro-btn'>Salvar</button>"
             "</form></div>"
             "<p><a class='link' href='/usuarios'>← Voltar</a></p>")
    return pagina("Editar usuário", corpo)


@app.route("/usuarios/excluir/<int:uid>", methods=["GET", "POST"])
@admin_required
def usuarios_excluir(uid):
    with banco() as conn:
        u = consultar(conn, "SELECT * FROM usuarios WHERE id=%s", (uid,))
        u = u[0] if u else None
    if u and u["username"] == "admin":
        return redirect("/usuarios?msg=Não é possível excluir o usuário admin")
    if u and request.method == "GET":
        return pagina("Excluir usuário", "<div class='card'><h2>Excluir usuário "
                      + esc(str(u["username"])) + "?</h2><form method='post'>"
                      "<button class='btn-danger'>Confirmar exclusão</button></form>"
                      "<p><a href='/usuarios'>Cancelar</a></p></div>")
    if u:
        with banco() as conn:
            executar(conn, "DELETE FROM reset_tokens WHERE usuario_id=%s", (uid,))
            executar(conn, "DELETE FROM usuarios WHERE id=%s", (uid,))
    return redirect("/usuarios?msg=Usuário excluído")


@app.route("/excluir/<int:cid>", methods=["GET", "POST"])
@admin_required
def excluir(cid):
    with banco() as conn:
        c = consultar(conn, "SELECT * FROM conexoes WHERE id=%s", (cid,))
        c = c[0] if c else None
    if not c:
        return pagina("Não encontrada",
                      "<div class='card'><h2>Loja não encontrada</h2></div>")
    if request.method == "POST":
        with banco() as conn:
            for tabela in ("full_observations", "dados_conta", "anuncios", "pedidos", "pedidos_itens",
                           "metricas", "desempenho_historico", "perguntas",
                           "envios", "promocoes", "promocoes_itens",
                           "ads_campanhas", "ads_metricas", "ads_dia",
                           "ads_campanha_dia", "erros"):
                executar(conn, "DELETE FROM " + tabela + " WHERE conexao_id=%s", (cid,))
            executar(conn, "DELETE FROM conexoes WHERE id=%s", (cid,))
        corpo = ("<div class='card' style='max-width:520px;margin:40px auto;"
                 "text-align:center'>"
                 "<h2>Loja excluída com sucesso!</h2>"
                 "<div class='sub'>Todos os dados dessa loja foram removidos</div>"
                 "<p style='margin-top:16px'><a class='link' href='/painel'>"
                 "Voltar ao painel</a></p></div>")
        return pagina("Excluída", corpo)
    corpo = ("<div class='card' style='max-width:520px;margin:40px auto;"
             "text-align:center'>"
             "<h2>Excluir loja de " + esc(str(c["cliente"])) + "?</h2>"
             "<div class='sub'>Essa ação remove a conexão e todos os dados "
             "da loja. Não pode ser desfeita.</div>"
             "<form method='post' action='/excluir/" + str(cid)
             + "' style='margin-top:16px'>"
             "<button class='btn-danger' style='padding:13px 24px;font-size:15px'>"
             "Sim, excluir</button>"
             "</form>"
             "<p style='margin-top:14px'><a class='link' href='/painel'>"
             "← Cancelar</a></p>"
             "</div>")
    return pagina("Excluir loja", corpo)


@app.route("/ads/<int:cid>/campanha/<path:campanha_id>")
@login_required
def painel_ads_campanha(cid, campanha_id):
    """Campaign detail: scoped DB-only reporting, no on-demand API calls."""
    if not pode_ver_loja(cid):
        return pagina("Acesso negado", "<div class='card'>Acesso restrito</div>"), 403
    if not campanha_id or len(campanha_id) > 120 or "/" in campanha_id:
        return pagina("Campanha inválida", "<div class='card'>Código inválido.</div>"), 400
    from sales_metrics import reporting_today
    hoje = reporting_today()
    dias = request.args.get("dias", "30")
    if dias not in ("7", "15", "30", "60", "90"):
        dias = "30"
    inicio = (hoje - datetime.timedelta(days=int(dias)*2-1)).isoformat()
    with banco() as conn:
        loja = consultar(conn, "SELECT cliente FROM conexoes WHERE id=%s", (cid,))
        camp = consultar(conn, """SELECT campanha_id,nome,status,orcamento,roas_objetivo,tipo,atualizado_em
            FROM ads_campanhas WHERE conexao_id=%s AND campanha_id=%s""", (cid,campanha_id))
        if loja and camp:
            diario = consultar(conn, """SELECT data,gasto,faturamento,atualizado_em,
                vendas_atribuidas,outras_vendas,faturamento_sem_ads
                FROM ads_campanha_dia WHERE conexao_id=%s AND campanha_id=%s
                AND data BETWEEN %s AND %s ORDER BY data""",
                (cid,campanha_id,inicio,hoje.isoformat()))
            historico_alteracoes = consultar(conn, """SELECT campo,valor_anterior,
                valor_novo,observado_em,origem FROM ads_campanha_alteracoes
                WHERE conexao_id=%s AND campanha_id=%s ORDER BY observado_em DESC,id DESC LIMIT 50""",
                (cid,campanha_id))
    if not loja:
        return pagina("Não encontrada", "<div class='card'>Loja não encontrada</div>"), 404
    if not camp:
        return pagina_loja("Campanha não encontrada", cid, loja[0]["cliente"],
            "ads", "<div class='card'>Campanha não encontrada para esta loja.</div>"), 404
    from ads_campaign_detail import render
    corpo = render(cid, camp[0], diario, int(dias), hoje,
                   changes=historico_alteracoes)
    return pagina_loja("Campanha Ads", cid, loja[0]["cliente"], "ads", corpo)


@app.route("/ads/<int:cid>")
@login_required
def painel_ads(cid):
    """Read-only ads overview using stored daily API metrics only."""
    if not pode_ver_loja(cid):
        return pagina("Acesso negado", "<div class='card'>Acesso restrito</div>"), 403
    from urllib.parse import urlencode
    from sales_metrics import reporting_today
    from ads_overview_design import header, alert_card, daily_chart, ratio, budget_progress
    hoje = reporting_today()
    dias = request.args.get("dias", "30")
    if dias not in ("7", "30", "90"):
        dias = "30"
    inicio = hoje - datetime.timedelta(days=int(dias) - 1)
    fim = hoje
    de, ate = request.args.get("de", ""), request.args.get("ate", "")
    if de or ate:
        try:
            inicio = datetime.date.fromisoformat(de)
            fim = datetime.date.fromisoformat(ate)
            if fim < inicio or fim > hoje or (fim - inicio).days > 89:
                raise ValueError()
        except ValueError:
            return pagina("Ads", "<div class='card'>Período inválido: selecione até 90 dias, sem datas futuras.</div>"), 400
    inicio_iso, fim_iso = inicio.isoformat(), fim.isoformat()
    with banco() as conn:
        lojas = consultar(conn, "SELECT id, cliente FROM conexoes WHERE id=%s", (cid,))
        if not lojas:
            return pagina("Não encontrada", "<div class='card'>Loja não encontrada</div>"), 404
        dias_ads = consultar(conn, """
            SELECT data, gasto, faturamento, impressoes, cliques, atualizado_em
            FROM ads_dia WHERE conexao_id=%s AND data BETWEEN %s AND %s
            ORDER BY data DESC
        """, (cid, inicio_iso, fim_iso))
    # Local business day and observation freshness matter for budget pacing.
    from zoneinfo import ZoneInfo
    agora_local = datetime.datetime.now(ZoneInfo("America/Sao_Paulo"))
    ontem = (hoje - datetime.timedelta(days=1)).isoformat()
    with banco() as conn:
        alert_rows = consultar(conn, """
            SELECT d.campanha_id,
                COALESCE(NULLIF(c.nome,''), 'Campanha ' || d.campanha_id) AS nome,
                c.status, c.orcamento, c.roas_objetivo,
                MAX(e.detectado_em) AS esgotado_em,
                COALESCE(SUM(d.gasto) FILTER (WHERE d.data BETWEEN %s AND %s),0) AS gasto_periodo,
                COALESCE(SUM(d.faturamento) FILTER (WHERE d.data BETWEEN %s AND %s),0) AS fat_periodo,
                COALESCE(SUM(d.gasto) FILTER (WHERE d.data=%s),0) AS gasto_hoje,
                MAX(d.atualizado_em) FILTER (WHERE d.data=%s) AS hoje_atualizado,
                COALESCE(SUM(d.gasto) FILTER (WHERE d.data=%s),0) AS gasto_ontem,
                COALESCE(SUM(d.faturamento) FILTER (WHERE d.data=%s),0) AS fat_ontem,
                COUNT(*) FILTER (WHERE d.data=%s) AS tem_ontem
            FROM ads_campanha_dia d
            LEFT JOIN ads_campanhas c ON c.conexao_id=d.conexao_id AND c.campanha_id=d.campanha_id
            LEFT JOIN ads_orcamento_esgotado e ON e.conexao_id=d.conexao_id AND e.campanha_id=d.campanha_id AND e.data=%s
            WHERE d.conexao_id=%s AND d.data BETWEEN %s AND %s
            GROUP BY d.campanha_id,c.nome,c.status,c.orcamento,c.roas_objetivo
        """, (inicio_iso,fim_iso,inicio_iso,fim_iso,hoje.isoformat(),
              hoje.isoformat(),ontem,ontem,ontem,hoje.isoformat(),
              cid,min(inicio_iso,ontem),max(fim_iso,hoje.isoformat())))
    alertas_orcamento, alertas_risco, alertas_roas, alertas_sem_venda = [], [], [], []
    campanhas_configuradas = sum(1 for camp in alert_rows if camp["status"] is not None)
    for camp in alert_rows:
        nome = esc(str(camp["nome"]))
        ativa = str(camp["status"] or "").lower() in ("active","running")
        budget = float(camp["orcamento"] or 0)
        gasto_hoje = float(camp["gasto_hoje"] or 0)
        updated = camp["hoje_atualizado"]
        dados_recentes = (updated is not None and
                         0 <= time.time() - float(updated) <= 2 * 3600)
        if ativa and budget > 0 and dados_recentes:
            consumo = gasto_hoje / budget
            info = "R$ %.2f de R$ %.2f (%.0f%%)" % (gasto_hoje,budget,consumo*100)
            info = info.replace(".", ",")
            if consumo >= 1:
                alertas_orcamento.append((nome,info))
            elif consumo >= 0.8 and agora_local.hour < 18:
                alertas_risco.append((nome,info+" · atenção antes das 18h"))
        gasto = float(camp["gasto_periodo"] or 0)
        fat = float(camp["fat_periodo"] or 0)
        alvo = float(camp["roas_objetivo"] or 0)
        if ativa and alvo > 0 and gasto > 0 and fat/gasto >= 1.3*alvo:
            alertas_roas.append((nome,"ROAS %.2fx · objetivo %.2fx" % (fat/gasto,alvo)))
        if camp["tem_ontem"] and float(camp["gasto_ontem"] or 0) > 0 and float(camp["fat_ontem"] or 0) == 0:
            alertas_sem_venda.append((nome,"R$ %.2f investidos ontem · nenhuma venda atribuída" %
                                     float(camp["gasto_ontem"] or 0)))
    def quadro_alertas(titulo, descricao, itens, vazio):
        return alert_card(titulo, descricao, itens, vazio, opportunity='ROAS' in titulo)
    alertas_html = ("<div class='ka-section-title'><h3>Oportunidades e atenção</h3><span>Leituras recentes de Ads</span></div>"
        + ("<div class='aviso'>Configuração de campanhas ainda não sincronizada nesta loja. "
           "Alertas de orçamento e ROAS objetivo ficam pendentes.</div>"
           if not campanhas_configuradas else "")
        + "<div class='ka-alert-grid'>"
        + quadro_alertas("Orçamento diário esgotado",
            "Campanhas ativas cujo gasto observado hoje atingiu o orçamento diário; leitura recente (até 2h).",
            alertas_orcamento,"Nenhuma campanha com limite atingido e dados recentes.")
        + quadro_alertas("Risco de orçamento esgotar cedo",
            "Mais de 80% do orçamento consumido antes das 18h, com leitura recente. "
            "É um sinal de atenção, não uma previsão de horário.",
            alertas_risco,"Nenhuma campanha nesta condição com dados recentes.")
        + quadro_alertas("ROAS pelo menos 30% acima do objetivo",
            "Considera o ROAS do período selecionado e a meta cadastrada.",
            alertas_roas,"Nenhuma campanha identificada ou objetivo ainda não importado.")
        + quadro_alertas("Sem vendas de Ads ontem",
            "Campanhas com investimento ontem e nenhuma venda atribuída.",
            alertas_sem_venda,"Nenhuma campanha nesta condição.") + "</div>")
    titulo = "Publicidade · Product Ads"
    abertura = header(cid, lojas[0]["cliente"], inicio, fim)
    filtro = ("<section class='card'><h2>Filtros de publicidade</h2>"
             "<form method='get' action='/ads/" + str(cid) + "' "
             "style='display:flex;gap:10px;align-items:end;flex-wrap:wrap'>"
             "<label>Período <select name='dias'>"
             + "".join("<option value='" + n + "'" + (" selected" if dias == n else "")
                       + ">" + n + " dias</option>" for n in ("7","30","90"))
             + "</select></label><label>De <input type='date' name='de' value='"
             + esc(de) + "'></label><label>Até <input type='date' name='ate' value='"
             + esc(ate) + "'></label><button class='btn' type='submit'>Aplicar filtros</button>"
             "</form><p class='muted'>Período exibido: " + inicio.strftime("%d/%m")
             + " a " + fim.strftime("%d/%m") + "</p></section>")
    if not dias_ads:
        corpo = (abertura + filtro
                 + "<section class='card'><h3>Sem dados de publicidade para este período</h3>"
                 "<p>Confira se Product Ads está ativo, se a sincronização da loja terminou "
                 "e se existe investimento no intervalo selecionado.</p>"
                 "<p>Use Atualizar Ads acima para carregar as campanhas e os resultados.</p></section>")
        return pagina_loja("Ads", cid, lojas[0]["cliente"], "ads", corpo + alertas_html + "</div>")
    def soma(campo):
        return sum(float(row[campo] or 0) for row in dias_ads)
    gasto = soma("gasto")
    receita = soma("faturamento")
    impressoes = int(soma("impressoes"))
    cliques = int(soma("cliques"))
    roas = receita / gasto if gasto > 0 else None
    acos = gasto / receita * 100 if receita > 0 else None
    ctr = cliques / impressoes * 100 if impressoes > 0 else None
    cpc = gasto / cliques if cliques > 0 else None
    def brl(valor):
        return "R$ " + format(float(valor), ",.2f").replace(",", "X").replace(".", ",").replace("X", ".")
    # TACOS requires verified total shop sales for the *same* selected period.
    # Reuse the sales overview classifier and coverage rules; never substitute
    # attributed Ads revenue (which would be ACOS, not TACOS).
    faturamento_loja = None
    motivo_tacos = None
    try:
        from overview import load_overview
        argumentos_vendas = {
            "periodo": "personalizado",
            "inicio": inicio_iso,
            "fim": fim_iso,
        }
        loja_vendas = [{"id": cid, "cliente": lojas[0]["cliente"], "status": "ativa"}]
        dados_vendas = load_overview(loja_vendas, argumentos_vendas, banco, consultar)
        resumo_vendas = dados_vendas["reports"][0]
        atual = resumo_vendas["current"]
        dias_venda = (fim - inicio).days + 1
        # A previous-period sync is not required to evaluate the selected
        # current period, but completeness, integrity and freshness are.
        if (not atual["complete"] or atual["invalid"] or atual["unallocated_dates"]
                or atual["revenue"] is None):
            motivo_tacos = "Histórico de vendas incompleto ou com inconsistências no período."
        else:
            faturamento_loja = float(atual["revenue"])
    except (ValueError, KeyError, TypeError, IndexError, OSError) as exc:
        app.logger.warning("ads_tacos_sales_unavailable loja=%s tipo=%s", cid, type(exc).__name__)
        motivo_tacos = "Não foi possível validar o faturamento total da loja."
    tacos_geral = gasto / faturamento_loja * 100 if faturamento_loja and faturamento_loja > 0 else None
    # Filter and page campaign comparisons without changing the sales baseline.
    filtro_nome = (request.args.get("campanha", "") or "").strip()[:80]
    situacao = request.args.get("situacao", "todas")
    if situacao not in ("todas", "com_gasto", "sem_vendas", "com_vendas", "ativas_sem_vendas"):
        situacao = "todas"
    ordem_camp = request.args.get("ordem_campanha", "gasto")
    if ordem_camp not in ("gasto", "faturamento", "roas", "acos", "tacos"):
        ordem_camp = "gasto"
    linhas_tacos = []
    for camp in alert_rows:
        nome = str(camp["nome"])
        investimento = float(camp["gasto_periodo"] or 0)
        atribuida = float(camp["fat_periodo"] or 0)
        if filtro_nome and filtro_nome.casefold() not in nome.casefold():
            continue
        if situacao == "ativas_sem_vendas" and (str(camp["status"] or "").lower() not in ("active", "running") or atribuida > 0):
            continue
        if situacao == "com_gasto" and investimento <= 0:
            continue
        if situacao == "sem_vendas" and not (investimento > 0 and atribuida <= 0):
            continue
        if situacao == "com_vendas" and atribuida <= 0:
            continue
        roas_camp = atribuida / investimento if investimento > 0 else None
        acos_camp = investimento / atribuida * 100 if atribuida > 0 else None
        tacos_camp = investimento / faturamento_loja * 100 if faturamento_loja and faturamento_loja > 0 else None
        linhas_tacos.append((nome, investimento, atribuida, roas_camp, acos_camp, tacos_camp, str(camp["campanha_id"])))
    ordem_indices = {"gasto":1, "faturamento":2, "roas":3, "acos":4, "tacos":5}
    indice = ordem_indices[ordem_camp]
    linhas_tacos.sort(key=lambda item: (item[indice] is None,
        -(item[indice] or 0), item[0].casefold()))
    try:
        pagina_tacos = max(1, int(request.args.get("pagina_campanhas", "1")))
    except (TypeError, ValueError):
        pagina_tacos = 1
    try:
        tamanho_tacos = int(request.args.get("por_pagina_campanhas", 20))
    except (TypeError, ValueError):
        tamanho_tacos = 20
    if tamanho_tacos not in (10, 20, 50):
        tamanho_tacos = 20
    paginas_tacos = max(1, (len(linhas_tacos) + tamanho_tacos - 1) // tamanho_tacos)
    pagina_tacos = min(pagina_tacos, paginas_tacos)
    def url_tacos(numero):
        params = request.args.to_dict(flat=True)
        params["pagina_campanhas"] = numero
        return "/ads/" + str(cid) + "?" + urlencode(params)
    linhas_tacos_pagina = linhas_tacos[
        (pagina_tacos - 1) * tamanho_tacos:pagina_tacos * tamanho_tacos]
    tacos_campanhas_html = (
        "<section id='ka-campaign-list' class='card kc-campaigns'><h3>Desempenho por campanha</h3>"
        "<p class='muted'>Compare investimento, retorno e participação no faturamento total da loja.</p>"
        "<form method='get' action='/ads/" + str(cid) + "' "
        "style='display:flex;gap:12px;align-items:end;flex-wrap:wrap;margin:14px 0'>"
        + "".join("<input type='hidden' name='" + esc(k,quote=True) + "' value='"
                  + esc(str(request.args.get(k, "")),quote=True) + "'>"
                  for k in ("dias","de","ate") if request.args.get(k))
        + "<label>Buscar campanha <input name='campanha' placeholder='Nome da campanha' value='"
        + esc(filtro_nome,quote=True) + "'></label>"
        + "<label>Situação <select name='situacao'>"
        + "".join("<option value='" + valor + "'" + (" selected" if situacao == valor else "")
            + ">" + rotulo + "</option>" for valor,rotulo in (
                ("todas","Todas"),("com_gasto","Com investimento"),
                ("sem_vendas","Com gasto e sem vendas"),("com_vendas","Com vendas"),("ativas_sem_vendas","Ativas sem vendas")))
        + "</select></label>"
        + "<label>Ordenar por <select name='ordem_campanha'>"
        + "".join("<option value='" + valor + "'" + (" selected" if ordem_camp == valor else "")
            + ">" + rotulo + "</option>" for valor,rotulo in (
                ("gasto","Maior investimento"),("faturamento","Maior faturamento Ads"),
                ("roas","Maior ROAS"),("acos","Maior ACOS"),("tacos","Maior TACOS")))
        + "</select></label>"
        + "<label>Por página <select name='por_pagina_campanhas'>"
        + "".join("<option value='" + str(n) + "'" + (" selected" if tamanho_tacos == n else "")
            + ">" + str(n) + "</option>" for n in (10,20,50))
        + "</select></label><button class='btn' type='submit'>Aplicar</button></form>")
    if faturamento_loja and faturamento_loja > 0:
        tacos_campanhas_html += ("<p>Faturamento total validado da loja: <strong>"
                                 + brl(faturamento_loja) + "</strong></p>")
    else:
        tacos_campanhas_html += ("<div class='aviso'>TACOS indisponível: "
            + esc(motivo_tacos or "Não há faturamento total positivo validado.")
            + "</div>")
    tacos_campanhas_html += (
        "<div class='kc-table-wrap'><table class='kc-campaign-table'><thead><tr>"
        "<th>Campanha</th><th>Investimento</th><th>Faturamento Ads</th>"
        "<th>ROAS</th><th>ACOS</th><th>TACOS</th></tr></thead><tbody>")
    from urllib.parse import quote
    budget_by_id = {str(c["campanha_id"]): c for c in alert_rows}
    for nome, investimento, atribuida, roas_camp, acos_camp, tacos_valor, camp_id in linhas_tacos_pagina:
        detalhe = "/ads/" + str(cid) + "/campanha/" + quote(camp_id, safe="")
        url_segura = esc(detalhe, quote=True)
        budget_row = budget_by_id.get(camp_id, {})
        budget_html = budget_progress(budget_row.get("orcamento"), budget_row.get("gasto_hoje"), budget_row.get("hoje_atualizado"), time.time(), budget_row.get("esgotado_em"))
        tacos_campanhas_html += (
            "<tr class='kc-campaign-row' tabindex='0' role='link' data-href='"
            + url_segura + "' aria-label='Abrir campanha " + esc(nome,quote=True) + "'>"
            "<td data-label='Campanha'><a class='kc-campaign-name' href='" + url_segura
            + "'>" + esc(nome) + "</a>" + budget_html + "</td>"
            "<td data-label='Investimento'>" + brl(investimento) + "</td>"
            "<td data-label='Faturamento Ads'>" + brl(atribuida) + "</td>"
            "<td data-label='ROAS'>" + (ratio(roas_camp, "x") if roas_camp is not None else "—") + "</td>"
            "<td data-label='ACOS'>" + (ratio(acos_camp, "%") if acos_camp is not None else "—") + "</td>"
            "<td data-label='TACOS'>" + (ratio(tacos_valor, "%") if tacos_valor is not None else "—")
            + "</td></tr>")
    if not linhas_tacos:
        tacos_campanhas_html += "<tr><td colspan='6'>Nenhuma campanha atende aos filtros.</td></tr>"
    tacos_campanhas_html += ("</tbody></table></div><nav aria-label='Paginação de campanhas'>"
        + "Exibindo " + str(len(linhas_tacos_pagina)) + " de " + str(len(linhas_tacos))
        + " campanhas · Página " + str(pagina_tacos) + " de " + str(paginas_tacos))
    if pagina_tacos > 1:
        tacos_campanhas_html += (" · <a href='" + esc(url_tacos(pagina_tacos - 1),quote=True)
            + "'>Anterior</a>")
    if pagina_tacos < paginas_tacos:
        tacos_campanhas_html += (" · <a href='" + esc(url_tacos(pagina_tacos + 1),quote=True)
            + "'>Próxima</a>")
    tacos_campanhas_html += """<style>
    .kc-campaigns .kc-table-wrap{width:100%;overflow-x:auto}
    .kc-campaigns .kc-campaign-table{width:100%;border-collapse:collapse}
    .kc-campaigns .kc-campaign-table th{font-size:11px;letter-spacing:.03em;white-space:nowrap;padding:13px 12px;background:#f5f7fc}
    .kc-campaigns .kc-campaign-table td{padding:17px 12px;border-bottom:1px solid #e8edf5;vertical-align:middle;font-variant-numeric:tabular-nums;font-size:12px}
    .kc-campaigns .kc-campaign-table .kc-campaign-row{cursor:pointer;transition:background .15s ease,box-shadow .15s ease}
    .kc-campaigns .kc-campaign-row:hover{background:#f0f5ff;box-shadow:inset 3px 0 #5269e7}
    .kc-campaigns .kc-campaign-row:focus-visible{outline:3px solid #6265dc;outline-offset:-3px;background:#f0f5ff}
    .kc-campaigns a.kc-campaign-name,.kc-campaigns a.kc-campaign-name:visited{display:block;color:#253657!important;text-decoration:none!important;font-weight:750;line-height:1.45;font-size:12px;overflow-wrap:anywhere}
    .kc-campaigns .kc-campaign-row:hover a.kc-campaign-name{color:#415fdb!important}
    .kc-campaigns .kc-campaign-row a.kc-campaign-name:focus-visible{outline:2px solid #5966dc;outline-offset:3px;border-radius:3px}
    @media(max-width:760px){
      .kc-campaigns .kc-table-wrap{overflow:visible}
      .kc-campaigns .kc-campaign-table,.kc-campaigns .kc-campaign-table tbody{display:block;width:100%}
      .kc-campaigns .kc-campaign-table thead{position:absolute;width:1px;height:1px;overflow:hidden;clip:rect(0 0 0 0);white-space:nowrap}
      .kc-campaigns .kc-campaign-row{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:12px;padding:15px;border:1px solid #e1e8f3;border-radius:13px;margin:0 0 11px;background:#fff;box-shadow:0 3px 12px #2436550a}
      .kc-campaigns .kc-campaign-table td{display:block;padding:0;border:0;font-size:13px;min-width:0;overflow-wrap:anywhere}
      .kc-campaigns .kc-campaign-table td::before{content:attr(data-label);display:block;margin-bottom:5px;color:#7c8aa2;font-size:10px;font-weight:750;letter-spacing:.02em}
      .kc-campaigns .kc-campaign-table td:first-child{grid-column:1/-1;border-bottom:1px solid #ebeff6;padding-bottom:11px}
      .kc-campaigns .kc-campaign-table td:first-child::before{display:none}
      .kc-campaigns a.kc-campaign-name,.kc-campaigns a.kc-campaign-name:visited{font-size:14px!important}
    }
    @media(prefers-reduced-motion:reduce){.kc-campaigns .kc-campaign-row{transition:none}}
    </style><script>
    (function(){var table=document.querySelector('.kc-campaigns .kc-campaign-table');if(!table)return;
      table.addEventListener('click',function(e){var row=e.target.closest('tr[data-href]');
        if(!row||e.target.closest('a,button,input,select'))return;window.location.assign(row.dataset.href);
      });
      table.addEventListener('keydown',function(e){var row=e.target.closest('tr[data-href]');
        if(!row||e.target!==row||!(e.key==='Enter'||e.key===' '))return;
        e.preventDefault();window.location.assign(row.dataset.href);
      });
    })();
    </script>"""
    tacos_campanhas_html += (
        "</nav><p class='muted'>ROAS = faturamento atribuído ÷ investimento; "
        "ACOS = investimento ÷ faturamento atribuído; "
        "TACOS = investimento da campanha ÷ faturamento total da loja.</p></section>")
    def card(t, val, nota=""):
        return ("<div class='stat'><div class='num'>" + esc(str(val))
                + "</div><div class='lab'>" + esc(t) + "</div>"
                + ("<small class='muted'>" + esc(nota) + "</small>" if nota else "")
                + "</div>")
    corpo = (abertura + filtro
             + "<div class='ka-kpis'>"
             + card("Investimento", brl(gasto))
             + card("Faturamento atribuído", brl(receita))
             + card("ROAS", ratio(roas, "x") if roas is not None else "—")
             + card("ACOS", ratio(acos, "%") if acos is not None else "—")
             + card("TACOS geral", ratio(tacos_geral, "%") if tacos_geral is not None else "—", "Ads ÷ faturamento total da loja")
             + card("Impressões", f"{impressoes:,}".replace(",", "."))
             + card("Cliques", f"{cliques:,}".replace(",", "."))
             + card("CTR", ratio(ctr, "%") if ctr is not None else "—")
             + card("CPC médio", brl(cpc) if cpc is not None else "—")
             + "</div>")
    corpo += alertas_html + tacos_campanhas_html
    cobertura = (fim - inicio).days + 1
    if len(dias_ads) < cobertura:
        corpo += ("<div class='aviso'>Cobertura parcial: " + str(len(dias_ads)) + " de "
                  + str(cobertura) + " dias com registros de Ads. "
                  "Os totais abaixo representam somente os dias registrados.</div>")
    corpo += daily_chart(dias_ads)
    try:
        por_pagina = int(request.args.get("por_pagina", 20))
    except (ValueError, TypeError):
        por_pagina = 20
    if por_pagina not in (10,20,50):
        por_pagina = 20
    try:
        pag = max(1, int(request.args.get("pagina", 1)))
    except (ValueError, TypeError):
        pag = 1
    total = len(dias_ads)
    paginas = max(1, (total + por_pagina - 1) // por_pagina)
    pag = min(pag, paginas)
    def link(numero):
        return "/ads/" + str(cid) + "?" + urlencode({
            "dias": dias, "de": de, "ate": ate,
            "por_pagina": por_pagina, "pagina": numero})
    bloco = dias_ads[(pag - 1) * por_pagina: pag * por_pagina]
    corpo += ("<section class='card'><h3>Detalhamento diário</h3>"
              "<form method='get' action='/ads/" + str(cid) + "'>"
              + "".join("<input type='hidden' name='" + k + "' value='" + esc(str(v)) + "'>"
                        for k, v in (("dias",dias),("de",de),("ate",ate)))
              + "<label>Linhas por página <select name='por_pagina'>"
              + "".join("<option value='" + str(n) + "'" + (" selected" if n == por_pagina else "")
                        + ">" + str(n) + "</option>" for n in (10,20,50))
              + "</select></label> <button class='btn' type='submit'>Aplicar</button></form>"
              "<div style='overflow-x:auto'><table><thead><tr><th>Dia</th>"
              "<th>Investimento</th><th>Faturamento Ads</th><th>ROAS</th>"
              "<th>Impressões</th><th>Cliques</th></tr></thead><tbody>")
    for row in bloco:
        g = float(row["gasto"] or 0)
        f = float(row["faturamento"] or 0)
        corpo += ("<tr><td>" + datetime.date.fromisoformat(str(row["data"])).strftime("%d/%m") + "</td><td>" + brl(g)
                  + "</td><td>" + brl(f) + "</td><td>"
                  + (ratio(f/g, "x") if g > 0 else "—") + "</td><td>"
                  + str(int(row["impressoes"] or 0)) + "</td><td>"
                  + str(int(row["cliques"] or 0)) + "</td></tr>")
    corpo += "</tbody></table></div><nav aria-label='Paginação de Ads'>"
    corpo += "Página " + str(pag) + " de " + str(paginas) + " · " + str(total) + " dias"
    if pag > 1:
        corpo += " · <a href='" + esc(link(pag - 1)) + "'>Anterior</a>"
    if pag < paginas:
        corpo += " · <a href='" + esc(link(pag + 1)) + "'>Próxima</a>"
    corpo += ("</nav><p class='muted'>Indicadores de Product Ads importados da API. "
              "Vendas atribuídas não representam o faturamento total da loja. "
              "ROAS = vendas atribuídas ÷ investimento; ACOS = investimento ÷ vendas atribuídas.</p>"
              "<p><a class='link' href='#ka-campaign-list'>Escolher uma campanha para abrir sua análise →</a></p></section>")
    return pagina_loja("Ads", cid, lojas[0]["cliente"], "ads", corpo + "</div>")


@app.route("/campanhas/<int:cid>")
@login_required
def campanhas(cid):
    """Compatibility link: Publicidade is the only campaign list."""
    if not pode_ver_loja(cid):
        return pagina("Acesso negado", "<div class='card'>Acesso restrito</div>"), 403
    from urllib.parse import urlencode
    from sales_metrics import reporting_today
    params = request.args.to_dict(flat=True)
    raw = params.pop("sem_vendas", "0")
    try:
        days = int(raw)
    except (TypeError, ValueError):
        days = 0
    if 0 < days < 90:
        today = reporting_today()
        params.update(situacao="ativas_sem_vendas", de=(today-datetime.timedelta(days=days)).isoformat(), ate=today.isoformat())
    return redirect("/ads/" + str(cid) + ("?"+urlencode(params) if params else ""))


@app.route("/promocoes/meta/<int:cid>",methods=["POST"])
@login_required
def salvar_meta_promocoes(cid):
    if not pode_ver_loja(cid):
        return 'Acesso restrito',403
    from client_onboarding import validate_tax
    try:
        meta=validate_tax(request.form.get('mc_meta'))
    except ValueError:
        return redirect('/promocoes/'+str(cid)+'?meta_erro=1')
    with banco() as conn:
        executar(conn,"UPDATE conexoes SET mc_meta_pct=%s,mc_meta_atualizada_em=%s WHERE id=%s",
                 (meta,int(time.time()),cid))
    return redirect('/promocoes/'+str(cid)+'?meta_ok=1')


def carregar_deal_e_anuncio(cid,pid,item_id=None):
    """Tenant-checked local campaign and product; no cross-shop item injection."""
    from promo_participation import safe_promo,safe_item
    safe_promo(pid)
    if item_id is not None:
        safe_item(item_id)
    with banco() as conn:
        campaign=consultar(conn,"SELECT * FROM promocoes WHERE conexao_id=%s AND promocao_id=%s AND tipo='DEAL'",(cid,pid))
        shop=consultar(conn,"SELECT * FROM conexoes WHERE id=%s",(cid,))
        item=(consultar(conn,"SELECT * FROM anuncios WHERE conexao_id=%s AND item_id=%s",(cid,item_id))
              if item_id else [])
    if not campaign or not shop or (item_id and not item):
        raise ValueError('Campanha ou anúncio não encontrado nesta loja.')
    return campaign[0],shop[0],item[0] if item else None


@app.route('/promocoes/candidatos/<int:cid>/<pid>')
@login_required
def promocoes_candidatos(cid,pid):
    if not pode_ver_loja(cid):
        return 'Acesso restrito',403
    from promo_participation import candidate_table
    try:
        campaign,shop,_=carregar_deal_e_anuncio(cid,pid)
    except ValueError:
        return 'Campanha indisponível',404
    token=token_atual(shop)
    if not token:
        return pagina_loja('Promoções',cid,shop['cliente'],'promocoes',
                           '<div class="aviso">Reconecte a conta do Mercado Livre para consultar candidatos.</div>'),503
    try:
        offset=max(0,min(int(request.args.get('offset','0')),10000))
    except ValueError:
        offset=0
    data,error=api_get(token,'/seller-promotions/promotions/'+pid+'/items',
                       {'promotion_type':'DEAL','app_version':'v2','limit':50,'offset':offset})
    if error or not isinstance(data,dict) or not isinstance(data.get('results'),list):
        return pagina_loja('Promoções',cid,shop['cliente'],'promocoes',
                           '<div class="aviso">Não foi possível consultar os candidatos agora.</div>'),503
    ids=[str(x.get('id')) for x in data['results'] if isinstance(x,dict) and x.get('id')]
    with banco() as conn:
        names=consultar(conn,'SELECT item_id,titulo FROM anuncios WHERE conexao_id=%s AND item_id=ANY(%s)',(cid,ids))
    html='<p><a class="link" href="/promocoes/'+str(cid)+'">Voltar às promoções</a></p>'
    html+=candidate_table(cid,pid,data['results'],{a['item_id']:a['titulo'] for a in names})
    paging=data.get('paging') or {}
    if len(data['results'])==50 or offset+len(data['results'])<int(paging.get('total') or 0):
        html+='<p><a class="filtro-btn" href="/promocoes/candidatos/'+str(cid)+'/'+pid+'?offset='+str(offset+50)+'">Próximos 50 candidatos</a></p>'
    return pagina_loja('Candidatos',cid,shop['cliente'],'promocoes',html)


def validar_adesao_deal(cid,pid,item_id,raw_price):
    """Re-read live ML candidate and financial configuration on every POST."""
    from promo_participation import parse_price,find_deal,check_candidate,margin_pct
    price=parse_price(raw_price)
    campaign,shop,item=carregar_deal_e_anuncio(cid,pid,item_id)
    if str(item.get('status') or '').lower()!='active':
        raise ValueError('Anúncio não está ativo.')
    if item.get('custo') is None or shop.get('mc_meta_pct') is None:
        raise ValueError('Cadastre o custo do produto e a meta de MC antes de participar.')
    if shop.get('aliquota_imposto') is None or shop.get('comissao_pct') is None:
        raise ValueError('Configure imposto e comissão antes de participar.')
    if session.get('tipo')=='cliente':
        with banco() as conn:
            finance=consultar(conn,"SELECT imposto_confirmado_em,custos_importados_em FROM cliente_onboarding_financeiro WHERE usuario_id=%s AND conexao_id=%s",(session['usuario_id'],cid))
        if not finance or not finance[0].get('imposto_confirmado_em') or not finance[0].get('custos_importados_em'):
            raise ValueError('Conclua a configuração financeira antes de participar.')
    token=token_atual(shop)
    if not token:
        raise ValueError('Autorização do Mercado Livre indisponível.')
    rows,error=api_get(token,'/seller-promotions/items/'+item_id,{'app_version':'v2'})
    if error:
        raise ValueError('Não foi possível validar o convite atualizado no Mercado Livre.')
    deal=find_deal(rows,pid)
    check_candidate(deal,price)
    mc=margin_pct(price,item['custo'],shop['aliquota_imposto'],shop['comissao_pct'])
    if mc < float(shop['mc_meta_pct']):
        raise ValueError('A margem estimada fica abaixo da meta definida para a loja.')
    return token,price,mc,shop,deal


@app.route('/promocoes/participar/<int:cid>/<pid>/<item_id>',methods=['GET','POST'])
@login_required
def participar_deal(cid,pid,item_id):
    if not pode_ver_loja(cid):
        return 'Acesso restrito',403
    from promo_participation import confirmation_form,find_deal,is_participating
    try:
        campaign,shop,item=carregar_deal_e_anuncio(cid,pid,item_id)
    except ValueError:
        return 'Campanha ou anúncio indisponível',404
    with banco() as conn:
        previous=consultar(conn,"SELECT status FROM promocoes_adesoes WHERE conexao_id=%s AND promocao_id=%s AND item_id=%s",(cid,pid,item_id))
    if previous:
        from promo_participation import find_deal,is_participating
        state=previous[0]['status']
        # Read-only reconciliation; never resend a request with an uncertain outcome.
        token=token_atual(shop)
        if token:
            rows,error=api_get(token,'/seller-promotions/items/'+item_id,{'app_version':'v2'})
            if not error and is_participating(find_deal(rows,pid)):
                state='confirmed'
                if previous[0]['status']!='confirmed':
                    with banco() as conn:
                        executar(conn,"""UPDATE promocoes_adesoes SET status='confirmed',
                            resposta=%s,atualizado_em=%s WHERE conexao_id=%s
                            AND promocao_id=%s AND item_id=%s""",
                            ('Participação verificada novamente no Mercado Livre.',
                             int(time.time()),cid,pid,item_id))
        label={'confirmed':'Confirmada no Mercado Livre','rejected':'Recusada pelo Mercado Livre',
               'uncertain':'Aguardando confirmação','sending':'Envio em andamento'}.get(state,'Em análise')
        return pagina_loja('Participação',cid,shop['cliente'],'promocoes',
            '<div class="card"><h2>'+esc(label)+'</h2><p>Campanha: '+esc(pid)+
            ' · Anúncio: '+esc(item_id)+'</p><p>Verificamos novamente o Mercado Livre. '
            'Para evitar adesões duplicadas, esta solicitação não será reenviada.</p>'
            '<p><a class="link" href="/promocoes/'+str(cid)+'">Voltar às promoções</a></p></div>')
    if request.method=='GET':
        token=token_atual(shop)
        if not token:
            return pagina_loja('Participação',cid,shop['cliente'],'promocoes',
                '<div class="aviso">Conta do Mercado Livre sem autorização válida.</div>'),503
        rows,error=api_get(token,'/seller-promotions/items/'+item_id,{'app_version':'v2'})
        if error:
            return pagina_loja('Participação',cid,shop['cliente'],'promocoes',
                '<div class="aviso">Não foi possível validar o convite no Mercado Livre.</div>'),503
        deal=find_deal(rows,pid)
        if not deal or str(deal.get('status','')).lower()!='candidate':
            return pagina_loja('Participação',cid,shop['cliente'],'promocoes',
                '<div class="aviso">O anúncio não é candidato a esta campanha.</div>'),409
        suggested=deal.get('suggested_discounted_price')
        default=esc(str(request.args.get('preco') or suggested or ''),quote=True)
        content=('<h2>Analisar participação</h2><p class="sub">'+esc(str(item.get('titulo') or item_id))+'</p>'
                 '<form method="get"><label>Preço promocional (R$)</label> '
                 '<input type="number" name="preco" min="0.01" step="0.01" required class="filtro" value="'+default+'"> '
                 '<button type="submit" class="filtro-btn">Calcular MC</button></form>')
        if request.args.get('preco'):
            try:
                _,price,mc,_,_=validar_adesao_deal(cid,pid,item_id,request.args['preco'])
                content+=confirmation_form(cid,pid,item_id,price,mc,shop['mc_meta_pct'],deal.get('original_price') or item['preco'] or 0)
            except ValueError as e:
                content+='<div class="aviso" role="alert">'+esc(str(e))+'</div>'
        return pagina_loja('Analisar participação',cid,shop['cliente'],'promocoes',content)
    if request.form.get('confirmar')!='sim':
        return 'Confirmação explícita obrigatória',400
    try:
        token,price,mc,shop,deal=validar_adesao_deal(cid,pid,item_id,request.form.get('deal_price'))
    except ValueError as e:
        return pagina_loja('Participação não enviada',cid,shop['cliente'],'promocoes',
                          '<div class="aviso" role="alert">'+esc(str(e))+'</div>'),409
    now=int(time.time())
    with banco() as conn:
        cursor=executar(conn,"""INSERT INTO promocoes_adesoes
            (conexao_id,promocao_id,item_id,usuario_id,preco,margem_pct,status,criado_em,atualizado_em)
            VALUES (%s,%s,%s,%s,%s,%s,'sending',%s,%s)
            ON CONFLICT DO NOTHING""",(cid,pid,item_id,session['usuario_id'],price,mc,now,now))
        reserved=cursor.rowcount==1
    if not reserved:
        return 'Já existe uma solicitação para esta campanha e anúncio.',409
    # The user authorized a real API write. Never retry automatically after timeout.
    state='uncertain'
    detail='Envio realizado; confirmação ainda não verificada.'
    try:
        response=requests.post(API+'/seller-promotions/items/'+item_id,
            params={'app_version':'v2'},
            headers={'Authorization':'Bearer '+token,'Content-Type':'application/json'},
            json={'promotion_id':pid,'promotion_type':'DEAL','deal_price':float(price)},timeout=25)
        if response.status_code not in (200,201,202,204):
            state='rejected'
            detail='Mercado Livre recusou a solicitação (HTTP '+str(response.status_code)+').'
        else:
            fresh,error=api_get(token,'/seller-promotions/items/'+item_id,{'app_version':'v2'})
            if not error and is_participating(find_deal(fresh,pid)):
                state='confirmed'
                detail='Participação confirmada pelo Mercado Livre.'
    except requests.RequestException:
        detail='O envio não pôde ser confirmado. Verifique no Mercado Livre antes de tentar novamente.'
    with banco() as conn:
        executar(conn,"""UPDATE promocoes_adesoes SET status=%s,resposta=%s,atualizado_em=%s
            WHERE conexao_id=%s AND promocao_id=%s AND item_id=%s""",
            (state,detail,int(time.time()),cid,pid,item_id))
    message=('<div class="sucesso">' if state=='confirmed' else '<div class="aviso">')+esc(detail)+'</div>'
    return pagina_loja('Resultado da participação',cid,shop['cliente'],'promocoes',
        message+'<p><a class="link" href="/promocoes/'+str(cid)+'">Voltar às promoções</a></p>')


@app.route("/promocoes/<int:cid>")
@login_required
def promocoes(cid):
    if not pode_ver_loja(cid):
        return pagina("Acesso negado",
                      "<div class='card'><h2>Acesso restrito</h2></div>")
    hoje = datetime.date.today()
    with banco() as conn:
        c = consultar(conn, "SELECT * FROM conexoes WHERE id=%s", (cid,))
        c = c[0] if c else None
        if c:
            promos = consultar(conn, "SELECT * FROM promocoes "
                                     "WHERE conexao_id=%s ORDER BY fim DESC", (cid,))
            historico_adesoes = consultar(conn, """
                SELECT h.*,a.titulo FROM promocoes_adesoes h
                LEFT JOIN anuncios a ON a.conexao_id=h.conexao_id AND a.item_id=h.item_id
                WHERE h.conexao_id=%s ORDER BY h.criado_em DESC LIMIT 100
            """, (cid,))
            itens_promo = consultar(conn, """
                SELECT pi.promocao_id, pi.item_id, pi.preco_promocional,
                       a.titulo, a.preco, a.custo, a.quantidade, a.status
                FROM promocoes_itens pi
                LEFT JOIN anuncios a
                  ON a.conexao_id=pi.conexao_id AND a.item_id=pi.item_id
                WHERE pi.conexao_id=%s
            """, (cid,))
    if not c:
        return pagina("Não encontrada",
                      "<div class='card'><h2>Loja não encontrada</h2></div>")
    ultima = ultima_atualizacao(cid)
    itens_por_promo = {}
    for it in itens_promo:
        itens_por_promo.setdefault(it["promocao_id"], []).append(it)
    meta=c.get("mc_meta_pct")
    meta_txt=("%.2f" % meta) if meta is not None else ""
    corpo = "<h2 style='margin-bottom:10px'>Promoções</h2>" \
        + barra_atualizar(cid, ultima, "promocoes")
    corpo += ("""<div class='card' style='border:1px solid #d6e5f4;border-left:5px solid #2464a0'>
        <h2 style='margin-bottom:7px'>Meta de Margem de Contribuição (MC)</h2>
        <p class='sub'>Defina a MC mínima desejada para avaliar cada promoção.
        Verde: dentro da meta. Vermelho: abaixo da meta.
        Para campanhas tradicionais DEAL, você pode enviar a participação manualmente após confirmar o preço e a MC. Nenhuma adesão é automática.</p>
        <form method='post' action='/promocoes/meta/""" + str(cid) + """'
        style='display:flex;flex-wrap:wrap;align-items:center;gap:10px;margin-top:12px'>
        <label for='mc_meta' style='font-weight:700'>MC mínima desejada (%)</label>
        <input id='mc_meta' name='mc_meta' type='number' step='0.01' min='0' max='100'
        required class='filtro' style='width:110px' value='""" + meta_txt + """'>
        <button type='submit' class='filtro-btn'>Salvar meta</button></form></div>""")
    from promo_history import render_history
    corpo += render_history(historico_adesoes,cid)
    if request.args.get("meta_erro"):
        corpo += "<div class='aviso'>Informe uma meta de MC entre 0% e 100%.</div>"
    if request.args.get("meta_ok"):
        corpo += "<div class='sucesso'>Meta de MC atualizada.</div>"
    if not promos:
        corpo += ("<div class='card' style='text-align:center;padding:40px'>"
                  "<h2>Nenhuma promoção carregada ainda</h2>"
                  "<div class='sub'>Use 'Atualizar agora' acima para buscar "
                  "as promoções da loja</div></div>")
        return pagina_loja("Promoções", cid, c["cliente"], "promocoes", corpo)
    ativas = sum(1 for p in promos if str(p["status"]) in ("active", "ACTIVE"))
    candidatas = sum(1 for p in promos if str(p["status"]) in ("candidate", "CANDIDATE"))
    finalizadas = len(promos) - ativas - candidatas
    corpo += ("<div class='grid' style='margin-bottom:18px'>"
              "<div class='stat'><div class='num'>" + str(ativas)
              + "</div><div class='lab'>Promoções ativas</div></div>"
              "<div class='stat'><div class='num'>" + str(candidatas)
              + "</div><div class='lab'>Candidatas</div></div>"
              "<div class='stat'><div class='num'>" + str(finalizadas)
              + "</div><div class='lab'>Finalizadas</div></div>"
              "</div>")
    alertas = []
    for p in promos:
        if str(p["status"]) not in ("active", "ACTIVE"):
            continue
        fim = p["fim"]
        if not fim:
            continue
        try:
            d_fim = datetime.datetime.fromisoformat(fim[:10]).date()
        except Exception:
            continue
        dias = (d_fim - hoje).days
        if dias <= 3:
            alertas.append((p["nome"] or p["promocao_id"], dias, "urgente"))
        elif dias <= 7:
            alertas.append((p["nome"] or p["promocao_id"], dias, "atencao"))
    if alertas:
        corpo += ("<div class='card'><h2>⏰ Alertas de expiração</h2>"
                  "<div class='sub'>Promoções ativas que estão terminando</div>")
        for nome, dias, grav in alertas:
            if dias <= 0:
                txt = "expira hoje"
            else:
                txt = "expira em " + str(dias) + " dia(s)"
            corpo += ("<div class='alerta " + grav
                      + "' style='margin-bottom:6px'>" + esc(nome) + " — " + txt + "</div>")
        corpo += "</div>"
    else:
        corpo += ("<div class='card'><h2>⏰ Alertas de expiração</h2>"
                  "<div class='sub'>Nenhuma promoção ativa perto de expirar</div>"
                  "<span class='alerta ok'>OK</span></div>")
    for p in promos:
        st = str(p["status"])
        if st in ("active", "ACTIVE"):
            classe, st_txt = "ativa", "ATIVA"
        elif st in ("candidate", "CANDIDATE"):
            classe, st_txt = "candidata", "CANDIDATA"
        else:
            classe, st_txt = "finalizada", "FINALIZADA"
        itens = itens_por_promo.get(p["promocao_id"], [])
        corpo += ("<div class='promo-item " + classe + "'>"
                  "<div class='titulo'>"
                  + esc(p["nome"] or ("Promoção " + str(p["promocao_id"])))
                  + " <span class='badge'>" + st_txt + "</span></div>"
                  "<div class='meta'>Tipo: " + esc(str(p["tipo"] or "-"))
                  + " · Início: " + esc(str(p["inicio"] or "-")[:10])
                  + " · Fim: " + esc(str(p["fim"] or "-")[:10])
                  + " · Itens: " + str(len(itens)) + "</div>")
        if p.get('tipo')=='DEAL':
            corpo += ('<p style="margin:10px 0"><a class="filtro-btn" href="/promocoes/candidatos/'
                      +str(cid)+'/'+esc(str(p['promocao_id']))+'">Ver candidatos e participar</a></p>')
        if itens:
            corpo += ("<table><thead><tr><th>Anúncio</th><th>Preço normal</th>"
                      "<th>Preço promo</th><th>Desconto</th><th>Custo</th>"
                      "<th>Contribuição estimada</th></tr></thead><tbody>")
            for it in itens:
                preco = it["preco"] or 0
                preco_promo = it["preco_promocional"]
                custo = it["custo"]
                if preco and preco_promo is not None:
                    desc = round((preco - preco_promo) * 100.0 / preco, 1)
                    desc_html = "<span class='alerta atencao'>-" + str(desc) + "%</span>"
                else:
                    desc_html = "<span class='tag sem'>—</span>"
                if custo is not None and preco_promo is not None:
                    if preco_promo:
                        imposto = float(c.get("aliquota_imposto") or 0) / 100
                        comissao = float(c.get("comissao_pct") or 0) / 100
                        margem = round((preco_promo - custo - preco_promo * (imposto + comissao)) * 100.0 / preco_promo, 1)
                    else:
                        margem = 0
                    if meta is None:
                        margem_html = ("<span class='alerta atencao'>%.1f%% · defina a meta</span>" % margem)
                    elif margem < float(meta):
                        margem_html = ("<span class='alerta urgente'>%.1f%% · abaixo da meta</span>" % margem)
                    else:
                        margem_html = ("<span class='alerta ok'>%.1f%% · dentro da meta</span>" % margem)
                elif custo is None:
                    margem_html = ("<a class='link' href='/custos/"
                                   + str(cid) + "'>+ custo</a>")
                else:
                    margem_html = "<span class='tag sem'>sem preço promo</span>"
                if custo is not None:
                    custo_html = "R$ %.2f" % custo
                else:
                    custo_html = ("<a class='link' href='/custos/"
                                  + str(cid) + "'>+ definir</a>")
                if preco_promo is not None:
                    preco_promo_html = "R$ %.2f" % preco_promo
                else:
                    preco_promo_html = "<span class='tag sem'>—</span>"
                corpo += ("<tr><td>"
                          + esc(str(it["titulo"] or it["item_id"]))
                          + "<br><small class='muted'>" + str(it["item_id"])
                          + "</small></td>"
                          "<td>R$ %.2f" % preco + "</td><td>"
                          + preco_promo_html + "</td><td>" + desc_html
                          + "</td><td>" + custo_html + "</td><td>"
                          + margem_html + "</td></tr>")
            corpo += "</tbody></table>"
        else:
            corpo += "<div class='muted'>Sem itens associados a esta promoção.</div>"
        corpo += "</div>"
    corpo += ("<div class='muted'>Para o comparativo de margem, cadastre o "
              "<b>custo</b> dos anúncios na <a class='link' href='/custos/"
              + str(cid) + "'>tela de Custos</a> (individual ou por planilha). "
              "O painel calcula uma contribuição estimada descontando custo, "
              "imposto e comissão configurada. Frete, tarifa fixa, publicidade "
              "e outros custos não estão incluídos; confira a rentabilidade "
              "real antes de aderir a qualquer promoção.</div>")
    return pagina_loja("Promoções", cid, c["cliente"], "promocoes", corpo)


@app.route("/exportar_demandas/<int:cid>")
@login_required
def exportar_demandas(cid):
    if not pode_ver_loja(cid):
        return "Acesso negado", 403
    with banco() as conn:
        c = consultar(conn, "SELECT * FROM conexoes WHERE id=%s", (cid,))
        c = c[0] if c else None
        if c:
            anuncios = consultar(conn, "SELECT * FROM anuncios WHERE conexao_id=%s", (cid,))
            metricas = consultar(conn, "SELECT item_id, visitas, conversao "
                                       "FROM metricas WHERE conexao_id=%s", (cid,))
            promos_ativas = consultar(conn, """SELECT DISTINCT pi.item_id
                FROM promocoes_itens pi
                JOIN promocoes p ON p.conexao_id=pi.conexao_id
                  AND p.promocao_id=pi.promocao_id
                WHERE pi.conexao_id=%s AND p.status IN ('active','ACTIVE')""", (cid,))
    if not c:
        return "Loja não encontrada", 404
    itens_promocao = {r["item_id"] for r in promos_ativas}
    visitas_por_item = {}  # API indisponível; não reutilizar zeros antigos
    conv_por_item = {}  # conversão sem visitas verificadas não é confiável
    linhas = []
    for a in anuncios:
        a["visitas"] = None
        demandas_lista = montar_demandas_anuncio(
            a, a["item_id"] in itens_promocao, conv_por_item.get(a["item_id"]))
        for acao, grav in demandas_lista:
            linhas.append([c["cliente"], a["item_id"], a["titulo"], grav, acao])
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Cliente", "Item", "Anúncio", "Gravidade", "Demanda"])
    for l in linhas:
        writer.writerow(l)
    nome_arq = "demandas_" + str(c["cliente"]).replace(" ", "_") + ".csv"
    return Response(buf.getvalue(), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": "attachment; filename=" + nome_arq})


@app.route("/demandas/<int:cid>")
@login_required
def demandas(cid):
    if not pode_ver_loja(cid):
        return pagina("Acesso negado",
                      "<div class='card'><h2>Acesso restrito</h2></div>")
    limite = request.args.get("limite", type=int, default=20)
    if limite < 1:
        limite = 20
    if limite > 200:
        limite = 200
    filtro = request.args.get("filtro", "").strip()
    with banco() as conn:
        c = consultar(conn, "SELECT * FROM conexoes WHERE id=%s", (cid,))
        c = c[0] if c else None
        if c:
            anuncios = consultar(conn, "SELECT * FROM anuncios WHERE conexao_id=%s", (cid,))
            metricas = consultar(conn, "SELECT item_id, visitas, conversao "
                                       "FROM metricas WHERE conexao_id=%s", (cid,))
            promos_ativas = consultar(conn, """SELECT DISTINCT pi.item_id
                FROM promocoes_itens pi
                JOIN promocoes p ON p.conexao_id=pi.conexao_id
                  AND p.promocao_id=pi.promocao_id
                WHERE pi.conexao_id=%s AND p.status IN ('active','ACTIVE')""", (cid,))
    if not c:
        return pagina("Não encontrada",
                      "<div class='card'><h2>Loja não encontrada</h2></div>")
    ultima = ultima_atualizacao(cid)
    itens_promocao = {r["item_id"] for r in promos_ativas}
    visitas_por_item = {}  # API indisponível; não reutilizar zeros antigos
    conv_por_item = {}  # conversão sem visitas verificadas não é confiável
    itens_demandas = []
    for a in anuncios:
        a["visitas"] = None
        demandas_lista = montar_demandas_anuncio(
            a, a["item_id"] in itens_promocao, conv_por_item.get(a["item_id"]))
        if not demandas_lista:
            continue
        urgentes = sum(1 for _, g in demandas_lista if g == "urgente")
        itens_demandas.append({
            "item_id": a["item_id"],
            "titulo": a["titulo"],
            "demandas": demandas_lista,
            "urgentes": urgentes,
            "total": len(demandas_lista)
        })
    itens_demandas.sort(key=lambda x: (-x["urgentes"], -x["total"]))
    corpo = ("<div class='barra-topo'><h2>Central de Demandas</h2>"
             "<div style='display:flex;gap:10px;flex-wrap:wrap;align-items:center'>"
             "<a class='btn-acoes cinza' href='/exportar_demandas/"
             + str(cid) + "'>⬇ Exportar CSV</a>"
             "<a class='btn-acoes cinza' href='#' onclick='window.print(); "
             "return false;'>🖨 Imprimir/PDF</a>"
             "</div></div>"
             + barra_atualizar(cid, ultima, "demandas"))
    total_anuncios = len(anuncios)
    com_demanda = len(itens_demandas)
    urgentes_total = sum(i["urgentes"] for i in itens_demandas)
    corpo += ("<div class='grid' style='margin-bottom:18px'>"
              "<div class='stat'><div class='num'>" + str(total_anuncios)
              + "</div><div class='lab'>Total de anúncios</div></div>"
              "<div class='stat'><div class='num'>" + str(com_demanda)
              + "</div><div class='lab'>Precisam de ação</div></div>"
              "<div class='stat'><div class='num'>" + str(urgentes_total)
              + "</div><div class='lab'>Demandas urgentes</div></div>"
              "</div>")
    corpo += ("<div class='card no-print'><h2>Filtros</h2>"
              "<div class='sub'>Principais pontos de alerta por anúncio, "
              "ordenados por prioridade</div>"
              "<form method='get' action='/demandas/" + str(cid)
              + "' style='display:flex;flex-wrap:wrap;align-items:center;gap:8px'>"
              "<label style='font-size:13px;color:#6b7280'>Mostrar:</label>"
              "<input type='number' name='limite' value='" + str(limite)
              + "' min='1' max='200' class='filtro' style='width:90px'>"
              "<span style='color:#6b7280;font-size:13px'>anúncios</span>"
              "<button class='filtro-btn'>Aplicar</button>"
              "</form>"
              "<div style='margin-top:10px'><b style='font-size:13px;"
              "color:#6b7280'>Filtrar por tipo:</b> "
              "<a class='filtro' href='/demandas/" + str(cid) + "?limite="
              + str(limite) + "' style='text-decoration:none'>Todos</a>"
              "<a class='filtro' href='/demandas/" + str(cid) + "?limite="
              + str(limite) + "&filtro=estoque' style='text-decoration:none'>Estoque</a>"
              "<a class='filtro' href='/demandas/" + str(cid) + "?limite="
              + str(limite) + "&filtro=fotos' style='text-decoration:none'>Fotos</a>"
              "<a class='filtro' href='/demandas/" + str(cid) + "?limite="
              + str(limite) + "&filtro=clips' style='text-decoration:none'>Clips</a>"
              "<a class='filtro' href='/demandas/" + str(cid) + "?limite="
              + str(limite) + "&filtro=tag' style='text-decoration:none'>Tags</a>"
              "<a class='filtro' href='/demandas/" + str(cid) + "?limite="
              + str(limite) + "&filtro=conversao' style='text-decoration:none'>Conversão</a>"
              "</div></div>")
    if not itens_demandas:
        corpo += ("<div class='card' style='text-align:center;padding:40px'>"
                  "<h2>Nenhuma demanda encontrada</h2>"
                  "<div class='sub'>Todos os anúncios carregados estão em dia. "
                  "Use 'Atualizar agora' acima para buscar os dados mais recentes."
                  "</div></div>")
        return pagina_loja("Demandas", cid, c["cliente"], "demandas", corpo)
    if filtro:
        filtrados = []
        for i in itens_demandas:
            textos = [d.lower() for d, _ in i["demandas"]]
            if filtro == "estoque" and any("estoque" in t for t in textos):
                filtrados.append(i)
            elif filtro == "fotos" and any("foto" in t for t in textos):
                filtrados.append(i)
            elif filtro == "clips" and any("clip" in t or "vídeo" in t
                                           or "video" in t for t in textos):
                filtrados.append(i)
            elif filtro == "tag" and any("tag" in t for t in textos):
                filtrados.append(i)
            elif filtro == "conversao" and any("conversão" in t for t in textos):
                filtrados.append(i)
        itens_demandas = filtrados
        if not itens_demandas:
            corpo += ("<div class='card' style='text-align:center;padding:30px'>"
                      "<h2>Nenhum anúncio com esse tipo de demanda</h2></div>")
            return pagina_loja("Demandas", cid, c["cliente"], "demandas", corpo)
    itens_demandas = itens_demandas[:limite]
    corpo += ("<div class='card'><h2>Anúncios que precisam de ação "
              "<span class='badge'>" + str(len(itens_demandas)) + "</span></h2>")
    for i in itens_demandas:
        if i["urgentes"] > 0:
            classe = "urgente"
        else:
            classe = "atencao"
        corpo += ("<div class='demanda-item " + classe + "'>"
                  "<div class='titulo'>" + esc(i["titulo"])
                  + " <small class='muted'>(" + str(i["item_id"]) + ")</small>"
                  " <span class='badge'>" + str(i["total"]) + " demandas</span></div>")
        for acao, grav in i["demandas"]:
            if grav == "urgente":
                icone = "🔴 "
            else:
                icone = "🟡 "
            corpo += ("<div class='acao " + grav + "'>" + icone + esc(acao) + "</div>")
        corpo += "</div>"
    corpo += "</div>"
    return pagina_loja("Demandas", cid, c["cliente"], "demandas", corpo)
@app.route("/diagnostico/<int:cid>/<item_id>")
@login_required
def diagnostico_saude_anuncio(cid, item_id):
    if not pode_ver_loja(cid):
        return pagina("Acesso negado", "<div class='card'><h2>Acesso restrito</h2></div>"), 403
    with banco() as conn:
        lojas = consultar(conn, "SELECT * FROM conexoes WHERE id=%s", (cid,))
        anuncios_encontrados = consultar(
            conn, "SELECT * FROM anuncios WHERE conexao_id=%s AND item_id=%s",
            (cid, item_id))
        pedidos_verificados = consultar(conn, """
            SELECT conexao_id, pedido_id, status, total, moeda, criado_em,
                   pagamentos_json, unidades, vendas_dia_coleta
            FROM pedidos WHERE conexao_id=%s AND vendas_verificado_em IS NOT NULL
        """, (cid,))
        itens_pedidos = consultar(conn, """
            SELECT pi.conexao_id, pi.pedido_id, pi.item_id, pi.quantidade,
                   pi.preco_unitario
            FROM pedidos_itens pi JOIN pedidos p
              ON p.conexao_id=pi.conexao_id AND p.pedido_id=pi.pedido_id
            WHERE pi.conexao_id=%s AND p.vendas_verificado_em IS NOT NULL
        """, (cid,))
        cobertura = consultar(conn,
            "SELECT dia FROM vendas_cobertura WHERE conexao_id=%s", (cid,))
        pendentes = consultar(conn, """
            SELECT COUNT(*) AS n FROM pedidos
            WHERE conexao_id=%s AND vendas_verificado_em IS NULL
        """, (cid,))
    if not lojas:
        return pagina("Loja não encontrada", "<div class='card'>Loja não encontrada</div>"), 404
    if not anuncios_encontrados:
        return pagina_loja("Diagnóstico", cid, lojas[0]["cliente"], "anuncios",
                          "<div class='card'>Anúncio não encontrado nesta loja.</div>"), 404
    item = anuncios_encontrados[0]
    resultados = diagnose_listing(item)
    vendas_7 = listing_sales_summary(
        pedidos_verificados, itens_pedidos, [r["dia"] for r in cobertura],
        item_id, days=7)
    vendas_30 = listing_sales_summary(
        pedidos_verificados, itens_pedidos, [r["dia"] for r in cobertura],
        item_id, days=30)
    if pendentes and pendentes[0]["n"]:
        for vendas in (vendas_7, vendas_30):
            vendas.update(orders=None, units=None, gross_items=None, ticket=None,
                          reason="Há pedidos antigos ainda não verificados nesta loja.")
    from listing_diagnostic_design import hero, priorities
    cabecalho = hero(item, cid, item_id, resultados)
    if session.get('tipo') == 'admin':
        from urllib.parse import quote
        cabecalho += "<p><a class='btn-acoes' href='/anuncios/"+str(int(cid))+"/"+quote(item_id,safe='')+"/tarefas'>Planejar tarefas do anúncio</a></p>"
    corpo = cabecalho
    corpo += ("<div class='card'><h3>Indicadores de vendas por anúncio</h3>"
              "<p class='muted'>Somente pedidos pagos e aprovados, com cobertura "
              "diária e vínculo completo dos itens. Valor bruto dos itens = "
              "preço unitário registrado na venda × quantidade. "
              "Não equivale ao valor líquido recebido, nem inclui rateio de "
              "frete, cupons ou taxas.</p>"
              "<div style='overflow-x:auto'><table><thead><tr>"
              "<th>Período</th><th>Pedidos</th><th>Unidades</th>"
              "<th>Valor dos itens</th><th>Ticket médio por pedido</th>"
              "</tr></thead><tbody>")
    for periodo, vendas in (("7 dias", vendas_7), ("30 dias", vendas_30)):
        if vendas["orders"] is None:
            corpo += ("<tr><td>" + periodo + "</td><td colspan='4'>"
                      "<b>Indisponível</b> — "
                      + esc(vendas["reason"] or "Sem dados validados.")
                      + "</td></tr>")
        else:
            valor = "R$ %.2f" % vendas["gross_items"]
            ticket = ("R$ %.2f" % vendas["ticket"]
                      if vendas["ticket"] is not None else "—")
            corpo += ("<tr><td>" + periodo + "</td><td>"
                      + str(vendas["orders"]) + "</td><td>"
                      + str(vendas["units"]) + "</td><td>"
                      + valor + "</td><td>" + ticket + "</td></tr>")
    corpo += ("</tbody></table></div><p class='muted'>"
              "Ticket médio = valor bruto dos itens deste anúncio ÷ "
              "pedidos que contêm este anúncio. Quando não há pedidos, "
              "o ticket não é calculado. A cobertura de visitas e a relação "
              "pedidos/visitas são detalhadas na investigação acima.</p></div>")
    # Compare complete, non-overlapping periods using the same paid-order
    # verification gates as the headline indicators.
    corpo += ("<div class='card'><h3>Evolução das vendas por anúncio</h3>"
              "<p class='muted'>Comparação com o período imediatamente anterior, "
              "sem sobreposição de dias. Valores em reais são brutos dos itens "
              "vendidos, não repasses líquidos.</p>"
              "<div style='overflow-x:auto'><table><thead><tr>"
              "<th>Comparação</th><th>Pedidos</th><th>Unidades</th>"
              "<th>Valor dos itens</th></tr></thead><tbody>")
    for dias, atual in ((7, vendas_7), (30, vendas_30)):
        fim_anterior = datetime.datetime.fromisoformat(
            atual["start"]).date() - datetime.timedelta(days=1)
        anterior = listing_sales_summary(
            pedidos_verificados, itens_pedidos, [r["dia"] for r in cobertura],
            item_id, end=fim_anterior, days=dias)
        if pendentes and pendentes[0]["n"]:
            anterior.update(orders=None, units=None, gross_items=None,
                            ticket=None, reason="Há pedidos antigos não verificados.")
        evolucao = comparison(atual, anterior)
        if not evolucao["available"]:
            corpo += ("<tr><td>" + str(dias) + " dias vs. anteriores</td>"
                      "<td colspan='3'><b>Indisponível</b> — "
                      + esc(evolucao["reason"]) + "</td></tr>")
            continue
        celulas = []
        for campo in ("orders", "units", "gross_items"):
            valor = evolucao[campo]
            if valor["status"] == "novas_vendas":
                rotulo = "Novas vendas (base anterior zero)"
            elif valor["status"] == "sem_movimento":
                rotulo = "Sem movimento nos dois períodos"
            else:
                rotulo = ("%+.1f%%" % valor["percent"])
            if campo == "gross_items":
                atual_txt = "R$ %.2f" % valor["current"]
                anterior_txt = "R$ %.2f" % valor["previous"]
            else:
                atual_txt = str(int(valor["current"]))
                anterior_txt = str(int(valor["previous"]))
            celulas.append("<td><b>" + atual_txt + "</b><br>"
                           "<small class='muted'>Antes: " + anterior_txt
                           + " · " + rotulo + "</small></td>")
        corpo += ("<tr><td>" + str(dias) + " dias vs. anteriores</td>"
                  + "".join(celulas) + "</tr>")
    corpo += ("</tbody></table></div><p class='muted'>"
              "Sem cobertura verificada em qualquer período, a comparação "
              "fica indisponível. Quando a base anterior é zero, não exibimos "
              "crescimento percentual artificial.</p></div>")
    # Observations are snapshots: unknown dates never inherit a guessed state.
    with banco() as conn:
        full_rows = consultar(conn, """SELECT observed_at,state,quantity,logistics
            FROM full_observations
            WHERE conexao_id=%s AND item_id=%s
              AND observed_at >= %s
            ORDER BY observed_at ASC""",
            (cid,item_id,int(time.time()) - 32*86400))
    from sales_decline import render as render_decline
    from sales_metrics import reporting_today
    fim_investigacao = reporting_today() - datetime.timedelta(days=1)
    ids_investigacao = sorted({str(i["item_id"]) for i in itens_pedidos if i.get("item_id")} | {str(item_id)})
    investigacao = ranking_metrics(
        pedidos_verificados, itens_pedidos, [r["dia"] for r in cobertura],
        ids_investigacao, end=fim_investigacao, days=7,
        pending=bool(pendentes and pendentes[0]["n"]))
    corpo += render_decline(investigacao, [item], cid, {}, fim_investigacao, days=7, item_id=item_id, observations=full_rows)
    from investigation_evidence import load as load_evidence, build as build_evidence, render as render_evidence
    with banco() as conn:
        visitas_evidencia, ads_evidencia = load_evidence(
            conn, consultar, cid, item_id, fim_investigacao)
    from offer_history import load as load_offers, render as render_offers
    with banco() as conn:
        ofertas = load_offers(conn,consultar,cid,item_id,fim_investigacao)
    evidencia = build_evidence(
        investigacao, item_id, visitas_evidencia, full_rows, ads_evidencia,
        fim_investigacao, offers=ofertas['own'])
    painel_evidencias = render_evidence(evidencia, cid, fim_investigacao)
    # Reuse stored manual and automatically detected changes; one bounded read.
    # No API calls, no guessing when a change actually happened.
    from automatic_history import HISTORY
    from investigation_change_context import read as read_change_context, render as render_change_context
    boundary = fim_investigacao - datetime.timedelta(days=6)
    with banco() as conn:
        alteracoes = consultar(conn, """
            SELECT event_date, kind, note, created_at, source, field
            FROM """ + HISTORY + """ WHERE conexao_id=%s AND item_id=%s
              AND void_at IS NULL AND event_date BETWEEN %s AND %s
            ORDER BY event_date DESC,created_at DESC LIMIT 100
        """, (cid,item_id,(boundary-datetime.timedelta(days=3)).isoformat(),
              (boundary+datetime.timedelta(days=3)).isoformat()))
    contexto = read_change_context(alteracoes, fim_investigacao,
                                  sales_verified=evidencia['sales_verified'],
                                  declining=evidencia['declining'])
    painel_evidencias += render_change_context(contexto)
    painel_ofertas = render_offers(ofertas,cid,item_id,fim_investigacao,args=request.args)
    from competitive_sales_context import render as render_competitive_sales
    painel_ofertas += render_competitive_sales(investigacao,item_id,ofertas,fim_investigacao)
    from listing_timeline import render as render_timeline
    linha_tempo = render_timeline(item_id,pedidos_verificados,itens_pedidos,
        [r["dia"] for r in cobertura],full_rows,visitas_evidencia,ofertas['own'],
        fim_investigacao,pending=bool(pendentes and pendentes[0]["n"]),args=request.args)
    corpo = cabecalho + painel_evidencias + linha_tempo + painel_ofertas + corpo[len(cabecalho):]
    corpo += priorities(resultados)
    return pagina_loja("Diagnóstico do anúncio", cid, lojas[0]["cliente"], "anuncios", corpo)


@app.route("/anuncios/<int:cid>")
@login_required
def anuncios(cid):
    if not pode_ver_loja(cid):
        return pagina("Acesso negado",
                      "<div class='card'><h2>Acesso restrito</h2></div>")
    with banco() as conn:
        c = consultar(conn, "SELECT * FROM conexoes WHERE id=%s", (cid,))
        c = c[0] if c else None
        anuncios = consultar(conn, "SELECT * FROM anuncios "
                                   "WHERE conexao_id=%s ORDER BY vendidos DESC NULLS LAST, item_id", (cid,))
        erros = consultar(conn, "SELECT * FROM erros WHERE conexao_id=%s AND resolvido_em IS NULL", (cid,))
    if not c:
        return pagina("Não encontrada",
                      "<div class='card'><h2>Loja não encontrada</h2></div>")
    ultima = ultima_atualizacao(cid)
    from listing_catalog import render_catalog
    corpo = render_catalog(cid, anuncios, request.args,
                           refresh=barra_atualizar(cid, ultima, "anuncios"),
                           errors=erros, admin=session.get("tipo") == "admin",
                           freight=frete_esperado)
    return pagina_loja("Anúncios", cid, c["cliente"], "anuncios", corpo)


@app.route("/ranking-anuncios/<int:cid>")
@login_required
def ranking_anuncios(cid):
    """Store-level verified listing ranking; never use lifetime vendidos."""
    if not pode_ver_loja(cid):
        return pagina("Acesso negado", "<div class='card'>Acesso restrito</div>"), 403
    dias = request.args.get("dias", "7")
    if dias not in ("7", "30"):
        dias = "30"
    from sales_metrics import reporting_today
    fim_analise = reporting_today() - datetime.timedelta(days=1)
    inicio_sql, fim_sql, inicio_cobertura, fim_cobertura = ranking_query_bounds(int(dias), today=fim_analise)
    sinais_financeiros = {}
    visitas_comparacao = {}
    with banco() as conn:
        executar(conn, 'SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY')
        lojas = consultar(conn, "SELECT * FROM conexoes WHERE id=%s", (cid,))
        if not lojas:
            return pagina("Não encontrada", "<div class='card'>Loja não encontrada</div>"), 404
        anuncios_loja = consultar(conn, """
            SELECT item_id, titulo, status, quantidade, atualizado_em FROM anuncios
            WHERE conexao_id=%s ORDER BY item_id
        """, (cid,))
        pedidos = consultar(conn, """
            SELECT conexao_id, pedido_id, status, total, moeda, criado_em,
                   pagamentos_json, unidades, vendas_dia_coleta
            FROM pedidos
            WHERE conexao_id=%s AND vendas_verificado_em IS NOT NULL
              AND (LEFT(criado_em,10) BETWEEN %s AND %s
                   OR vendas_dia_coleta BETWEEN %s AND %s)
        """, (cid, inicio_sql, fim_sql, inicio_cobertura, fim_cobertura))
        itens = consultar(conn, """
            SELECT pi.conexao_id, pi.pedido_id, pi.item_id,
                   pi.quantidade, pi.preco_unitario
            FROM pedidos_itens pi JOIN pedidos p
              ON p.conexao_id=pi.conexao_id AND p.pedido_id=pi.pedido_id
            WHERE pi.conexao_id=%s AND p.vendas_verificado_em IS NOT NULL
              AND (LEFT(p.criado_em,10) BETWEEN %s AND %s
                   OR p.vendas_dia_coleta BETWEEN %s AND %s)
        """, (cid, inicio_sql, fim_sql, inicio_cobertura, fim_cobertura))
        cobertura = consultar(conn, """
            SELECT dia FROM vendas_cobertura
            WHERE conexao_id=%s AND dia BETWEEN %s AND %s
        """, (cid, inicio_cobertura, fim_cobertura))
        pendentes = consultar(conn, """
            SELECT EXISTS (
                SELECT 1 FROM pedidos WHERE conexao_id=%s
                AND vendas_verificado_em IS NULL
            ) AS n
        """, (cid,))
        resumo = ranking_metrics(
            pedidos, itens, [r["dia"] for r in cobertura],
            sorted({str(a["item_id"]) for a in anuncios_loja} | {str(i["item_id"]) for i in itens if i.get("item_id")}),
            end=fim_analise, days=int(dias), pending=bool(pendentes and pendentes[0]["n"]))
        if request.args.get('modo') != 'prioridades':
            from listing_comparison import selection as comparison_selection, load_visits as load_comparison_visits
            selected_comparison=comparison_selection(resumo,anuncios_loja,request.args)
            if selected_comparison['available']:
                visitas_comparacao=load_comparison_visits(conn,consultar,cid,selected_comparison['ids'],fim_analise,days=int(dias))
        if resumo['available'] and session.get('tipo') == 'admin':
            from profitability_priority import load_signals as load_financial_signals
            sinais_financeiros = load_financial_signals(
                conn,consultar,cid,resumo,pedidos,itens,fim_analise,int(dias))
    corpo = ("<h2>Ranking de vendas por anúncio</h2>"
             "<p class='muted'>Últimos " + dias + " dias encerrados comparados aos " + dias + " anteriores. "
             "Somente pedidos pagos e aprovados com dados verificados. "
             "Valor bruto dos itens, não valor líquido recebido. "
             "Sinais de tráfego e eficiência dependem da cobertura diária do anúncio.</p>")
    corpo += ("<div class='card'><p>Período: "
              "<a class='filtro' href='/ranking-anuncios/" + str(cid)
              + "?dias=7'>7 dias</a> "
              "<a class='filtro' href='/ranking-anuncios/" + str(cid)
              + "?dias=30'>30 dias</a></p></div>")
    if request.args.get("modo") != "prioridades":
        from listing_comparison import render as render_listing_comparison
        corpo += render_listing_comparison(cid,resumo,anuncios_loja,request.args,fim_analise,days=int(dias),visits=visitas_comparacao)
        from listing_abc import render as render_listing_abc
        corpo += render_listing_abc(cid,resumo,anuncios_loja,request.args,fim_analise,days=int(dias))
    from listing_priority_queue import prioritize as prioritize_listings, render as render_priorities
    compact_priorities = request.args.get("modo") == "prioridades"
    if not resumo["available"] and compact_priorities:
        return render_priorities(prioritize_listings(resumo), anuncios_loja, cid, days=int(dias), compact=True)
    if not resumo["available"]:
        corpo += ("<div class='card'><b>Ranking indisponível</b><p>"
                  + esc(resumo["reason"] or "Dados não verificados.")
                  + "</p><p>Não interpretamos dados incompletos como vendas zeradas."
                  "</p></div>")
        return pagina_loja("Ranking de anúncios", cid, lojas[0]["cliente"],
                          "desempenho", corpo)
    from sales_decline import render as render_decline
    from investigation_evidence import load_bulk
    with banco() as conn:
        evidencias_queda = load_bulk(conn, consultar, cid, resumo, fim_analise, days=int(dias))
    from listing_priority_sources import listing_signals
    sinais_verificados = listing_signals(anuncios_loja, resumo, int(dias))
    for item_id, financial in sinais_financeiros.items():
        sinais_verificados.setdefault(item_id, {}).update(financial)
    prioridades = render_priorities(prioritize_listings(
        resumo, evidencias_queda, verified_signals_by_item=sinais_verificados),
        anuncios_loja, cid, days=int(dias), compact=compact_priorities)
    if compact_priorities:
        response = Response(prioridades, mimetype="text/html")
        response.headers["Cache-Control"] = "private, max-age=60"
        return response
    corpo += prioridades
    if session.get('tipo') == 'admin':
        from financial_priority_queue import render as render_financial_queue
        from profitability_view import CSS as profitability_css
        corpo += profitability_css + render_financial_queue(
            sinais_financeiros, anuncios_loja, cid,
            fim_analise-datetime.timedelta(days=int(dias)-1), fim_analise)
    corpo += render_decline(resumo, anuncios_loja, cid, request.args, fim_analise,
                            days=int(dias), evidence_by_item=evidencias_queda)
    kpis = listing_kpis(resumo["rows"])
    valor = "R$ %.2f" % kpis["value"]
    medio = ("R$ %.2f" % kpis["average_value_per_listing_order"]
             if kpis["average_value_per_listing_order"] is not None else "—")
    concentracao = ("%.1f%%" % kpis["top_listing_share_pct"]
                    if kpis["top_listing_share_pct"] is not None else "—")
    corpo += ("<p class='muted'>Evolução de anúncios com base de vendas anterior: "
              + str(kpis["rising"]) + " em alta, "
              + str(kpis["falling"]) + " em queda e "
              + str(kpis["newly_selling"]) + " com novas vendas.</p>")
    corpo += ("<div class='grid' style='margin-bottom:18px'>"
              "<div class='stat'><div class='num'>" + valor
              + "</div><div class='lab'>Valor bruto dos itens</div></div>"
              "<div class='stat'><div class='num'>" + str(kpis["units"])
              + "</div><div class='lab'>Unidades vendidas</div></div>"
              "<div class='stat'><div class='num'>" + str(kpis["selling_listings"])
              + " / " + str(kpis["listed"])
              + "</div><div class='lab'>Anúncios com vendas</div></div>"
              "<div class='stat'><div class='num'>" + medio
              + "</div><div class='lab'>Valor médio por pedido/anúncio</div></div>"
              "<div class='stat'><div class='num'>" + concentracao
              + "</div><div class='lab'>Participação do anúncio líder</div></div>"
              "</div>")
    sinais = listing_opportunities(resumo["rows"])
    if sinais:
        nomes = {str(a["item_id"]): str(a["titulo"] or a["item_id"])
                 for a in anuncios_loja}
        corpo += ("<div class='card'><h3>Oportunidades para investigar</h3>"
                  "<p class='muted'>Alertas baseados apenas em pedidos verificados. "
                  "Não comprovam queda de visitas, conversão ou demanda.</p><ul>")
        for sinal in sinais:
            corpo += ("<li><a href='/diagnostico/" + str(cid) + "/"
                      + esc(sinal["item_id"]) + "'>"
                      + esc(nomes.get(sinal["item_id"], sinal["item_id"]))
                      + "</a>: " + esc(sinal["detail"]) + " ("
                      + str(sinal["previous_orders"]) + " → "
                      + str(sinal["orders"]) + " pedidos)."
                      + "<details><summary>O que verificar</summary><ul>"
                      + "".join("<li>" + esc(passo) + "</li>"
                                for passo in opportunity_actions(sinal["kind"]))
                      + "</ul></details></li>")
        corpo += "</ul></div>"
    opcoes = {"valor": "Maior valor vendido", "pedidos": "Mais pedidos",
              "alta": "Maior crescimento", "queda": "Maior queda",
              "participacao": "Maior participação"}
    ordem = request.args.get("ordem", "valor")
    if ordem not in opcoes:
        ordem = "valor"
    filtros = {"todos": "Todos", "alta": "Em alta", "queda": "Em queda",
               "novas": "Novas vendas", "sem_vendas": "Sem vendas"}
    filtro = request.args.get("filtro", "todos")
    if filtro not in filtros:
        filtro = "todos"
    registros = [(a, resumo["rows"][str(a["item_id"])]) for a in anuncios_loja]
    if filtro == "alta":
        registros = [(a, m) for a, m in registros
                     if m["previous_orders"] > 0 and m["orders"] > m["previous_orders"]]
    elif filtro == "queda":
        registros = [(a, m) for a, m in registros
                     if m["previous_orders"] > 0 and m["orders"] < m["previous_orders"]]
    elif filtro == "novas":
        registros = [(a, m) for a, m in registros
                     if m["orders"] > 0 and m["previous_orders"] == 0]
    elif filtro == "sem_vendas":
        registros = [(a, m) for a, m in registros if m["orders"] == 0]
    if ordem == "valor":
        registros.sort(key=lambda x: (-x[1]["value"], str(x[0]["item_id"])))
    elif ordem == "pedidos":
        registros.sort(key=lambda x: (-x[1]["orders"], str(x[0]["item_id"])))
    elif ordem == "alta":
        registros.sort(key=lambda x: (
            x[1]["growth_pct"] is None, -(x[1]["growth_pct"] or 0),
            str(x[0]["item_id"])))
    elif ordem == "queda":
        registros.sort(key=lambda x: (
            x[1]["growth_pct"] is None, x[1]["growth_pct"] or 0,
            str(x[0]["item_id"])))
    else:
        registros.sort(key=lambda x: (-x[1]["value"], str(x[0]["item_id"])))
    corpo += "<div class='card'><p>Ordenar por: "
    for chave, titulo in opcoes.items():
        corpo += ("<a class='filtro' href='/ranking-anuncios/" + str(cid)
                  + "?dias=" + dias + "&filtro=" + filtro + "&ordem=" + chave + "'>" + titulo + "</a> ")
    corpo += "</p><p>Filtrar por: "
    for chave, titulo in filtros.items():
        corpo += ("<a class='filtro' href='/ranking-anuncios/" + str(cid)
                  + "?dias=" + dias + "&ordem=" + ordem + "&filtro=" + chave
                  + "'>" + titulo + "</a> ")
    corpo += ("</p><p class='muted'>Exibindo até 100 de "
              + str(len(registros)) + " anúncios. "
              "Crescimento percentual requer vendas no período anterior; "
              "anúncios com base zero não recebem porcentagem artificial."
              "</p><div style='overflow-x:auto'><table><thead><tr>"
              "<th>Anúncio</th><th>Pedidos</th><th>Unidades</th>"
              "<th>Valor dos itens</th><th>Participação</th><th>Ticket médio</th>"
              "<th>Pedidos anteriores</th><th>Evolução</th>"
              "</tr></thead><tbody>")
    for anuncio, m in registros[:100]:
        if m["growth_pct"] is None:
            evolucao = ("Novas vendas" if m["trend"] == "novas"
                        else "Sem base percentual")
        else:
            evolucao = "%+.1f%%" % m["growth_pct"]
        share = ("%.1f%%" % (m["value"] * 100 / kpis["value"])
                 if kpis["value"] else "—")
        if m["orders"] and not m["previous_orders"]:
            destaque = "<span class='alerta ok'>Novas vendas</span>"
        elif m["growth_pct"] is not None and m["growth_pct"] < 0:
            destaque = "<span class='alerta atencao'>Em queda</span>"
        elif m["growth_pct"] is not None and m["growth_pct"] > 0:
            destaque = "<span class='alerta ok'>Em alta</span>"
        elif m["value"] == max((r["value"] for _, r in registros), default=0) and m["value"] > 0:
            destaque = "<span class='badge'>Líder em vendas</span>"
        else:
            destaque = ""
        ticket = "R$ %.2f" % m["ticket"] if m["ticket"] is not None else "—"
        corpo += ("<tr><td><b>" + esc(str(anuncio["titulo"] or "Sem título"))
                  + "</b><br><small>" + esc(str(anuncio["item_id"]))
                  + "</small> " + destaque + "<br><a href='/diagnostico/" + str(cid) + "/"
                  + esc(str(anuncio["item_id"])) + "'>Ver diagnóstico</a></td>"
                  "<td>" + str(m["orders"]) + "</td><td>" + str(m["units"])
                  + "</td><td>R$ %.2f" % m["value"]
                  + "</td><td>" + share + "</td><td>" + ticket + "</td><td>"
                  + str(m["previous_orders"]) + "</td><td>"
                  + esc(evolucao) + "</td></tr>")
    corpo += "</tbody></table></div></div>"
    return pagina_loja("Ranking de anúncios", cid, lojas[0]["cliente"],
                      "desempenho", corpo)


@app.route("/desempenho/<int:cid>")
@login_required
def desempenho(cid):
    """Sem visitas da API, não apresentar zeros ou conversões artificiais."""
    if not pode_ver_loja(cid):
        return pagina("Acesso negado",
                      "<div class='card'><h2>Acesso restrito</h2></div>"), 403
    hoje = datetime.date.today()
    dias = request.args.get("dias", type=int, default=30)
    dias = dias if dias in (7, 30, 90) else 30
    de = request.args.get("de", "").strip()
    ate = request.args.get("ate", "").strip()
    try:
        d_inicio = datetime.date.fromisoformat(de) if de and ate else hoje - datetime.timedelta(days=dias - 1)
        d_fim = datetime.date.fromisoformat(ate) if de and ate else hoje
    except ValueError:
        d_inicio = hoje - datetime.timedelta(days=29)
        d_fim = hoje
    if d_fim < d_inicio:
        d_inicio, d_fim = d_fim, d_inicio
    if d_fim > hoje:
        d_fim = hoje
    if (d_fim - d_inicio).days > 365:
        d_inicio = d_fim - datetime.timedelta(days=365)
    ini, fim = d_inicio.isoformat(), d_fim.isoformat()
    with banco() as conn:
        lojas = consultar(conn, "SELECT * FROM conexoes WHERE id=%s", (cid,))
        if not lojas:
            return pagina("Não encontrada", "<div class='card'><h2>Loja não encontrada</h2></div>"), 404
        c = lojas[0]
        # O campo anuncios.vendidos é acumulado desde a criação do anúncio,
        # NÃO corresponde às vendas do período. Usar pedidos registrados.
        linhas = consultar(conn, """
            SELECT a.item_id, a.titulo, a.status, a.preco,
                   COUNT(DISTINCT CASE WHEN p.status='paid'
                       AND SUBSTRING(p.fechado_em,1,10) BETWEEN %s AND %s
                       THEN p.pedido_id END) AS pedidos
            FROM anuncios a
            LEFT JOIN pedidos_itens pi ON pi.conexao_id=a.conexao_id
                AND pi.item_id=a.item_id
            LEFT JOIN pedidos p ON p.conexao_id=pi.conexao_id
                AND p.pedido_id=pi.pedido_id
            WHERE a.conexao_id=%s
            GROUP BY a.item_id, a.titulo, a.status, a.preco
            ORDER BY pedidos DESC, a.item_id
        """, (ini, fim, cid))
        grafico_f = consultar(conn, """
            SELECT SUBSTRING(fechado_em,1,10) AS dia,
                   COALESCE(SUM(total),0) AS total
            FROM pedidos
            WHERE conexao_id=%s AND status='paid'
              AND fechado_em IS NOT NULL
              AND SUBSTRING(fechado_em,1,10) BETWEEN %s AND %s
            GROUP BY SUBSTRING(fechado_em,1,10)
        """, (cid, ini, fim))
    ultima = ultima_atualizacao(cid)
    link_ranking = ("<p><a class='btn-acoes cinza' href='/ranking-anuncios/"
                    + str(cid) + "'>Ranking de vendas por anúncio →</a></p>")
    corpo = (link_ranking + "<h2 style='margin-bottom:10px'>Desempenho</h2>"
             + barra_atualizar(cid, ultima, "desempenho")
             + "<div class='aviso'><b>Visitas e conversão indisponíveis.</b> "
             "O Mercado Livre não está fornecendo visitas confiáveis pela API. "
             "Os zeros do histórico antigo não representam ausência de tráfego. "
             "Nenhuma taxa de conversão é estimada.</div>"
             "<div class='card'><h2>Período de análise</h2>"
             "<form method='get' action='/desempenho/" + str(cid)
             + "' style='display:flex;flex-wrap:wrap;gap:8px'>"
             "<a class='filtro' href='/desempenho/" + str(cid) + "?dias=7'>7 dias</a>"
             "<a class='filtro' href='/desempenho/" + str(cid) + "?dias=30'>30 dias</a>"
             "<a class='filtro' href='/desempenho/" + str(cid) + "?dias=90'>90 dias</a>"
             "<input type='date' name='de' value='" + esc(ini) + "' class='filtro'>"
             "<input type='date' name='ate' value='" + esc(fim) + "' class='filtro'>"
             "<button class='filtro-btn'>Aplicar</button></form>"
             "<div class='muted'>Período: " + d_inicio.strftime("%d/%m/%Y")
             + " a " + d_fim.strftime("%d/%m/%Y") + "</div></div>")
    fat_por_dia = {r["dia"]: float(r["total"] or 0) for r in grafico_f}
    rotulos, fat_serie = [], []
    dia = d_inicio
    while dia <= d_fim:
        rotulos.append(dia.strftime("%d/%m"))
        fat_serie.append(round(fat_por_dia.get(dia.isoformat(), 0), 2))
        dia += datetime.timedelta(days=1)
    total_pedidos = sum(int(r["pedidos"] or 0) for r in linhas)
    total_fat = round(sum(fat_serie), 2)
    corpo += ("<div class='grid' style='margin-bottom:18px'>"
              "<div class='stat'><div class='num'>Indisponível</div><div class='lab'>Visitas</div></div>"
              "<div class='stat'><div class='num'>Indisponível</div><div class='lab'>Conversão</div></div>"
              "<div class='stat'><div class='num'>" + str(total_pedidos)
              + "</div><div class='lab'>Pedidos pagos por anúncio (base legada)</div></div>"
              "<div class='stat'><div class='num'>R$ %.2f" % total_fat
              + "</div><div class='lab'>Faturamento registrado (base legada)</div></div>"
              "</div><div class='card'><h2>Faturamento diário registrado</h2>"
              "<div class='grafico-box'><canvas id='graficoDesempenho'></canvas></div></div>"
              "<script src='https://cdn.jsdelivr.net/npm/chart.js@4.4.1/dist/chart.umd.min.js'></script>"
              "<script>new Chart(document.getElementById('graficoDesempenho'),{"
              "type:'line',data:{labels:" + json.dumps(rotulos)
              + ",datasets:[{label:'Faturamento registrado (R$)',data:"
              + json.dumps(fat_serie)
              + ",borderColor:'#3483FA',tension:0.3}]},"
              "options:{responsive:true,maintainAspectRatio:false}});</script>")
    from performance_cards import render_performance_cards
    corpo += render_performance_cards(cid, linhas, request.args, dias=dias, de=de, ate=ate)
    return pagina_loja("Desempenho", cid, c["cliente"], "desempenho", corpo)


def br(v, decimais=2):
    if v is None:
        return None
    s = format(float(v), ",." + str(decimais) + "f")
    s = s.replace(",", "X").replace(".", ",").replace("X", ".")
    return s


def var_html(hoje, ontem, fmt="R$"):
    if ontem is None or ontem == 0:
        return "<div class='var estavel'>— sem base anterior</div>"
    delta = hoje - ontem
    pct = (delta * 100.0 / ontem) if ontem else 0
    if fmt == "n":
        un = ""
        dec = 0
    else:
        un = "R$ "
        dec = 2
    if delta > 0:
        return ("<div class='var subiu'>▲ +" + un + br(delta, dec)
                + " (+" + br(pct, 1) + "%)</div>")
    elif delta < 0:
        return ("<div class='var caiu'>▼ " + un + br(abs(delta), dec)
                + " (" + br(pct, 1) + "%)</div>")
    return "<div class='var estavel'>— 0%</div>"

@app.route('/kyra')
@login_required
@admin_required
def kyra_central():
    from kyra_central import central_rows, render_central
    with banco() as conn:
        shops = consultar(conn, "SELECT id,cliente FROM conexoes WHERE status='ativa' ORDER BY cliente,id")
        runs = consultar(conn, """SELECT DISTINCT ON (conexao_id)
            conexao_id,day,status,result FROM keiper_ai_runs
            ORDER BY conexao_id,day DESC""")
    body = render_central(central_rows(shops, runs))
    response = Response(pagina('KEVIA · Central de Inteligência', body), mimetype='text/html')
    response.headers['Cache-Control'] = 'no-store'
    return response


@app.route('/kyra/revisar/<int:cid>', methods=['POST'])
@login_required
@admin_required
def kyra_revisar(cid):
    from hmac import compare_digest
    from kyra_reviews import ALLOWED
    if not pode_ver_loja(cid):
        return 'Acesso restrito', 403
    csrf = request.form.get('csrf') or ''
    if not csrf or not compare_digest(csrf, session.get('kyra_review_csrf') or ''):
        return 'Solicitação inválida', 403
    code = request.form.get('code') or ''
    state = request.form.get('state') or ''
    if not code or len(code) > 80 or state not in ALLOWED:
        return 'Dados inválidos', 400
    with banco() as conn:
        rows = consultar(conn, '''SELECT result FROM keiper_ai_runs
            WHERE conexao_id=%s ORDER BY day DESC LIMIT 1''', (cid,))
        from kevia_finance import review_findings
        findings = review_findings(rows[0].get('result')) if rows else []
        if code not in {str(f.get('code')) for f in findings if isinstance(f, dict)}:
            return 'Recomendação não encontrada', 404
        executar(conn, '''INSERT INTO kyra_reviews(conexao_id,code,state)
            VALUES(%s,%s,%s) ON CONFLICT(conexao_id,code)
            DO UPDATE SET state=EXCLUDED.state,updated_at=now()''', (cid,code,state))
    return redirect('/ia-keiper/'+str(cid), code=303)


@app.route('/ia-keiper/<int:cid>')
@login_required
def painel_ia_keiper(cid):
    if not pode_ver_loja(cid):
        return 'Acesso restrito', 403
    from contextlib import closing
    from keiper_ai import render
    from keiper_ai_budget import budget_status
    with closing(banco()) as conn, conn:
        lojas = consultar(conn, 'SELECT cliente FROM conexoes WHERE id=%s', (cid,))
        if not lojas:
            return 'Loja não encontrada', 404
        runs = consultar(conn, '''SELECT day,status,result,api_status FROM keiper_ai_runs
            WHERE conexao_id=%s ORDER BY day DESC LIMIT 30''', (cid,))
        from kevia_finance import visible_result
        runs = [dict(row,result=visible_result(row.get('result'),admin=session.get('tipo')=='admin'))
                for row in runs]
        campaign_runs = consultar(conn, '''SELECT day,status,result FROM keiper_ai_campaign_runs
            WHERE conexao_id=%s ORDER BY day DESC LIMIT 1''', (cid,))
        budget = budget_status(conn, consultar) if session.get('tipo') == 'admin' else None
        saved_reviews = consultar(conn, 'SELECT code,state FROM kyra_reviews WHERE conexao_id=%s', (cid,)) if session.get('tipo') == 'admin' else []
        result = runs[0].get('result') if runs else None
        daily = consultar(conn, '''SELECT data,gasto,faturamento FROM ads_dia
            WHERE conexao_id=%s AND data BETWEEN %s AND %s ORDER BY data''',
            (cid, result['previous_start'], result['end'])) if result else []
    run = runs[0] if runs else {}
    corpo = render(run.get('result'), run.get('day'), run.get('status'), run.get('api_status'), budget, cid=cid, daily=daily, campaign_run=campaign_runs[0] if campaign_runs else None,admin_financial=session.get('tipo')=='admin')
    if run.get('status') in ('running', 'rules_ready'):
        corpo += "<p>Execução em andamento ou interrompida. A próxima análise ocorrerá no próximo dia de serviço ativo.</p>"
    from kyra_history import render_history
    corpo += render_history(runs,admin_financial=session.get('tipo')=='admin')
    if session.get('tipo') == 'admin':
        from kyra_reviews import render_review
        token = session.get('kyra_review_csrf')
        if not token:
            token = secrets.token_urlsafe(32)
            session['kyra_review_csrf'] = token
        from kevia_finance import review_findings
        findings = review_findings(run.get('result'))
        corpo += render_review(cid, findings, {r['code']:r['state'] for r in saved_reviews}, token)
    return pagina_loja('KEVIA', cid, lojas[0]['cliente'], 'ia', corpo)


@app.route("/painel/<int:cid>/operacao")
@login_required
def painel_detalhe(cid):
    if not pode_ver_loja(cid):
        return pagina("Acesso negado",
                      "<div class='card'><h2>Acesso restrito</h2></div>")
    hoje_iso = datetime.date.today().isoformat()
    ontem_iso = (datetime.date.today()
                 - datetime.timedelta(days=1)).isoformat()
    inicio_semana = (datetime.date.today() - datetime.timedelta(days=6)).isoformat()
    c, vendas = calcular_vendas(cid, inicio=inicio_semana)
    if not c:
        return pagina("Não encontrada",
                      "<div class='card'><h2>Loja não encontrada</h2></div>")
    with banco() as conn:
        cobertura_vendas = consultar(conn, """
            SELECT dia FROM vendas_cobertura
            WHERE conexao_id=%s AND dia >= %s
        """, (cid, inicio_semana))
        dias_cobertos = {r["dia"] for r in cobertura_vendas}
        conta = consultar(conn, "SELECT * FROM dados_conta WHERE conexao_id=%s", (cid,))
        conta = conta[0] if conta else None
        ads_dias = consultar(conn, "SELECT * FROM ads_dia WHERE conexao_id=%s "
                                   "AND data >= %s ORDER BY data",
                             (cid, (datetime.date.today()
                                    - datetime.timedelta(days=6)).isoformat()))
        perg_pend = consultar(conn, "SELECT COUNT(*) AS n FROM perguntas "
                                    "WHERE conexao_id=%s AND status != 'ANSWERED'", (cid,))
        envios_hoje = consultar(conn, "SELECT COUNT(*) AS n FROM envios "
                                      "WHERE conexao_id=%s AND SUBSTRING(data,1,10)=%s",
                                (cid, hoje_iso))
        ult = consultar(conn, "SELECT MAX(atualizado_em) AS ult FROM anuncios "
                              "WHERE conexao_id=%s", (cid,))
    ultima = ult[0]["ult"] if ult and ult[0]["ult"] else None
    perguntas_pendentes = perg_pend[0]["n"] if perg_pend else 0
    envios_dia = envios_hoje[0]["n"] if envios_hoje else 0
    atualizando = False
    if (ultima is None or (int(time.time()) - ultima) > GATILHO_AUTO_SEG) \
            and c["user_id"]:
        enfileirar_atualizacao(cid)
        atualizando = cid in ATUALIZANDO
    por_dia = {}
    for v in vendas:
        dia = v["data"][:10]
        d = por_dia.setdefault(dia, {"fat": 0.0, "qtd": 0,
                                     "lucro": 0.0, "pendente": False})
        d["fat"] += v["receita"]
        d["qtd"] += 1
        if v["sem_itens"] or v["sem_custo"]:
            d["pendente"] = True
        else:
            d["lucro"] += v["mc"]
    hoje_info = por_dia.get(hoje_iso, {"fat": 0.0, "qtd": 0,
                                       "lucro": 0.0, "pendente": False})
    ontem_info = por_dia.get(ontem_iso, {"fat": 0.0, "qtd": 0,
                                         "lucro": 0.0, "pendente": False})
    fat_dia = hoje_info["fat"]
    qtd_dia = hoje_info["qtd"]
    lucro_dia = hoje_info["lucro"]
    tem_pendente_hoje = hoje_info["pendente"]
    ad_hoje = None
    ad_ontem = None
    for r in ads_dias:
        if r["data"] == hoje_iso:
            ad_hoje = r
        if r["data"] == ontem_iso:
            ad_ontem = r
    gasto_ads = ad_hoje["gasto"] if ad_hoje and ad_hoje["gasto"] is not None else None
    fat_ads = ad_hoje["faturamento"] if ad_hoje and ad_hoje["faturamento"] is not None else None
    gasto_ads_ontem = ad_ontem["gasto"] if ad_ontem and ad_ontem["gasto"] is not None else None
    fat_ads_ontem = ad_ontem["faturamento"] if ad_ontem and ad_ontem["faturamento"] is not None else None
    ult_txt = hora_br(ultima)
    corpo = "<h2 style='margin-bottom:14px'>Resumo do dia</h2>"
    corpo += ("<div style='display:flex;gap:12px;align-items:center;"
              "flex-wrap:wrap;margin-bottom:14px'>"
              "<a class='btn-acoes' href='/atualizar/resumo/"
              + str(cid) + "'>🔄 Atualizar agora</a>"
              "<span style='color:#6b7280;font-size:13px'>"
              "Última atualização: " + ult_txt + "</span></div>")
    if atualizando:
        corpo += ("<div class='aviso'>Atualização automática em andamento — "
                  "os dados de hoje e de Ads aparecem em instantes. "
                  "A página vai recarregar sozinha.</div>")
    if perguntas_pendentes > 0:
        corpo += ("<div class='aviso'>💬 Você tem <b>"
                  + str(perguntas_pendentes)
                  + " pergunta(s) pendente(s)</b> para responder nos anúncios. "
                  "Responder rápido aumenta a chance de fechar a venda.</div>")
    if lucro_dia <= 0:
        lucro_html = "<span class='alerta urgente'>R$ %.2f</span>" % lucro_dia
    else:
        lucro_html = "R$ %.2f" % lucro_dia
    gasto_ads_var = ""
    if gasto_ads is not None:
        gasto_ads_var = var_html(gasto_ads, gasto_ads_ontem)
    fat_ads_var = ""
    if fat_ads is not None:
        fat_ads_var = var_html(fat_ads, fat_ads_ontem)
    corpo += ("<div class='grid' style='margin-bottom:18px'>"
              "<div class='stat'><div class='num'>R$ %.2f" % fat_dia
              + "</div><div class='lab'>Faturamento do dia</div>"
              + var_html(fat_dia, ontem_info["fat"]) + "</div>"
              "<div class='stat'><div class='num'>" + str(qtd_dia)
              + "</div><div class='lab'>Vendas do dia</div>"
              + var_html(qtd_dia, ontem_info["qtd"], "n") + "</div>"
              "<div class='stat'><div class='num'>" + lucro_html
              + "</div><div class='lab'>Lucro do dia (MC)</div>"
              + var_html(lucro_dia, ontem_info["lucro"]) + "</div>"
              "<div class='stat'><div class='num'>" + str(envios_dia)
              + "</div><div class='lab'>Envios do dia</div></div>"
              "<div class='stat'><div class='num'>"
              + (("R$ %.2f" % gasto_ads) if gasto_ads is not None
                 else "<span class='tag sem'>—</span>")
              + "</div><div class='lab'>Gasto Ads do dia</div>"
              + gasto_ads_var + "</div>"
              "<div class='stat'><div class='num'>"
              + (("R$ %.2f" % fat_ads) if fat_ads is not None
                 else "<span class='tag sem'>—</span>")
              + "</div><div class='lab'>Faturamento Ads do dia</div>"
              + fat_ads_var + "</div>"
              "</div>")
    if tem_pendente_hoje:
        corpo += ("<div class='aviso'>Há vendas de hoje sem custo cadastrado — "
                  "o lucro do dia pode estar incompleto. "
                  "<a class='link' href='/custos/" + str(cid)
                  + "'>Cadastrar custos</a>.</div>")
    dias_incompletos = [
        (datetime.date.today() - datetime.timedelta(days=i)).isoformat()
        for i in range(7) if (datetime.date.today() - datetime.timedelta(days=i)).isoformat()
        not in dias_cobertos
    ]
    if dias_incompletos:
        corpo += ("<div class='aviso'><b>Histórico de vendas em sincronização.</b> "
                  "Os dias sem cobertura confirmada podem estar incompletos. "
                  "Um traço não significa zero vendas.</div>")
    rotulos7, fat7, lucro7 = [], [], []
    d = datetime.date.today() - datetime.timedelta(days=6)
    while d <= datetime.date.today():
        iso = d.isoformat()
        info = por_dia.get(iso, {"fat": 0.0, "qtd": 0,
                                 "lucro": 0.0, "pendente": False})
        rotulos7.append(d.strftime("%d/%m"))
        fat7.append(round(info["fat"], 2) if iso in dias_cobertos
                    else (round(info["fat"], 2) if info["qtd"] else None))
        lucro7.append(round(info["lucro"], 2) if iso in dias_cobertos
                      else (round(info["lucro"], 2) if info["qtd"] else None))
        d += datetime.timedelta(days=1)
    corpo += ("<div class='card'><h2>Evolução dos últimos 7 dias</h2>"
              "<div class='sub'>Faturamento e lucro (MC) por dia</div>"
              "<div class='grafico-box'>"
              "<canvas id='graficoResumo'></canvas></div></div>")
    corpo += ("<script src='https://cdn.jsdelivr.net/npm/"
              "chart.js@4.4.1/dist/chart.umd.min.js'></script>")
    corpo += ("<script>var r7 = " + json.dumps(rotulos7) + ";"
              "var f7 = " + json.dumps(fat7) + ";"
              "var l7 = " + json.dumps(lucro7) + ";"
              "new Chart(document.getElementById('graficoResumo'), {"
              "type:'line',"
              "data:{labels:r7,datasets:["
              "{label:'Faturamento (R$)',data:f7,borderColor:'#3483FA',"
              "backgroundColor:'rgba(52,131,250,0.12)',tension:0.3,fill:true},"
              "{label:'Lucro (MC R$)',data:l7,borderColor:'#1a9d5c',"
              "backgroundColor:'rgba(26,157,92,0.12)',tension:0.3,fill:true}]},"
              "options:{responsive:true,maintainAspectRatio:false,"
              "plugins:{legend:{position:'top'}},"
              "scales:{y:{type:'linear',beginAtZero:true,"
              "title:{display:true,text:'R$'}}}}});</script>")
    ads_map = {r["data"]: r for r in ads_dias}
    linhas_html = ""
    d = datetime.date.today() - datetime.timedelta(days=6)
    while d <= datetime.date.today():
        iso = d.isoformat()
        info = por_dia.get(iso, {"fat": 0.0, "qtd": 0,
                                 "lucro": 0.0, "pendente": False})
        ad = ads_map.get(iso)
        if ad and ad["gasto"] is not None:
            gasto_txt = "R$ %.2f" % ad["gasto"]
        else:
            gasto_txt = "<span class='tag sem'>—</span>"
        if ad and ad["faturamento"] is not None:
            fatads_txt = "R$ %.2f" % ad["faturamento"]
        else:
            fatads_txt = "<span class='tag sem'>—</span>"
        if info["pendente"]:
            lucro_txt = "<span class='tag sem'>custo pendente</span>"
        else:
            lucro_txt = "R$ %.2f" % info["lucro"]
        completo = iso in dias_cobertos
        receita_txt = ("R$ %.2f" % info["fat"]
                       if completo or info["qtd"] else "—")
        vendas_txt = (str(info["qtd"]) if completo or info["qtd"] else "—")
        if not completo and not info["qtd"]:
            lucro_txt = "<span class='tag sem'>—</span>"
        badge = ""
        if not completo:
            badge += " <span class='tag sem'>parcial</span>"
        if iso == hoje_iso:
            badge = " <span class='badge'>hoje</span>"
        linhas_html += ("<tr><td><b>" + d.strftime("%d/%m") + "</b>"
                        + badge + "</td>"
                        "<td>" + receita_txt + "</td><td>"
                        + vendas_txt + "</td><td>" + lucro_txt
                        + "</td><td>" + gasto_txt + "</td><td>"
                        + fatads_txt + "</td></tr>")
        d += datetime.timedelta(days=1)
    corpo += ("<div class='card'><h2>Últimos 7 dias</h2>"
              "<div class='sub'>Faturamento, vendas, lucro (MC) e Ads por dia</div>"
              "<div style='overflow-x:auto'><table>"
              "<thead><tr><th>Dia</th><th>Faturamento</th><th>Vendas</th>"
              "<th>Lucro (MC)</th><th>Gasto Ads</th>"
              "<th>Faturamento Ads</th></tr></thead>"
              "<tbody>" + linhas_html + "</tbody></table></div>"
              "<div class='muted'>Lucro (MC) = Receita − Custo − Imposto − "
              "Comissão − Frete, calculado por venda. Vendas sem custo "
              "cadastrado aparecem como 'custo pendente'.</div></div>")
    if conta:
        corpo += ("<div class='card'><h2>Conta</h2>"
                  "<div class='sub'>Informações do vendedor</div>"
                  "<div class='grid'>"
                  "<div class='stat'><div class='num'>"
                  + esc(str(conta["nickname"]))
                  + "</div><div class='lab'>Nickname</div></div>"
                  "<div class='stat'><div class='num'>"
                  + esc(str(conta["reputacao"]))
                  + "</div><div class='lab'>Reputação</div></div>"
                  "<div class='stat'><div class='num'>"
                  + str(conta["pontos"])
                  + "</div><div class='lab'>Vendas concluídas</div></div>"
                  "</div></div>")
    if atualizando:
        corpo += ("<script>setTimeout(function(){ location.reload(); }, "
                  "90000);</script>")
    # Load deeper priority analysis separately so the store summary stays fast.
    priority_url = "/ranking-anuncios/%s?dias=7&modo=prioridades" % int(cid)
    corpo = ("<div id='kp-auto-priorities' aria-live='polite'>"
             "<div class='card'><h3>Prioridades dos anúncios</h3>"
             "<p class='muted'>Investigando quedas de vendas verificadas…</p></div></div>"
             "<script>(function(){"
             "var node=document.getElementById('kp-auto-priorities');"
             "fetch("+json.dumps(priority_url)+",{credentials:'same-origin'})"
             ".then(function(r){if(!r.ok)throw Error('priority-load');return r.text()})"
             ".then(function(h){if(node)node.innerHTML=h})"
             ".catch(function(){if(node)node.innerHTML="
             "'<div class=\\'card\\'><h3>Prioridades indisponíveis</h3>' +"
             "'<p>Abra o ranking para consultar as quedas verificadas.</p>' +"
             "'<a href=\\'/ranking-anuncios/"+str(int(cid))+"?dias=7\\'>Ver ranking de anúncios</a></div>'})"
             "})();</script>") + corpo
    corpo = ("<div class='card'><h3>KEVIA</h3><p>Nome provisório, aguardando análise jurídica de marca.</p><p>Consulte as recomendações diárias e suas evidências.</p>"
             "<a href='/ia-keiper/"+str(int(cid))+"'>Ver recomendações</a></div>") + corpo
    return pagina_loja("Resumo", cid, c["cliente"], "resumo", corpo)
@app.route("/atualizar/<tipo>/<int:cid>")
@login_required
def atualizar_tipo(tipo, cid):
    if not pode_ver_loja(cid):
        return pagina("Acesso negado",
                      "<div class='card'><h2>Acesso restrito</h2></div>")
    from section_refresh import SCOPES
    if tipo in SCOPES:
        # Opening a legacy refresh URL only displays the explicit POST button.
        return redirect(("/ads/" if tipo == "campanhas" else "/" + tipo + "/") + str(cid))
    if tipo != "resumo":
        return pagina("Seção inválida", "<div class='card'>Seção inválida</div>"), 400
    enfileirar_atualizacao(cid)
    return redirect("/painel/" + str(cid))


@app.route("/atualizar")
@login_required
def atualizar():
    cid = request.args.get("cid", type=int)
    with banco() as conn:
        if cid:
            c = consultar(conn, "SELECT * FROM conexoes WHERE id=%s", (cid,))
        else:
            c = consultar(conn, "SELECT * FROM conexoes "
                                "WHERE status='ativa' ORDER BY id LIMIT 1")
        c = c[0] if c else None
    if not c:
        return pagina("Sem loja",
                      "<div class='card'><h2>Nenhuma loja conectada</h2>"
                      "<p><a class='link' href='/conectar'>Conectar</a></p></div>")
    if not pode_ver_loja(c["id"]):
        return pagina("Acesso negado",
                      "<div class='card'><h2>Acesso restrito</h2></div>")
    enfileirar_atualizacao(c["id"])
    if cid:
        return redirect("/painel/" + str(cid))
    return redirect("/painel")


# ---------- CONEXÃO / TOKENS ----------

def renovar(refresh_token):
    resp = requests.post(TOKEN_URL, data={
        "grant_type": "refresh_token",
        "client_id": CLIENT_ID,
        "client_secret": CLIENT_SECRET,
        "refresh_token": reveal_token(refresh_token),
    }, timeout=25)
    return resp.json()


def token_atual(c):
    agora = int(time.time())
    if c["expires_at"] and c["expires_at"] > agora + 600:
        return reveal_token(c["access_token"])
    dados = renovar(c["refresh_token"])
    if "access_token" not in dados:
        registrar_erro(c["id"], "token", "Não foi possível renovar a autorização do Mercado Livre.")
        return None
    with banco() as conn:
        executar(conn, "UPDATE conexoes SET access_token=%s, "
                       "refresh_token=%s, expires_at=%s, status='ativa' "
                       "WHERE id=%s",
                 (seal_token(dados["access_token"]), seal_token(dados["refresh_token"]),
                  agora + dados["expires_in"], c["id"]))
    return dados["access_token"]


def job_refresh():
    agora = int(time.time())
    with banco() as conn:
        conexoes = consultar(conn, "SELECT * FROM conexoes "
                                   "WHERE status='ativa' AND expires_at < %s",
                             (agora + 1800,))
        for c in conexoes:
            try:
                dados = renovar(c["refresh_token"])
                executar(conn, "UPDATE conexoes SET access_token=%s, "
                               "refresh_token=%s, expires_at=%s, "
                               "status='ativa' WHERE id=%s",
                         (seal_token(dados["access_token"]), seal_token(dados["refresh_token"]),
                          agora + dados["expires_in"], c["id"]))
            except Exception:
                executar(conn, "UPDATE conexoes SET status='expirada' "
                               "WHERE id=%s", (c["id"],))


# ---------- PUXAR DADOS ----------

def api_get(token, path, params=None):
    try:
        r = requests.get(API + path,
                         headers={"Authorization": "Bearer " + (token or "")},
                         params=params, timeout=25)
        if r.status_code == 200:
            return r.json(), None
        return None, "HTTP " + str(r.status_code) + ": " + r.text[:180]
    except Exception as e:
        return None, str(e)


def registrar_erro(conexao_id, categoria, mensagem):
    with banco() as conn:
        executar(conn, """INSERT INTO erros (conexao_id, categoria, mensagem, quando)
                        VALUES (%s,%s,%s,%s)
                        ON CONFLICT (conexao_id, categoria)
                        DO UPDATE SET mensagem=EXCLUDED.mensagem,
                          quando=EXCLUDED.quando, resolvido_em=NULL""",
                 (conexao_id, categoria, (mensagem or "")[:300],
                  int(time.time())))



def resolver_erro(cid, categoria, iniciado_em):
    # A concurrent/new error is not cleared by an older successful collection.
    with banco() as conn:
        executar(conn, """UPDATE erros SET resolvido_em=%s
            WHERE conexao_id=%s AND categoria=%s AND quando < %s
            AND resolvido_em IS NULL""", (int(time.time()),cid,categoria,iniciado_em))


def gravar_historico(conexao_id):
    hoje = datetime.date.today().isoformat()
    agora = int(time.time())
    with banco() as conn:
        linhas = consultar(conn, "SELECT item_id, vendidos FROM anuncios "
                                 "WHERE conexao_id=%s", (conexao_id,))
        metricas = consultar(conn, "SELECT item_id, visitas FROM metricas "
                                   "WHERE conexao_id=%s", (conexao_id,))
    # Visitas da API indisponíveis: NULL significa desconhecido, nunca zero.
    for l in linhas:
        item_id = l["item_id"]
        visitas = None
        vendidos = l["vendidos"] or 0
        with banco() as conn:
            executar(conn, """INSERT INTO desempenho_historico
                            (conexao_id, item_id, data, visitas, vendidos, atualizado_em)
                            VALUES (%s,%s,%s,%s,%s,%s)
                            ON CONFLICT (conexao_id, item_id, data) DO UPDATE SET
                              visitas=EXCLUDED.visitas,
                              vendidos=EXCLUDED.vendidos,
                              atualizado_em=EXCLUDED.atualizado_em""",
                     (conexao_id, item_id, hoje, visitas, vendidos, agora))


def puxar_conta(conexao_id, token):
    dados, erro = api_get(token, "/users/me")
    if erro or not dados:
        registrar_erro(conexao_id, "conta", erro or "sem resposta")
        return False
    rep = dados.get("seller_reputation") or {}
    metricas = rep.get("metrics") or {}
    sales = metricas.get("sales") or {}
    pontos = sales.get("completed") or rep.get("transactions_completed")
    with banco() as conn:
        executar(conn, """INSERT INTO dados_conta
                        (conexao_id, user_id, nickname, nome, sobrenome,
                         reputacao, pontos, atualizado_em)
                        VALUES (%s,%s,%s,%s,%s,%s,%s,%s)
                        ON CONFLICT (conexao_id) DO UPDATE SET
                          user_id=EXCLUDED.user_id,
                          nickname=EXCLUDED.nickname, nome=EXCLUDED.nome,
                          sobrenome=EXCLUDED.sobrenome,
                          reputacao=EXCLUDED.reputacao,
                          pontos=EXCLUDED.pontos,
                          atualizado_em=EXCLUDED.atualizado_em""",
                 (conexao_id, dados.get("id"), dados.get("nickname"),
                  dados.get("first_name"), dados.get("last_name"),
                  rep.get("level_id"), pontos, int(time.time())))

    return True


def puxar_anuncios(conexao_id, token, user_id):
    ids = []
    scroll = None
    while True:
        params = {"limit": 100, "search_type": "scan"}
        if scroll:
            params["scroll_id"] = scroll
        dados, erro = get_json(token, "/users/" + str(user_id) + "/items/search", params, requests.get)
        if not erro and scroll and isinstance(dados, dict) and dados.get("results") is None:
            break  # Documented scan end sentinel.
        if erro or not isinstance(dados, dict) or not isinstance(dados.get("results"), list):
            registrar_erro(conexao_id, "anuncios", erro or "busca incompleta")
            return False
        resultados = dados["results"]
        if not resultados:
            break
        if set(resultados).intersection(ids):
            registrar_erro(conexao_id, "anuncios", "paginação repetida; dados anteriores preservados")
            return False
        ids.extend(resultados)
        novo_scroll = dados.get("scroll_id")
        total = (dados.get("paging") or {}).get("total")
        if total is not None and len(ids) >= total:
            break
        if not novo_scroll:
            if len(resultados) < 100:
                break
            registrar_erro(conexao_id, "anuncios", "paginação incompleta; dados anteriores preservados")
            return False
        scroll = novo_scroll
    agora = int(time.time())
    linhas = []
    full_rows = []
    change_rows = []
    completo = True
    for i in range(0, len(ids), 20):
        lote = ids[i:i + 20]
        dados, erro = get_json(token, "/items",
                              {"ids": ",".join(lote), "include_attributes": "all"}, requests.get)
        if erro or not isinstance(dados, list):
            registrar_erro(conexao_id, "anuncios",
                           erro or "lote sem resposta")
            completo = False
            if erro and any(code in erro for code in ("HTTP 401", "HTTP 403", "HTTP 429")):
                break
            continue
        for r in dados:
            if not isinstance(r, dict):
                completo = False
                continue
            if r.get("code", 200) != 200:
                completo = False
                registrar_erro(conexao_id, "anuncios", "item indisponível HTTP " + str(r.get("code")))
                continue
            det = r.get("body")
            if not isinstance(det, dict) or not det.get("id"):
                completo = False
                continue
            from automatic_history import listing_values
            change_rows.append((det['id'],int(time.time()),listing_values(det,user_id)))
            full_rows.append((json.dumps(full_metadata(det, user_id, int(time.time()))),
                              conexao_id, det['id']))
            fotos = 0
            pics = det.get("pictures")
            if isinstance(pics, list):
                fotos = len(pics)
            vid = det.get("video_id")
            if vid is None:
                clips = None
            elif vid:
                clips = 1
            else:
                clips = 0
            tags = det.get("tags") or []
            original_price = det.get("original_price")
            pesos = registered_weights(det)
            peso = legacy_weight(pesos)
            linhas.append((conexao_id, det.get("id"), det.get("title"),
                           det.get("price"), det.get("available_quantity"),
                           det.get("status"), det.get("sold_quantity"),
                           agora, fotos, clips,
                           json.dumps(tags, ensure_ascii=False),
                           peso, det.get("seller_custom_field"), original_price,
                           json.dumps(pesos, ensure_ascii=False),
                           det.get("secure_thumbnail") or (pics[0].get("secure_url") if isinstance(pics, list) and pics and isinstance(pics[0], dict) else None) or det.get("thumbnail")))
    with banco() as conn:
        for t in linhas:
            executar(conn, """INSERT INTO anuncios
                            (conexao_id, item_id, titulo, preco, quantidade,
                             status, vendidos, atualizado_em, fotos, clips,
                             tags, peso, sku, original_price, pesos_json, thumbnail)
                            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                            ON CONFLICT (conexao_id, item_id) DO UPDATE SET
                              titulo=EXCLUDED.titulo, preco=EXCLUDED.preco,
                              quantidade=EXCLUDED.quantidade,
                              status=EXCLUDED.status,
                              vendidos=EXCLUDED.vendidos,
                              atualizado_em=EXCLUDED.atualizado_em,
                              fotos=EXCLUDED.fotos, clips=EXCLUDED.clips,
                              tags=EXCLUDED.tags, peso=EXCLUDED.peso,
                              sku=EXCLUDED.sku, original_price=EXCLUDED.original_price,
                              pesos_json=EXCLUDED.pesos_json, thumbnail=EXCLUDED.thumbnail""", t)

        from automatic_history import record as record_changes
        for changed_item,stamp,values in change_rows:
            record_changes(conn,executar,conexao_id,changed_item,'listing',stamp,values)

        for metadata in full_rows:
            executar(conn, '''UPDATE anuncios SET full_metadata_json=%s::jsonb
                WHERE conexao_id=%s AND item_id=%s''', metadata)
        velhos = consultar(conn, "SELECT item_id FROM anuncios WHERE conexao_id=%s", (conexao_id,))
        atuais = set(ids)
        for row in velhos:
            if row["item_id"] not in atuais:
                executar(conn, "DELETE FROM anuncios WHERE conexao_id=%s AND item_id=%s",
                         (conexao_id, row["item_id"]))
    return completo


def puxar_metricas(conexao_id, token):
    return

def puxar_perguntas(conexao_id, token, user_id):
    dados, erro = api_get(token, "/questions/search",
                          {"seller_id": user_id, "api_version": 4,
                           "limit": 50})
    if erro or not dados:
        registrar_erro(conexao_id, "perguntas", erro or "sem resposta")
        return False
    from listing_questions import persist
    try:
        with banco() as conn:
            persist(conn, executar, conexao_id, dados, int(time.time()))
    except ValueError as exc:
        registrar_erro(conexao_id, "perguntas", str(exc))
        return False
    return True


def puxar_envios(conexao_id, token, user_id):
    # O antigo /shipments/search foi desativado pelo ML.
    # Novo fluxo: pega o envio de cada pedido recente e consulta o detalhe.
    dados, erro = api_get(token, "/orders/search",
                          {"seller": user_id, "sort": "date_desc",
                           "limit": 50})
    if erro or not dados:
        registrar_erro(conexao_id, "envios", erro or "sem resposta")
        return False
    resultados = dados.get("results") or []
    agora = int(time.time())
    completo = True
    pares = []
    for o in resultados:
        ship = o.get("shipping") or {}
        sid = ship.get("id")
        if sid:
            par = (str(o.get("id")), str(sid))
            if par not in pares:
                pares.append(par)
        if len(pares) >= 20:
            break
    with banco() as conn:
        executar(conn, "DELETE FROM envios WHERE conexao_id=%s",
                 (conexao_id,))
    for oid, sid in pares:
        det, det_erro = api_get(token, "/shipments/" + sid)
        if det_erro or not det:
            completo = False
            registrar_erro(conexao_id, "envios", det_erro or "envio sem resposta")
            continue
        with banco() as conn:
            executar(conn, """INSERT INTO envios
                            (conexao_id, envio_id, status, tracking,
                             pedido_id, data, custo_frete, atualizado_em)
                            VALUES (%s,%s,%s,%s,%s,%s,%s,%s)
                            ON CONFLICT (conexao_id, envio_id) DO UPDATE SET
                              status=EXCLUDED.status,
                              tracking=EXCLUDED.tracking,
                              pedido_id=EXCLUDED.pedido_id,
                              data=EXCLUDED.data,
                              custo_frete=EXCLUDED.custo_frete,
                              atualizado_em=EXCLUDED.atualizado_em""",
                     (conexao_id, sid, det.get("status"),
                      det.get("tracking_number"), oid,
                      det.get("date_created"), det.get("cost_to_seller"),
                      agora))
    return completo


def puxar_promocoes(conexao_id, token, user_id):
    return sync_promotions(conexao_id, token, banco, consultar, executar,
                           requests.get, registrar_erro, user_id=user_id)


def atualizar_anuncios_bg(cid, token, user_id, tipo):
    try:
        if tipo == "anuncios":
            puxar_anuncios(cid, token, user_id)
        puxar_promocoes(cid, token, user_id)
    except Exception as exc:
        registrar_erro(cid, "anuncios_atualizacao", "Falha na atualização: " + type(exc).__name__)
    finally:
        with LOCK_ATUALIZACAO:
            ATUALIZANDO.discard(cid)


@app.route("/diagnostico/anuncios/<int:cid>/<item_id>")
@login_required
def diagnostico_anuncio(cid, item_id):
    if session.get("tipo") != "admin":
        return jsonify({"error": "Acesso restrito"}), 403
    with banco() as conn:
        contas = consultar(conn, "SELECT * FROM conexoes WHERE id=%s", (cid,))
        itens = consultar(conn, "SELECT item_id FROM anuncios WHERE conexao_id=%s AND item_id=%s", (cid, item_id))
    if not contas or not itens:
        return jsonify({"error": "Anúncio não encontrado nesta loja"}), 404
    conta = contas[0]
    token = token_atual(conta)
    if not token:
        return jsonify({"error": "Conta precisa ser reconectada"}), 401
    return jsonify(inspect_listing(token, conta["user_id"], item_id, requests.get))


def coletar_secao(cid, secao, periodo=None):
    """Run only the collectors required by the requested page."""
    with banco() as conn:
        rows = consultar(conn, "SELECT * FROM conexoes WHERE id=%s", (cid,))
    if not rows:
        return False
    conta = rows[0]
    token = token_atual(conta)
    user_id = conta.get("user_id")
    if not token or not user_id:
        return False
    from section_refresh import collector_names
    for name in collector_names(secao):
        collector = globals()[name]
        result = (collector(cid, token, user_id, refresh_range=periodo)
                  if name == "puxar_ads" else collector(cid, token, user_id))
        if result is False:
            return False
    return True


def puxar_ads(conexao_id, token, user_id, refresh_range=None):
    from sales_metrics import reporting_today
    from ads_official_metrics import save_day
    # Use the same Brazilian reporting day as the dashboards, even after
    # midnight UTC. Freeze it for this entire collection.
    reporting_day = reporting_today()
    hoje = reporting_day.isoformat()
    # 1) Descobre o advertiser_id no formato novo da API
    try:
        r = requests.get(API + "/advertising/advertisers",
                         headers={"Authorization": "Bearer " + token,
                                  "Api-Version": "1"},
                         params={"product_id": "PADS"}, timeout=30)
    except Exception as e:
        registrar_erro(conexao_id, "ads", "advertiser: " + repr(e)[:150])
        return False
    if r.status_code != 200:
        registrar_erro(conexao_id, "ads",
                       "advertiser HTTP " + str(r.status_code)
                       + ": " + r.text[:120])
        return False
    advs = r.json()
    if isinstance(advs, dict):
        advs = advs.get("advertisers") or advs.get("results") or []
    if not advs:
        registrar_erro(conexao_id, "ads",
                       "conta sem Product Ads ativo (PADS)")
        return False
    adv_id = str(advs[0].get("advertiser_id"))
    site = str(advs[0].get("site_id") or "MLB")
    url = (API + "/advertising/" + site + "/advertisers/" + adv_id
           + "/product_ads/campaigns/search")
    cab = {"Authorization": "Bearer " + token, "api-version": "2"}

    def buscar_dia(dia):
        # Campanhas do dia com metricas (gasto e faturamento)
        offset = 0
        tot_g = tot_f = 0.0
        tot_i = tot_c = 0
        linhas = []
        while True:
            try:
                r = requests.get(url, headers=cab,
                                 params={"limit": 50, "offset": offset,
                                         "date_from": dia, "date_to": dia,
                                         "metrics": "clicks,prints,cost,"
                                                    "total_amount,advertising_items_quantity,"
                                                    "organic_items_quantity,organic_units_amount"},
                                 timeout=30)
            except Exception as e:
                registrar_erro(conexao_id, "ads",
                               dia + ": " + repr(e)[:120])
                return None
            if r.status_code != 200:
                registrar_erro(conexao_id, "ads",
                               dia + " HTTP " + str(r.status_code)
                               + ": " + r.text[:120])
                return None
            dados = r.json()
            resultados = dados.get("results") or []
            for c in resultados:
                m = c.get("metrics") or {}
                g = float(m.get("cost") or 0)
                f = float(m.get("total_amount") or 0)
                tot_i += int(m.get("prints") or 0)
                tot_c += int(m.get("clicks") or 0)
                tot_g += g
                tot_f += f
                linhas.append((str(c.get("id")), g, f, m))
            offset += 50
            total = (dados.get("paging") or {}).get("total") or 0
            if not resultados or offset >= total:
                break
        return (dia, tot_g, tot_f, tot_i, tot_c, linhas)

    # Current campaign configuration (budget and target ROAS) is separate
    # from the daily performance snapshot. Keep last known config if API fails.
    config_complete = True
    try:
        campanhas_atuais = []
        deslocamento = 0
        while True:
            resposta = requests.get(url, headers=cab,
                                    params={"limit": 50, "offset": deslocamento},
                                    timeout=30)
            if resposta.status_code != 200:
                raise ValueError("HTTP " + str(resposta.status_code))
            lote = resposta.json()
            registros = lote.get("results") or []
            campanhas_atuais.extend(registros)
            deslocamento += len(registros)
            total = (lote.get("paging") or {}).get("total") or 0
            if not registros or deslocamento >= total or len(registros) < 50:
                break
        with banco() as conn:
            for campanha in campanhas_atuais:
                camp_id = campanha.get("id")
                if camp_id is None:
                    continue
                budget = campanha.get("daily_budget")
                if budget is None:
                    budget = campanha.get("budget")
                target = campanha.get("roas_target")
                from ads_change_history import record_changes
                anterior = consultar(conn, """SELECT orcamento, roas_objetivo
                    FROM ads_campanhas WHERE conexao_id=%s AND campanha_id=%s
                    FOR UPDATE""", (conexao_id, str(camp_id)))
                record_changes(conn, executar, conexao_id, camp_id,
                    anterior[0] if anterior else None,
                    {"orcamento":budget,"roas_objetivo":target},
                    int(time.time()), "api")
                executar(conn, """
                    INSERT INTO ads_campanhas (
                        conexao_id, campanha_id, nome, status, tipo,
                        data_inicio, data_fim, orcamento, roas_objetivo, atualizado_em
                    ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                    ON CONFLICT (conexao_id,campanha_id) DO UPDATE SET
                        nome=EXCLUDED.nome, status=EXCLUDED.status,
                        tipo=EXCLUDED.tipo, orcamento=EXCLUDED.orcamento,
                        roas_objetivo=COALESCE(EXCLUDED.roas_objetivo,ads_campanhas.roas_objetivo),
                        atualizado_em=EXCLUDED.atualizado_em
                """, (conexao_id, str(camp_id), campanha.get("name"),
                      campanha.get("status"), campanha.get("strategy"),
                      campanha.get("date_created"), None, budget, target,
                      int(time.time())))
    except (requests.RequestException, ValueError, TypeError) as exc:
        config_complete = False
        registrar_erro(conexao_id, "ads_config", "Campanhas: " + str(exc)[:140])

    # 2) Hoje (igual a antes)
    res_hoje = buscar_dia(hoje)
    if res_hoje is None:
        return False
    _, tot_g, tot_f, tot_i, tot_c, linhas = res_hoje
    agora = int(time.time())
    with banco() as conn:
        for camp_id, g, f, metricas in linhas:
            save_day(conn, executar, conexao_id, camp_id, hoje, g, f, metricas, agora)
            executar(conn, """INSERT INTO ads_orcamento_esgotado(conexao_id,campanha_id,data,detectado_em)
                SELECT %s,%s,%s,%s FROM ads_campanhas WHERE conexao_id=%s AND campanha_id=%s
                AND orcamento>0 AND %s>=orcamento
                ON CONFLICT(conexao_id,campanha_id,data) DO NOTHING""",
                (conexao_id,str(camp_id),hoje,agora,conexao_id,str(camp_id),g))
        executar(conn,
                 """INSERT INTO ads_dia
                    (conexao_id, data, gasto, faturamento, impressoes,
                     cliques, atualizado_em)
                    VALUES (%s,%s,%s,%s,%s,%s,%s)
                    ON CONFLICT (conexao_id, data)
                    DO UPDATE SET gasto=EXCLUDED.gasto,
                      faturamento=EXCLUDED.faturamento,
                      impressoes=EXCLUDED.impressoes,
                      cliques=EXCLUDED.cliques, atualizado_em=EXCLUDED.atualizado_em""",
                 (conexao_id, hoje, tot_g, tot_f, tot_i, tot_c, agora))

    # 3) Fill the initial historical window and repair recent missing days
    # after it has 60+ records. A missed daily collection must not become
    # permanent. Keep API work bounded for established stores.
    from ads_coverage import daily_days_to_fetch
    with banco() as conn:
        reg = consultar(conn, "SELECT data FROM ads_dia WHERE conexao_id=%s AND data BETWEEN %s AND %s",
                        (conexao_id, (reporting_day-datetime.timedelta(days=89)).isoformat(),
                         (reporting_day-datetime.timedelta(days=1)).isoformat()))
    days_to_fetch = daily_days_to_fetch(
        [r["data"] for r in reg], reporting_day,
        initial=len(reg)<60, max_recent=3)
    # Enrich old snapshots gradually without an unbounded API backfill.
    # Mark an attempted successful response even if a metric is unavailable.
    with banco() as conn:
        pending = consultar(conn, """SELECT DISTINCT data FROM ads_campanha_dia
            WHERE conexao_id=%s AND data BETWEEN %s AND %s
            AND metricas_vendas_coletadas_em IS NULL ORDER BY data DESC LIMIT 3""",
            (conexao_id, (reporting_day-datetime.timedelta(days=89)).isoformat(),
             (reporting_day-datetime.timedelta(days=1)).isoformat()))
    days_to_fetch = list(dict.fromkeys(days_to_fetch + [r['data'] for r in pending]))
    if refresh_range is not None:
        # Explicit refresh repairs the period visible on the Ads page only.
        start, end = refresh_range
        days_to_fetch = [(end-datetime.timedelta(days=i)).isoformat()
                         for i in range((end-start).days+1)
                         if (end-datetime.timedelta(days=i)).isoformat() != hoje]
    history_complete = True
    if days_to_fetch:
        for dia_b in days_to_fetch:
            res = buscar_dia(dia_b)
            if res is None:
                history_complete = False
                continue
            _, g2, f2, i2, c2, linhas2 = res
            with banco() as conn:
                for camp_id, g, f, metricas in linhas2:
                    save_day(conn, executar, conexao_id, camp_id, dia_b, g, f, metricas, agora)
                executar(conn,
                         """INSERT INTO ads_dia
                            (conexao_id, data, gasto, faturamento, impressoes,
                             cliques, atualizado_em)
                            VALUES (%s,%s,%s,%s,%s,%s,%s)
                            ON CONFLICT (conexao_id, data)
                            DO UPDATE SET gasto=EXCLUDED.gasto,
                              faturamento=EXCLUDED.faturamento,
                              impressoes=EXCLUDED.impressoes,
                              cliques=EXCLUDED.cliques, atualizado_em=EXCLUDED.atualizado_em""",
                         (conexao_id, dia_b, g2, f2, i2, c2, agora))
    return (history_complete and config_complete) if refresh_range is not None else True

def puxar_pedidos(conexao_id, token, user_id):
    return sync_orders(conexao_id, token, user_id, banco, consultar,
                       executar, api_get, registrar_erro)


def puxar_full(conexao_id, token, user_id):
    return sync_full(conexao_id, token, banco, consultar, executar, requests.get)


def executar_atualizacao(cid):
    with store_sync_lock(banco, cid) as acquired:
        if not acquired:
            return False
        with banco() as conn:
            rows = consultar(conn, "SELECT * FROM conexoes WHERE id=%s", (cid,))
        if not rows:
            return False
        c = rows[0]
        token = token_atual(c)
        if not token:
            return False
        started = int(time.time())
        user_id = c["user_id"]
        stages = [("conta", lambda: puxar_conta(cid, token))]
        if user_id:
            stages.extend([
                # Atualizar pedidos primeiro: a fila é serial e anúncios podem demorar.
                ("pedidos", lambda: puxar_pedidos(cid, token, user_id)),
                ("anuncios", lambda: puxar_anuncios(cid, token, user_id)),
                ("full", lambda: puxar_full(cid, token, user_id)),
                ("perguntas", lambda: puxar_perguntas(cid, token, user_id)),
                ("envios", lambda: puxar_envios(cid, token, user_id)),
                ("promocoes", lambda: puxar_promocoes(cid, token, user_id)),
                ("ads", lambda: puxar_ads(cid, token, user_id)),
            ])
        # puxar_ads already refreshes daily metrics. No second/undefined collector.
        # Disabled visit collection is not reported as a successful refresh.
        stages.append(("historico", lambda: (gravar_historico(cid), True)[1]))
        falhas = []
        def coletar_etapa(nome, coleta):
            resultado = coleta()
            if resultado is not True:
                falhas.append(nome)
            return resultado
        def registrar_falha(conexao_id, categoria, mensagem):
            falhas.append(categoria)
            registrar_erro(conexao_id, categoria, mensagem)
        etapas = [(nome, lambda nome=nome, coleta=coleta: coletar_etapa(nome, coleta)) for nome, coleta in stages]
        ok = run_stages(cid, etapas, registrar_falha, resolver_erro, time.time)
        if ok:
            resolver_erro(cid, "atualizacao", started)
        else:
            registrar_erro(cid, "atualizacao", "Etapas com falha: " + ", ".join(dict.fromkeys(falhas)))
        return ok


def executar_atualizacao_bg(cid):
    try:
        executar_atualizacao(cid)
    except Exception as e:
        try:
            registrar_erro(cid, "atualizacao", "Falha na atualização: " + type(e).__name__)
        except Exception:
            pass
    finally:
        with LOCK_ATUALIZACAO:
            ATUALIZANDO.discard(cid)


def job_dados():
    # Recover expired durable leases before scheduling more work.
    # The legacy worker remains active until durable processing is enabled.
    try:
        recuperar_tarefas_persistentes()
    except Exception:
        pass  # Queue maintenance must not block store synchronization.
    # Mark abandoned execution records even if the process did not restart.
    # Do not let monitoring errors prevent normal sync scheduling.
    try:
        with banco() as conn:
            agora = int(time.time())
            executar(conn, """UPDATE sincronizacao_execucoes
                SET status='interrompida', fim_em=%s,
                    detalhe='Execucao sem conclusao por mais de 2 horas'
                WHERE status='executando' AND COALESCE(heartbeat_em, inicio_em) < %s""",
                (agora, agora - 7200))
    except Exception:
        pass
    with banco() as conn:
        conexoes = consultar(conn, """SELECT c.id FROM conexoes c
            WHERE c.status='ativa' AND c.cliente <> %s
            ORDER BY CASE WHEN EXISTS (
                SELECT 1 FROM generate_series(
                    (CURRENT_TIMESTAMP AT TIME ZONE 'America/Sao_Paulo')::date-14,
                    (CURRENT_TIMESTAMP AT TIME ZONE 'America/Sao_Paulo')::date-1,
                    interval '1 day') d(dia)
                WHERE NOT EXISTS (SELECT 1 FROM vendas_cobertura v
                    WHERE v.conexao_id=c.id AND v.dia=d.dia::date::text)
            ) THEN 0 ELSE 1 END, c.id""", ("CASAMIX_STORE",))
    for c in conexoes:
        # Use the existing queue so the scheduler does not abort subsequent shops
        # or overlap a collection already queued from the dashboard.
        enfileirar_atualizacao(c["id"])


register_crm(app, banco, consultar, executar, pagina)
register_mfa(app, banco, consultar, executar, pagina, auth_attempt_allowed)
register_full(app, login_required, pode_ver_loja, banco, consultar, pagina, pagina_loja, session)
register_competitors(app, login_required, pode_ver_loja, banco, consultar, executar, pagina, pagina_loja)
from intervention_routes import register_interventions
register_interventions(app, login_required, pode_ver_loja, banco, consultar, executar, pagina_loja)
from anuncios_progress_routes import register_progress
register_progress(app, login_required, pagina)
from listing_task_routes import register_tasks
register_tasks(app, login_required, pode_ver_loja, banco, consultar, executar, pagina_loja)
from listing_content_routes import register_content
register_content(app, login_required, pode_ver_loja, banco, consultar, executar, pagina_loja)
from listing_questions import register_questions
register_questions(app, login_required, pode_ver_loja, banco, consultar, executar, pagina_loja)
from profitability_routes import register_profitability
register_profitability(app, login_required, pode_ver_loja, banco, consultar, executar, pagina_loja)
register_management(app, login_required, pode_ver_loja, banco, consultar, pagina, pagina_loja,
                    is_updating=lambda cid: cid in ATUALIZANDO)

from section_refresh import register as register_section_refresh
register_section_refresh(app, login_required, pode_ver_loja, banco, consultar, executar, coletar_secao)

from startup_schema import initialize_with_retry
initialize_with_retry(iniciar_banco, app.logger)

# Official item visits: unique item/day records, initialized without backfilling
# or inventing missing days. Collection is activated separately.
with banco() as conn:
    executar(conn, """CREATE TABLE IF NOT EXISTS item_visits_daily (
        conexao_id INTEGER NOT NULL REFERENCES conexoes(id),
        item_id TEXT NOT NULL,
        dia DATE NOT NULL,
        visitas INTEGER NOT NULL CHECK (visitas >= 0),
        atualizado_em TIMESTAMPTZ NOT NULL DEFAULT NOW(),
        PRIMARY KEY (conexao_id, item_id, dia)
    )""")
    executar(conn, """CREATE INDEX IF NOT EXISTS idx_item_visits_daily_shop_day
        ON item_visits_daily (conexao_id, dia)""")

# Persist execution status so interrupted web workers remain diagnosable.
with banco() as conn:
    executar(conn, """CREATE TABLE IF NOT EXISTS sincronizacao_execucoes (
        id BIGSERIAL PRIMARY KEY,
        conexao_id INTEGER NOT NULL,
        inicio_em BIGINT NOT NULL,
        fim_em BIGINT,
        status TEXT NOT NULL,
        detalhe TEXT
    )""")
    executar(conn, "ALTER TABLE sincronizacao_execucoes ADD COLUMN IF NOT EXISTS heartbeat_em BIGINT")
    executar(conn, """UPDATE sincronizacao_execucoes
        SET status='interrompida', fim_em=%s,
            detalhe='Processo reiniciado antes da conclusão'
        WHERE status='executando' AND COALESCE(heartbeat_em, inicio_em) < %s""",
        (int(time.time()), int(time.time()) - 7200))

# Future durable queue: schema only. Current worker remains unchanged.
with banco() as conn:
    executar(conn, """CREATE TABLE IF NOT EXISTS sincronizacao_tarefas (
        id BIGSERIAL PRIMARY KEY,
        conexao_id INTEGER NOT NULL REFERENCES conexoes(id),
        status TEXT NOT NULL DEFAULT 'pendente'
            CHECK (status IN ('pendente','executando','concluida','falhou')),
        tentativas INTEGER NOT NULL DEFAULT 0,
        disponivel_em BIGINT NOT NULL,
        criada_em BIGINT NOT NULL,
        iniciada_em BIGINT,
        concluida_em BIGINT,
        lease_ate BIGINT,
        ultimo_erro TEXT
    )""")
    executar(conn, """CREATE INDEX IF NOT EXISTS idx_sync_tarefas_pendentes
        ON sincronizacao_tarefas (disponivel_em, id)
        WHERE status='pendente'""")
    executar(conn, """CREATE INDEX IF NOT EXISTS idx_sync_tarefas_leases
        ON sincronizacao_tarefas (lease_ate)
        WHERE status='executando'""")

def criar_tarefa_persistente(cid, agora=None):
    """Enqueue one pending job; duplicate pending jobs are harmless at this stage."""
    agora = int(time.time()) if agora is None else int(agora)
    with banco() as conn:
        cur = executar(conn, """INSERT INTO sincronizacao_tarefas
            (conexao_id, disponivel_em, criada_em)
            VALUES (%s,%s,%s) ON CONFLICT DO NOTHING RETURNING id""",
            (cid, agora, agora))
        row = cur.fetchone()
        return row[0] if row else None


def reservar_tarefa_persistente(agora=None, duracao_lease=7200):
    """Atomically claim a ready job across workers using SKIP LOCKED."""
    lojas = [int(part.strip()) for part in os.getenv("DURABLE_QUEUE_STORE_IDS", "").split(",")
             if part.strip().isdigit()]
    if not lojas:
        return None
    agora = int(time.time()) if agora is None else int(agora)
    with banco() as conn:
        rows = consultar(conn, """SELECT t.id, t.conexao_id
            FROM sincronizacao_tarefas t
            WHERE t.status='pendente' AND t.disponivel_em <= %s
              AND t.conexao_id = ANY(%s)
              AND EXISTS (SELECT 1 FROM conexoes c WHERE c.id=t.conexao_id AND c.status='ativa')
              AND NOT EXISTS (
                SELECT 1 FROM sincronizacao_tarefas a
                WHERE a.conexao_id=t.conexao_id AND a.status='executando'
              )
            ORDER BY t.disponivel_em,t.id
            FOR UPDATE OF t SKIP LOCKED LIMIT 1""", (agora, lojas))
        if not rows:
            return None
        tarefa = rows[0]
        try:
            executar(conn, """UPDATE sincronizacao_tarefas
                SET status='executando', tentativas=tentativas+1,
                    iniciada_em=%s, lease_ate=%s
                WHERE id=%s""", (agora, agora+duracao_lease, tarefa["id"]))
        except psycopg2.errors.UniqueViolation:
            conn.rollback()
            return None  # Another worker claimed a task for this store.
        return tarefa


def renovar_lease_tarefa_persistente(tarefa_id, agora=None, duracao_lease=7200):
    """Extend the lease only while the task remains executing."""
    agora = int(time.time()) if agora is None else int(agora)
    with banco() as conn:
        cur = executar(conn, """UPDATE sincronizacao_tarefas
            SET lease_ate=GREATEST(lease_ate, %s)
            WHERE id=%s AND status='executando' AND lease_ate >= %s""",
            (agora + duracao_lease, tarefa_id, agora))
        return cur.rowcount == 1


def finalizar_tarefa_persistente(tarefa_id, sucesso, erro=None, agora=None):
    agora = int(time.time()) if agora is None else int(agora)
    with banco() as conn:
        executar(conn, """UPDATE sincronizacao_tarefas
            SET status=%s, concluida_em=%s, lease_ate=NULL,
                ultimo_erro=%s WHERE id=%s AND status='executando'""",
            ("concluida" if sucesso else "falhou", agora,
             (erro or "")[:300], tarefa_id))


def recuperar_tarefas_persistentes(agora=None, max_tentativas=3):
    """Requeue expired leases with exponential backoff or mark exhausted jobs failed."""
    agora = int(time.time()) if agora is None else int(agora)
    with banco() as conn:
        cur = executar(conn, """UPDATE sincronizacao_tarefas
            SET status=CASE WHEN tentativas >= %s OR EXISTS (SELECT 1 FROM sincronizacao_tarefas outra WHERE outra.conexao_id=sincronizacao_tarefas.conexao_id AND outra.status='pendente' AND outra.id<>sincronizacao_tarefas.id) THEN 'falhou' ELSE 'pendente' END,
                disponivel_em=CASE WHEN tentativas >= %s THEN disponivel_em
                    ELSE %s + LEAST(3600, 60 * POWER(2, LEAST(tentativas, 6))::INTEGER) END,
                lease_ate=NULL,
                concluida_em=CASE WHEN tentativas >= %s THEN %s ELSE NULL END,
                ultimo_erro='Lease expirado; worker interrompido'
            WHERE status='executando' AND lease_ate < %s""",
            (max_tentativas, max_tentativas, agora, max_tentativas, agora, agora))
        return cur.rowcount


def reagendar_tarefa_persistente(tarefa_id, erro, agora=None, max_tentativas=3):
    """Retry failed work after backoff, with a hard retry cap."""
    agora = int(time.time()) if agora is None else int(agora)
    with banco() as conn:
        cur = executar(conn, """UPDATE sincronizacao_tarefas
            SET status=CASE WHEN tentativas >= %s OR EXISTS (SELECT 1 FROM sincronizacao_tarefas outra WHERE outra.conexao_id=sincronizacao_tarefas.conexao_id AND outra.status='pendente' AND outra.id<>sincronizacao_tarefas.id) THEN 'falhou' ELSE 'pendente' END,
                disponivel_em=CASE WHEN tentativas >= %s THEN disponivel_em
                    ELSE %s + LEAST(3600, 60 * POWER(2, LEAST(tentativas, 6))::INTEGER) END,
                lease_ate=NULL,
                concluida_em=CASE WHEN tentativas >= %s THEN %s ELSE NULL END,
                ultimo_erro=%s
            WHERE id=%s AND status='executando'""",
            (max_tentativas, max_tentativas, agora, max_tentativas,
             agora, str(erro)[:300], tarefa_id))
        return cur.rowcount


# Avoid accumulating duplicate pending tasks from repeated scheduler runs.
with banco() as conn:
    executar(conn, """CREATE UNIQUE INDEX IF NOT EXISTS
        idx_sync_tarefas_uma_pendente_por_loja
        ON sincronizacao_tarefas (conexao_id)
        WHERE status='pendente'""")

# Database-level invariant: at most one running durable task per store.
with banco() as conn:
    executar(conn, """CREATE UNIQUE INDEX IF NOT EXISTS
        idx_sync_tarefas_uma_execucao_por_loja
        ON sincronizacao_tarefas (conexao_id)
        WHERE status='executando'""")

def recuperar_tarefas_sem_coletor():
    """Recover interrupted selected stores only while owning their collection lock."""
    lojas = {int(part.strip()) for part in os.getenv("DURABLE_QUEUE_STORE_IDS", "").split(",")
             if part.strip().isdigit()}
    recuperadas = 0
    for cid in sorted(lojas):
        with store_sync_lock(banco, cid) as acquired:
            if not acquired:
                continue
            with banco() as conn:
                tarefas = consultar(conn, """SELECT id FROM sincronizacao_tarefas
                    WHERE conexao_id=%s AND status='executando'""", (cid,))
            for tarefa in tarefas:
                recuperadas += reagendar_tarefa_persistente(
                    tarefa["id"], "Coletor interrompido; retomada sob trava exclusiva")
            if tarefas:
                encerrar_execucoes_abandonadas(cid)
    return recuperadas


def job_fila_persistente():
    """Drain one durable task per tick; disabled unless explicitly enabled."""
    if os.getenv("DURABLE_QUEUE_WORKER_ENABLED", "").lower() != "true":
        return
    try:
        recuperar_tarefas_sem_coletor()
        processar_tarefa_persistente_uma()
    except Exception:
        # Keep legacy scheduling alive even if the durable queue is unavailable.
        app.logger.exception("Falha no worker da fila persistente")


def job_visitas_todas_contas():
    """One small rotating batch per run, protected against parallel workers."""
    if os.getenv("VISITS_AUTO_ENABLED", "0") != "1":
        return
    import psycopg2
    from copas_visits_store import collect
    lock_conn = None
    try:
        lock_conn = psycopg2.connect(os.environ["DATABASE_URL"], connect_timeout=10)
        lock_conn.autocommit = True
        with lock_conn.cursor() as cur:
            cur.execute("SELECT pg_try_advisory_lock(%s)", (880142,))
            if not cur.fetchone()[0]:
                return
            cur.execute("""
                CREATE TABLE IF NOT EXISTS item_visits_checked (
                    conexao_id INTEGER NOT NULL, item_id TEXT NOT NULL,
                    consultado_em TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                    PRIMARY KEY (conexao_id, item_id)
                )
            """)
            from visits_refresh import DUE_STORE_SQL, REFRESH_HOURS
            cur.execute(DUE_STORE_SQL, (REFRESH_HOURS,))
            record = cur.fetchone()
            if record:
                # Reuse the same refresh path as store synchronization before
                # the collector opens its own connection and reads the token.
                with banco() as conn:
                    contas = consultar(conn, "SELECT * FROM conexoes WHERE id=%s AND status='ativa'", (record[0],))
                if not contas or not token_atual(contas[0]):
                    app.logger.warning("visits_batch_auth_unavailable connection_id=%s",record[0])
                    return
                resultado = collect(limit=40, connection_id=record[0], rotate=True)
                if resultado.get("falhas"):
                    app.logger.warning("visits_batch_result %s",resultado)
                else:
                    app.logger.info("visits_batch_result %s",resultado)
    except Exception:
        app.logger.exception("visits_all_accounts_batch_failed")
    finally:
        if lock_conn is not None:
            try:
                with lock_conn.cursor() as cur:
                    cur.execute("SELECT pg_advisory_unlock(%s)", (880142,))
            except Exception:
                pass
            lock_conn.close()


def job_sincronizacao_lojas_novas():
    """Queue newly connected stores without local items, avoiding repeated jobs."""
    try:
        with banco() as conn:
            lojas = consultar(conn, """
                SELECT c.id FROM conexoes c
                WHERE c.status='ativa' AND c.access_token IS NOT NULL
                  AND c.user_id IS NOT NULL
                  AND NOT EXISTS (SELECT 1 FROM anuncios a WHERE a.conexao_id=c.id)
                  AND NOT EXISTS (
                      SELECT 1 FROM sincronizacao_execucoes e
                      WHERE e.conexao_id=c.id AND e.status='executando'
                  )
                ORDER BY c.id
                LIMIT 3
            """)
        for loja in lojas:
            enfileirar_atualizacao(loja["id"])
    except Exception:
        app.logger.exception("new_store_initial_sync_failed")


def job_concorrencia():
    try:
        collect_due(banco, consultar, executar, requests.get)
    except Exception:
        app.logger.exception("competitor_collection_failed")


def job_token_encryption():
    try:
        count = migrate_tokens(banco)
        if count:
            app.logger.info("ml_tokens_encrypted count=%s", count)
    except Exception:
        app.logger.error("ml_token_encryption_failed")


def job_keiper_ai():
    try:
        from keiper_ai_store import run_one
        run_one(banco, consultar, executar, app.logger)
    except Exception:
        app.logger.warning('keiper_ai_daily_unavailable')


def job_kyra_campaigns():
    try:
        from kyra_campaign_store import run_one
        run_one(banco, consultar, executar, app.logger)
    except Exception:
        app.logger.warning('kyra_campaign_daily_unavailable')


scheduler = BackgroundScheduler()
scheduler.add_job(job_kyra_campaigns, 'interval', minutes=1,
                  next_run_time=datetime.datetime.now() + datetime.timedelta(seconds=30),
                  id='kyra_campaign_daily', max_instances=1, coalesce=True)
scheduler.add_job(job_keiper_ai, 'interval', minutes=1,
                  id='keiper_ai_daily', max_instances=1, coalesce=True)
scheduler.add_job(job_token_encryption, "interval", minutes=1,
                  next_run_time=datetime.datetime.now() + datetime.timedelta(seconds=10),
                  max_instances=1, coalesce=True)
scheduler.add_job(job_concorrencia, "interval", minutes=1,
                  max_instances=1, coalesce=True)
scheduler.add_job(job_sincronizacao_lojas_novas, "interval", minutes=5,
                  max_instances=1, coalesce=True)
scheduler.add_job(job_visitas_todas_contas, "interval", minutes=1,
                  max_instances=1, coalesce=True)
scheduler.add_job(job_refresh, "interval", minutes=30)
scheduler.add_job(job_fila_persistente, "interval", minutes=1,
                  max_instances=1, coalesce=True)
scheduler.add_job(job_dados, "interval", minutes=60,
                  next_run_time=datetime.datetime.now() + datetime.timedelta(minutes=2),
                  max_instances=1, coalesce=True)
scheduler.start()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", 8000)))
