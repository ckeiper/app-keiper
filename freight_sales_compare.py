"""Compare report charges to the same freight calculation used by Vendas.

By default only aggregates leave this module. An explicit details list receives
minimal per-order evidence for the authorized store-scoped freight review.
Multiple movements per order are left unresolved rather than added or deduplicated.
"""
import io
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from freight_reconciliation import schema_columns, normalize, _amount
from freight_comparison import compare_freight, format_brl, money


def compare_sales_rows(rows, columns, references, payment_links, shipment_links, details=None):
    links = defaultdict(set)
    for link in payment_links:
        links[('operation', str(link['operation_id']))].add(str(link['pedido_id']))
    for link in shipment_links:
        links[('shipment', str(link['envio_id']))].add(str(link['pedido_id']))
    grouped = defaultdict(list)
    result = dict(candidate_rows=0, unmatched_rows=0, ambiguous_rows=0,
                  ambiguous_orders=0, unavailable_orders=0, invalid_orders=0,
                  compared_orders=0, real_orders=0, estimated_orders=0,
                  above=0, below=0, within=0)
    paid = expected = excess = below_total = Decimal('0.00')
    for row in rows:
        if normalize(row[columns['operation_type']]) not in ('venda', 'envio'):
            continue
        result['candidate_rows'] += 1
        operation = row[columns['operation_id']]
        shipment = row[columns['shipment_id']]
        candidates = (links.get(('operation', str(operation)), set())
                      | links.get(('shipment', str(shipment)), set())
                      | ({str(operation)} if str(operation) in references else set()))
        if not candidates:
            result['unmatched_rows'] += 1
            continue
        if len(candidates) != 1:
            result['ambiguous_rows'] += 1
            # Poison every candidate: a different singleton row cannot make this safe.
            for order_id in candidates:
                grouped[order_id].extend([None, None])
            continue
        grouped[next(iter(candidates))].append(row)
    for order_id, movements in grouped.items():
        if len(movements) != 1:
            result['ambiguous_orders'] += 1
            continue
        reference = references.get(order_id)
        if (not reference or not reference.get('frete_verificavel') or
                reference.get('frete_fonte') not in ('real', 'estimado')):
            result['unavailable_orders'] += 1
            continue
        row = movements[0]
        try:
            gross, discount, charged = [_amount(row[columns[key]]) for key in
                ('gross_shipping','shipping_discount','seller_shipping')]
            baseline = money(reference['frete'])
            if discount > gross or abs(gross - discount - charged) > Decimal('.01'):
                raise ValueError
        except (ValueError, InvalidOperation, KeyError):
            result['invalid_orders'] += 1
            continue
        comparison = compare_freight(charged, baseline)
        status = {'acima':'above','abaixo':'below','dentro_tolerancia':'within'}[comparison['status']]
        result[status] += 1
        result['compared_orders'] += 1
        result['real_orders' if reference['frete_fonte'] == 'real' else 'estimated_orders'] += 1
        paid += charged
        expected += baseline
        if comparison['status'] == 'acima':
            excess += charged - baseline
        elif comparison['status'] == 'abaixo':
            below_total += baseline - charged
        if details is not None and status != 'within':
            methods=[]
            if str(row[columns['operation_id']])==order_id:methods.append('pedido')
            if order_id in links.get(('operation',str(row[columns['operation_id']])),set()):methods.append('pagamento')
            if order_id in links.get(('shipment',str(row[columns['shipment_id']])),set()):methods.append('envio')
            shipment=row[columns['shipment_id']]
            details.append(dict(order_id=order_id,shipment_id=str(shipment) if shipment is not None else '',
                charged=format(charged,'.2f'),baseline=format(baseline,'.2f'),delta=format(charged-baseline,'.2f'),
                source=reference['frete_fonte'],status=status,matched_by=', '.join(methods),
                sale_date=str(reference.get('fechado_em') or '')))
    result.update(charged_total=str(paid),sales_total=str(expected),above_total=str(excess),below_total=str(below_total))
    return result


def audit_against_sales(cid, filename, payload, summary, banco, query, calculate,
                        parsed_rows=None, parsed_columns=None, details=None):
    if not summary.get('period_start'):
        return compare_sales_rows([],{}, {},[],[])
    _, sales = calculate(cid, inicio=summary['period_start'])
    refs = {str(row['pedido_id']): row for row in sales}
    payments, shipments = [], []
    if refs:
        with banco() as conn:
            payments = query(conn, '''SELECT pedido_id,payment->>'id' AS operation_id
                FROM pedidos CROSS JOIN LATERAL jsonb_array_elements(
                    CASE WHEN jsonb_typeof(pagamentos_json)='array' THEN pagamentos_json
                    ELSE '[]'::jsonb END) AS payment
                WHERE conexao_id=%s AND pedido_id=ANY(%s) AND payment->>'id' IS NOT NULL''',
                (cid,list(refs)))
            shipments = query(conn, '''SELECT pedido_id,envio_id FROM envios
                WHERE conexao_id=%s AND pedido_id=ANY(%s)''',(cid,list(refs)))
    per_order = defaultdict(set)
    for link in shipments:
        per_order[str(link['pedido_id'])].add(str(link['envio_id']))
    for order_id, identifiers in per_order.items():
        if len(identifiers) > 1 and order_id in refs:
            refs[order_id]['frete_verificavel'] = False
    if parsed_rows is not None and parsed_columns is not None:
        return compare_sales_rows(parsed_rows,parsed_columns,refs,payments,shipments,details=details)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(payload),read_only=True,data_only=False,keep_links=False)
    try:
        ws = wb.active
        columns = schema_columns(next(ws.iter_rows(min_row=3,max_row=3,max_col=45,values_only=True)))
        return compare_sales_rows(ws.iter_rows(min_row=4,max_col=45,values_only=True),
                                  columns,refs,payments,shipments,details=details)
    finally:
        wb.close()


def render_sales_comparison(result):
    """Highlight investigation priorities rather than presenting ten unranked counts."""
    def tile(label, number):
        return ('<div class="frete-metrica"><span>' + label + '</span><strong>'
                + str(int(number)) + '</strong></div>')
    attention = int(result['above']) + int(result['below'])
    return (
        '<section class="card frete-resultado"><h3>Comparação com Vendas</h3>'
        '<p>As divergências indicam pedidos para investigar, não cobranças indevidas comprovadas.</p>'
        '<div class="frete-metricas">'
        + tile('Pedidos comparados', result['compared_orders'])
        + tile('Com diferença', attention)
        + tile('Dentro da tolerância', result['within'])
        + tile('Sem vínculo na base', result['unmatched_rows'])
        + '</div><div class="frete-metricas">'
        + '<div class="frete-metrica"><span>Frete na planilha</span><strong>'
        + format_brl(result['charged_total']) + '</strong></div>'
        + '<div class="frete-metrica"><span>Frete em Vendas</span><strong>'
        + format_brl(result['sales_total']) + '</strong></div>'
        + '<div class="frete-metrica"><span>Diferenças positivas</span><strong>'
        + format_brl(result['above_total']) + '</strong></div>'
        + '</div><details class="frete-detalhes"><summary>Ver critérios e pendências da comparação</summary><ul>'
        + ''.join('<li>' + label + ': <b>' + str(int(result[key])) + '</b></li>'
          for label, key in (
              ('Referência real', 'real_orders'),
              ('Referência estimada', 'estimated_orders'),
              ('Planilha acima de Vendas', 'above'),
              ('Planilha abaixo de Vendas', 'below'),
              ('Múltiplos lançamentos ou vínculo ambíguo', 'ambiguous_orders'),
              ('Sem frete verificável em Vendas', 'unavailable_orders'),
              ('Valores de frete inválidos', 'invalid_orders')))
        + '</ul><p>Vínculo por pedido, pagamento ou envio, dentro da mesma loja. '
          'Custos de Vendas podem ser reais ou estimados. Tolerância de R$ 0,01. '
          'Diferenças não comprovam cobrança indevida.</p></details></section>'
    )
