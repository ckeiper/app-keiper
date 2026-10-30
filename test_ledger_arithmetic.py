import unittest
from ledger_arithmetic import empty,accumulate,render,signed_amount
import test_freight_reconciliation as fixtures
from freight_reconciliation import inspect_sales_reconciliation,FIELDS
from openpyxl import load_workbook
from io import BytesIO

class LedgerArithmeticTests(unittest.TestCase):
 def test_signed_adjustments_and_zero(self):
  r=empty(True)
  for row in [(100,10,90),(-100,-10,-90),(0,0,0)]:accumulate(r,*row)
  self.assertEqual(r['valid_rows'],3);self.assertEqual(r['inconsistent_rows'],0)
  self.assertEqual(r['net_total'],'0.00')
 def test_one_cent_and_offsetting_rows_stay_flagged(self):
  r=empty(True);accumulate(r,100,10,'89.99');accumulate(r,100,10,'90.01')
  self.assertEqual(r['inconsistent_rows'],2);self.assertEqual(r['difference_total'],'0.00')
 def test_missing_invalid_formula_and_fractional_cent(self):
  for v in [None,True,'','=1+2','NaN','Infinity','1.001','R$ 1,00','1e9999']:
   self.assertIsNone(signed_amount(v))
   r=empty(True);accumulate(r,v,1,1);self.assertEqual(r['invalid_rows'],1);self.assertEqual(r['valid_rows'],0)
 def test_no_new_result_for_legacy_and_no_zero_for_invalid(self):
  self.assertEqual(render(None),'')
  self.assertIn('indisponível',render(empty(False)))
  self.assertNotIn('<strong>',render(empty(True)))
 def test_duplicates_included_but_not_claimed_as_sales(self):
  r=empty(True);accumulate(r,100,10,90);accumulate(r,100,10,90)
  self.assertEqual(r['net_total'],'180.00')
  html=render(r);self.assertIn('incluindo repetições',html);self.assertIn('não são faturamento',html)
 def test_original_schema_optional_new_column_and_real_parser(self):
  helper=fixtures.SalesReconciliationTests();payload=helper.workbook([dict(total_fees=-10,net_after_fees=110)])
  old=inspect_sales_reconciliation('test.xlsx',payload);self.assertFalse(old['financial']['available'])
  w=load_workbook(BytesIO(payload));w.active.cell(3,len(FIELDS)+1,'Valor bruto (=)');w.active.cell(4,len(FIELDS)+1,100)
  out=BytesIO();w.save(out)
  r=inspect_sales_reconciliation('test.xlsx',out.getvalue())['financial']
  self.assertEqual(r['valid_rows'],1);self.assertEqual(r['inconsistent_rows'],0)
  self.assertEqual(r['fees_total'],'-10.00')
