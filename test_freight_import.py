import csv
import io
import unittest
from openpyxl import Workbook
from freight_import import (preview_reconciliation_upload,
                            render_reconciliation_preview, PreviewError)


class ReconciliationPreviewTests(unittest.TestCase):
    def test_xlsx_preview_does_not_guess_columns(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Conciliação"
        ws.append(["Relatório", None])
        ws.append(["Número da venda", "Repasse"])
        ws.append(["123", -12.34])
        data = io.BytesIO()
        wb.save(data)
        sheets = preview_reconciliation_upload("conciliacao.xlsx", data.getvalue())
        self.assertEqual(sheets[0]["name"], "Conciliação")
        self.assertEqual(sheets[0]["sample"][0][0], "Relatório")
        self.assertEqual(sheets[0]["sample"][1][0], "Número da venda")
        self.assertNotIn("divergencia", sheets[0])

    def test_csv_with_semicolon_and_html_is_escaped(self):
        sheets = preview_reconciliation_upload(
            "conciliação.csv", b"Venda;Valor\n123;<script>alert(1)</script>\n")
        html = render_reconciliation_preview(sheets)
        self.assertIn("&lt;script&gt;", html)
        self.assertNotIn("<script>", html)

    def test_rejects_unsupported_or_oversize(self):
        for filename, data in [("report.xls", b"old"), ("report.xlsx", b"invalid"),
                               ("report.csv", b""), ("report.csv", b"x" * (2 * 1024 * 1024 + 1))]:
            with self.subTest(filename=filename):
                with self.assertRaises(PreviewError):
                    preview_reconciliation_upload(filename, data)

    def test_multiple_sheets_and_blank_values(self):
        wb = Workbook()
        wb.active.append(["A", None])
        wb.create_sheet("Detalhes").append(["Venda", 2])
        data = io.BytesIO()
        wb.save(data)
        sheets = preview_reconciliation_upload("a.xlsx", data.getvalue())
        self.assertEqual(len(sheets), 2)
        self.assertEqual(sheets[0]["sample"][0][1], "")
        self.assertEqual(sheets[1]["sample"][0][0], "Venda")


if __name__ == "__main__":
    unittest.main()
